# -*- coding: utf-8 -*-
"""把主宣讲里「后续设想」页换成两栏版（小古文 / 主题阅读）。"""
import copy, os, sys
from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))          # <项目根>\写作课相关宣传文件\国庆后宣讲\脚本
ROOT = os.path.abspath(os.path.join(HERE, '..', '..', '..'))  # 按相对位置定位项目根，不写死盘符
sys.path.insert(0, os.path.join(ROOT, '.claude', 'skills', 'laojohn-ppt', 'tools'))
from pptx_text_edit import set_text, prune_timing, normalize_paragraphs

F = os.path.join(HERE, '..', '1-国庆后宣讲-主宣讲.pptx')
R_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
P_NS = '{http://schemas.openxmlformats.org/presentationml/2006/main}'

prs = Presentation(F)
slides = list(prs.slides)


def title_of(s):
    return '\n'.join(x.text_frame.text for x in s.shapes if x.has_text_frame)


old = [s for s in slides if '来自老师们的建议' in title_of(s)]
tpl = [s for s in slides if '先看现状：我们和同行' in title_of(s)]
assert len(old) == 1 and len(tpl) == 1
old, tpl = old[0], tpl[0]

new = prs.slides.add_slide(tpl.slide_layout)
for ph in list(new.placeholders):
    ph._element.getparent().remove(ph._element)
rid_map = {}
for rid, rel in tpl.part.rels.items():
    if rel.is_external or rel.reltype.endswith(('/slideLayout', '/notesSlide', '/tags')):
        continue
    rid_map[rid] = new.part.relate_to(rel._target, rel.reltype)
tree = new.shapes._spTree
for el in tpl.shapes._spTree:
    if el.tag.endswith('}nvGrpSpPr') or el.tag.endswith('}grpSpPr'):
        continue
    tree.append(copy.deepcopy(el))
for cd in list(tree.iter(P_NS + 'custDataLst')):
    cd.getparent().remove(cd)
for el in tree.iter():
    for attr in ('{%s}embed' % R_NS, '{%s}id' % R_NS):
        if el.get(attr) in rid_map:
            el.set(attr, rid_map[el.get(attr)])
bg = tpl._element.find('.//' + P_NS + 'bg')
if bg is not None:
    new._element.find(P_NS + 'cSld').insert(0, copy.deepcopy(bg))

sh = {x.shape_id: x for x in new.shapes}
set_text(sh[2], '后续设想')
set_text(sh[3], '来自老师们的建议：小古文、主题阅读'); sh[3].width = 1500 * 9144
set_text(sh[6], '有老师建议后续开发这两类课程。总部先列为后续设想，评估校区需求后再定。')
set_text(sh[8], '小古文')
set_text(sh[9], '统编教材从三年级起选入文言文，孩子初读小古文，常卡在读不顺、读不懂。\n'
                '可以围绕教材篇目和同类短文，练朗读、断句和理解，打好文言基础。')
set_text(sh[12], '主题阅读')
set_text(sh[13], '围绕一个主题，把几篇文章或几本书放在一起读，比较、归纳，再表达自己的看法。\n'
                 '与整本书阅读互补：整本书阅读读深一本，主题阅读读宽一类。')
set_text(sh[16], '目前处在收集建议阶段。【待填：是否纳入计划、大致时间】欢迎伙伴们继续提需求。')
new.notes_slide.notes_text_frame.text = ('两类课程都还没有立项，这一页讲清方向和价值即可，不承诺时间。'
                                         '可以现场请伙伴们说说自己校区的需求。')

# 新页放到旧页位置，旧页删除
lst = prs.slides._sldIdLst
ids = list(lst)
by_part = {prs.part.related_part(e.get('{%s}id' % R_NS)): e for e in ids}
old_el, new_el = by_part[old.part], by_part[new.part]
lst.remove(new_el)
old_el.addprevious(new_el)
prs.part.drop_rel(old_el.get('{%s}id' % R_NS))
lst.remove(old_el)

prune_timing(prs)
normalize_paragraphs(prs)
prs.save(os.environ.get('OUT') or F)
print('saved', len(lst), 'slides; new page at', list(lst).index(new_el) + 1)

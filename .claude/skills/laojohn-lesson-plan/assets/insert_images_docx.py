# -*- coding: utf-8 -*-
"""
insert_images_docx —— docx 正文图「回插后处理」（跨技能共享件，见 CLAUDE.md §3）。

不碰共享 docx 引擎 md_to_laojohn_docx.py：先用它把详案 .md 出基础 docx（引擎把
`【图位:编号】` 原样印为文字），再用 python-docx 打开，把每个占位换成真图。
段内只有标记则就地换图；标记与正文（师话等）同段时**只剥标记、正文留在原段**，图另起一段——
整段清空会把老师要念的问题句一起吞掉。缺图则留灰字提示，不抛错。

课型差异全部收敛在顶部 PROFILES 表（同 style_front_page.py 的成例）：
**禁 fork、禁往下面的渲染原语里塞 `if profile == ...` 分支。**
改本文件须同时回归看图写话三课与读书会配图版（回归法见 SKILL/plan）。

用法：
  PYTHONUTF8=1 python insert_images_docx.py <详案.md> [--profile lesson|picture|writing] [...]
输出：<详案名><out_suffix>.docx（不覆盖无图版，便于对照）。
"""
import os
import re
import sys
import tempfile
import argparse
import subprocess

_ENGINE = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                       'md_to_laojohn_docx.py')

# 图注分隔符：`【图位:插-01｜马良掀翻大船】` 的 `｜` 之后是图注。全角，勿用半角 |
# （半角在 Markdown 表格里是列分隔符，占位若出现在表格单元格内会被引擎拆散）。
CAPTION_SEP = '｜'

# ---------------------------------------------------------------------------
# 课型档：差异全是纯数据。新增课型只在此加一档，不改下面任何函数。
# ---------------------------------------------------------------------------
PROFILES = {
    # 读书会详案（原书插图）。页眉 None＝不给引擎传 flag＝回落引擎默认
    # 「老约翰深度阅读／阅读·思辨·表达」。
    'lesson': {
        'header_left':   None,
        'header_right':  None,
        'code_prefixes': '插',
        'skip_prefixes': (),
        # None＝按 <项目根>/读书会原书插图/<书名>/ 解析（见 resolve_images_dir）
        'images_dir':    None,
        'width_cm':      11.0,   # A4 正文可用宽 16cm（边距各 2.5cm）的约 69%
        'max_height_cm': 12.0,   # 原书扫描件常为竖构图，限高免得一张图顶掉整页
        'caption':       True,
        'img_exts':      ('.png', '.jpg', '.jpeg'),
        'book_dir_name': '读书会原书插图',
        'name_from':     'title',   # 目录名＝H1 里的《书名》
    },
    # 看图写话。取图口径由 picture-writing/scripts/imgspec_parser.py 单独持有，
    # 经薄 shim 以 images_dir= 注入——本共享件绝不 import 那个私有模块
    # （反向依赖会让读书会线故障依赖看图写话文件树）。
    'picture': {
        'header_left':   '老约翰·看图写话',
        'header_right':  '从看懂一幅图，到写成一个故事',
        'code_prefixes': '主练备格锚例',
        'skip_prefixes': ('格-',),   # 格式图＝固定资产，手动插
        'images_dir':    None,
        'width_cm':      12.0,
        'max_height_cm': None,
        'caption':       False,
        'img_exts':      ('.png',),
        'book_dir_name': None,
        'name_from':     'title',   # 走 images_dir 注入，本字段不生效
    },
    # 同步习作（写作课）。用图＝统编教材习作页的原图（扫描/翻拍），版权材料不入库，
    # 口径同读书会原书插图：图不入 git、同目录 _图单.md 入库作凭据。
    'writing': {
        'header_left':   '老约翰·同步习作',
        'header_right':  '写清楚·写生动·有章法',
        'code_prefixes': '插',
        'skip_prefixes': (),
        'images_dir':    None,
        'width_cm':      11.0,
        'max_height_cm': 12.0,   # 教材翻拍常为竖构图，限高免得一张顶掉整页
        'caption':       True,
        'img_exts':      ('.jpg', '.jpeg', '.png'),   # 翻拍件多为 jpg，故 jpg 优先
        'book_dir_name': '写作课教材插图',
        # 写作课 H1 是《题目》，而图目录名带年级册（三上-第三单元-续写故事），取标题会解析错，
        # 故按文件名 stem 去尾缀取——与 §7 命名口径 <年级册>-第N单元-<题目>-写作课详案 对齐。
        'name_from':     'stem',
    },
}


def make_place_re(prefixes):
    """按课型的编号前缀集编译占位正则。

    group(1)=编号（如 插-01 / 主-01），group(2)=编号之后的尾串（可含图注）。
    尾串宽松吃到 `】`，以容忍 `【图位:主-01 — 见生图工单】` 这类人读旁注。
    """
    return re.compile(r'【图位[：:]\s*([' + prefixes + r']-\d+)([^】]*)】')


def split_caption(tail):
    """从占位尾串里取图注：只认 CAPTION_SEP 之后的文本，其余旁注一律忽略。"""
    if not tail or CAPTION_SEP not in tail:
        return ''
    return tail.split(CAPTION_SEP, 1)[1].strip()


# ---------------------------------------------------------------------------
# 路径解析
# ---------------------------------------------------------------------------
def resolve_book_name(md_path, name_from='title'):
    """取图目录名。取法由课型档的 name_from 选，两种取法本身课型无关。

    'title'（读书会）：详案首个 `# ` 标题里的《书名》；缺则回退文件名。
      与引擎 md_to_laojohn_docx.convert() 预扫封面书名同一口径（引擎未把这段抽成
      函数，故此处按同一规则实现；若引擎口径变更，此处须同步）。
    'stem'（写作课）：文件名去课型尾缀——写作课 H1 是《题目》、不含年级册，
      而目录名是 <年级册>-第N单元-<题目>，取标题必错。
    """
    stem = os.path.splitext(os.path.basename(md_path))[0]
    fallback = re.sub(r'-(课案详案|写作课详案)$', '', stem)
    if name_from != 'title':
        return fallback
    try:
        with open(md_path, encoding='utf-8') as f:
            for ln in f:
                s = ln.strip()
                if s.startswith('# '):
                    mb = re.search(r'《(.+?)》', s)
                    return mb.group(1).strip() if mb else ''
    except OSError:
        pass
    return fallback


def resolve_images_dir(md_path, prof, cli_dir=''):
    """取图目录：CLI 显式 > profile 固定值 > 按 book_dir_name 推。"""
    if cli_dir:
        return os.path.abspath(cli_dir)
    if prof['images_dir']:
        return prof['images_dir']
    if prof['book_dir_name']:
        # 惰性 import：引擎模块级即 import docx，须在 _ensure_docx() 之后才安全。
        sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
        import md_to_laojohn_docx as engine   # 同目录，复用其 PROJECT_ROOT
        book = resolve_book_name(md_path, prof.get('name_from', 'title'))
        if not book:
            return ''
        return os.path.join(engine.PROJECT_ROOT, prof['book_dir_name'], book)
    return ''


def resolve_image(images_dir, code, exts):
    """某编号的真图路径；按 exts 顺序找，找不到返回 None。"""
    if not images_dir:
        return None
    for ext in exts:
        p = os.path.join(images_dir, code + ext)
        if os.path.isfile(p):
            return p
    return None


# ---------------------------------------------------------------------------
# 渲染原语（课型无关，勿加课型分支）
# ---------------------------------------------------------------------------
def _ensure_docx():
    try:
        import docx  # noqa
        return
    except ImportError:
        pass
    for args in (
        [sys.executable, '-m', 'pip', 'install', '-q', 'python-docx'],
        [sys.executable, '-m', 'pip', 'install', '-q', '--user', 'python-docx'],
    ):
        try:
            subprocess.check_call(args)
            import docx  # noqa
            return
        except Exception:
            continue
    print('[错误] 无法自动安装 python-docx，请手动安装后重试。')
    sys.exit(1)


def build_base_docx(md_path, header_left=None, header_right=None):
    """调共享引擎产基础 docx（落临时文件，绝不碰用户已导出的同名 docx）。"""
    fd, base = tempfile.mkstemp(prefix='详案基础_', suffix='.docx')
    os.close(fd)
    cmd = [sys.executable, _ENGINE, md_path, base]
    if header_left:
        cmd += ['--header-left', header_left]
    if header_right:
        cmd += ['--header-right', header_right]
    subprocess.check_call(cmd)
    return base


def _clear_paragraph(p):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)


def _strip_placeholders(p, place_re):
    """就地剥除段内占位标记，保留同段其余正文（师话等），返回剥后文本。

    曾因整段清空导致与占位同段的师话被静默吞掉（老师拿到的配图版少了问题句）。
    优先逐 run 替换以保住加粗等格式；标记被引擎拆到多个 run 时回退为纯文本重建。
    """
    for r in p.runs:
        if place_re.search(r.text):
            r.text = place_re.sub('', r.text)
    if place_re.search(p.text):  # 跨 run 拆分，逐 run 剥不掉
        residual = place_re.sub('', p.text)
        _clear_paragraph(p)
        if residual.strip():
            p.add_run(residual.strip())
    return p.text.strip()


def _new_paragraph_after(p):
    """在 p 之后插入一个同级空段落，用于承载图片。"""
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph

    new_el = OxmlElement('w:p')
    p._element.addnext(new_el)
    np = Paragraph(new_el, p._parent)
    # 继承样式，避免图片段落掉回 Normal 造成间距突变
    try:
        np.style = p.style
    except Exception:
        pass
    return np


def _add_caption(after_p, text):
    """在 after_p 之后加一行图注：小一号、灰、居中、紧贴图片。

    图注段继承的样式常来自师话（整段 12pt 加粗），故字号/粗细/颜色必须显式设，
    否则会渲染成「加粗大字图注」。
    """
    from docx.shared import Pt, RGBColor

    cp = _new_paragraph_after(after_p)
    run = cp.add_run(text)
    run.font.size = Pt(10.5)
    run.font.bold = False
    run.font.italic = False
    try:
        run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    except Exception:
        pass
    cp.alignment = 1
    cp.paragraph_format.space_before = Pt(2)
    cp.paragraph_format.space_after = Pt(6)
    return cp


def _fit_width(img_path, width_cm, max_height_cm):
    """按宽插图会等比放大高度；超过 max_height_cm 时改按高反算宽。

    用 python-docx 自带的像素读取，不引入 PIL 新依赖。max_height_cm 为 None
    （或读不出尺寸）时原样返回宽度。
    """
    if not max_height_cm:
        return width_cm
    try:
        from docx.image.image import Image
        img = Image.from_file(img_path)
        pw, ph = img.px_width, img.px_height
        if not pw or not ph:
            return width_cm
        if width_cm * ph / pw > max_height_cm:
            return max_height_cm * pw / ph
    except Exception:
        pass
    return width_cm


def insert(md_path, base_docx, *, place_re, images_dir, width_cm,
           max_height_cm=None, skip_prefixes=(), caption=False,
           img_exts=('.png',), out_suffix='-配图'):
    from docx import Document
    from docx.shared import Cm, RGBColor

    doc = Document(base_docx)
    inserted, missing, fmt_skipped = [], [], []
    for p in doc.paragraphs:
        hits = [(m.group(1), m.group(2)) for m in place_re.finditer(p.text)]
        if not hits:
            continue
        # 同段可含多个占位（如一行里连列例-02/03/04）：逐个回插，不得只取首个
        fmt_skipped.extend(c for c, _ in hits
                           if any(c.startswith(s) for s in skip_prefixes))
        hits = [(c, t) for c, t in hits
                if not any(c.startswith(s) for s in skip_prefixes)]
        if not hits:
            continue
        # 剥标记而非清整段：段内若还有正文就把正文留在原段，图另起一段
        residual = _strip_placeholders(p, place_re)
        target = _new_paragraph_after(p) if residual else p
        if not residual:
            _clear_paragraph(target)
        p = target
        last_cap = ''
        for code, tail in hits:
            png = resolve_image(images_dir, code, img_exts)
            if png:
                run = p.add_run()
                run.add_picture(png, width=Cm(_fit_width(png, width_cm,
                                                         max_height_cm)))
                inserted.append(code)
            else:
                want = code + img_exts[0]
                run = p.add_run(f'【图位 {code} 待补：未找到 {want}】')
                try:
                    run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
                except Exception:
                    pass
                missing.append(code)
            if caption:
                cap = split_caption(tail)
                if cap:
                    # 图注是「图片段之后的独立段落」，插不到同段两图之间，故同段多图
                    # 只落最后一条图注。读书会约定占位独占一行（一行一图），不受影响；
                    # 若将来要给同段多图各配图注，须改成每图单独成段。
                    last_cap = cap
        try:
            p.alignment = 1  # 居中
        except Exception:
            pass
        if last_cap:
            _add_caption(p, last_cap)

    out = os.path.splitext(md_path)[0] + out_suffix + '.docx'
    try:
        doc.save(out)
    except PermissionError:
        print(f'[错误] 无法写入 {os.path.basename(out)}——'
              f'请先关闭 Word/WPS 里打开的该文件，然后重跑。')
        sys.exit(2)
    return out, inserted, missing, fmt_skipped


def scan_md(md_path, place_re, skip_prefixes=()):
    """扫详案 .md 里的占位编号（按出现顺序，去重保序）。"""
    with open(md_path, encoding='utf-8') as f:
        text = f.read()
    seen, codes = set(), []
    for m in place_re.finditer(text):
        c = m.group(1)
        if any(c.startswith(s) for s in skip_prefixes) or c in seen:
            continue
        seen.add(c)
        codes.append(c)
    return codes


def check_only(md_path, place_re, images_dir, img_exts, skip_prefixes=()):
    """只报缺口不出 docx：正文缺图 / 目录多图 / 图单缺行。返回缺口总数。"""
    codes = scan_md(md_path, place_re, skip_prefixes)
    no_file = [c for c in codes if not resolve_image(images_dir, c, img_exts)]

    on_disk = []
    if images_dir and os.path.isdir(images_dir):
        for fn in sorted(os.listdir(images_dir)):
            stem, ext = os.path.splitext(fn)
            if ext.lower() in img_exts and place_re.search(f'【图位:{stem}】'):
                on_disk.append(stem)
    orphan = [c for c in on_disk if c not in codes]

    sheet = os.path.join(images_dir, '_图单.md') if images_dir else ''
    if sheet and os.path.isfile(sheet):
        # 只认表格行（行首 `|`）里的编号：图单的说明文字也会写「插-01.png」当示例，
        # 整篇 substring 匹配会把没登记的编号误判成已登记。
        with open(sheet, encoding='utf-8') as f:
            rows = [ln for ln in f if ln.lstrip().startswith('|')]
        sheet_text = ''.join(rows)
        no_row = [c for c in codes if c not in sheet_text]
    else:
        no_row = list(codes) if codes else []

    print(f'取图目录：{images_dir or "(未解析出)"}')
    print(f'正文占位 {len(codes)} 个：{codes}')
    print(f'  缺图（正文有占位、目录无图）{len(no_file)}：{no_file}')
    if sheet and os.path.isfile(sheet):
        print(f'  图单缺行 {len(no_row)}：{no_row}')
    else:
        print(f'  图单缺失：{sheet or "(无目录)"} 不存在，{len(no_row)} 个编号无凭据')
    # 孤图只提示、不判失败：原书插图目录是**素材库**，详案按「少而精、只留能生思考的」
    # 选用一部分是常态，余下的留给老师自取。（看图写话线不同——那边每张生成图都该被
    # 用上，但那条线走自己的 imgspec 自检，不经过这里。）
    print(f'  未选用（目录有图、正文无占位）{len(orphan)} 张：素材库富余，正常')
    if orphan:
        print(f'    {orphan}')
    return len(no_file) + len(no_row)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('md_path')
    ap.add_argument('--profile', choices=sorted(PROFILES), default='lesson')
    ap.add_argument('--images-dir', default='', help='取图目录；不给则按课型档解析')
    ap.add_argument('--base-docx', default='', help='已有基础 docx；不给则现调引擎生成')
    # 凡 PROFILES 里有的字段，CLI 默认值一律 None——给了硬默认会永久盖住课型档，
    # 且完全静默（docx 打开只是图偏大，无任何报错）。
    ap.add_argument('--width-cm', type=float, default=None)
    ap.add_argument('--max-height-cm', type=float, default=None,
                    help='显式 0 表示不限高')
    ap.add_argument('--caption', dest='caption', action='store_true', default=None)
    ap.add_argument('--no-caption', dest='caption', action='store_false')
    ap.add_argument('--out-suffix', default='-配图')
    ap.add_argument('--check-only', action='store_true',
                    help='只报占位/图/图单三方缺口，不出 docx')
    ap.add_argument('--strict-missing', action='store_true',
                    help='有缺图时以退出码 3 结束（默认只提示）')
    args = ap.parse_args()

    prof = PROFILES[args.profile]
    place_re = make_place_re(prof['code_prefixes'])
    width = args.width_cm if args.width_cm is not None else prof['width_cm']
    max_h = (args.max_height_cm if args.max_height_cm is not None
             else prof['max_height_cm'])
    caption = args.caption if args.caption is not None else prof['caption']

    _ensure_docx()
    images_dir = resolve_images_dir(args.md_path, prof, args.images_dir)

    if args.check_only:
        gaps = check_only(args.md_path, place_re, images_dir,
                          prof['img_exts'], prof['skip_prefixes'])
        sys.exit(1 if gaps else 0)

    temp_base = not args.base_docx
    base = args.base_docx or build_base_docx(
        args.md_path, prof['header_left'], prof['header_right'])
    try:
        out, inserted, missing, fmt = insert(
            args.md_path, base, place_re=place_re, images_dir=images_dir,
            width_cm=width, max_height_cm=max_h,
            skip_prefixes=prof['skip_prefixes'], caption=caption,
            img_exts=prof['img_exts'], out_suffix=args.out_suffix)
    finally:
        if temp_base and os.path.exists(base):
            try:
                os.remove(base)
            except OSError:
                pass
    print(f'回插完成：{out}')
    print(f'  插入 {len(inserted)} 张：{inserted}')
    if missing:
        print(f'  缺图待补 {len(missing)}：{missing}（已在文中留灰字提示）')
    if fmt:
        print(f'  跳过格式图 {len(set(fmt))}：{sorted(set(fmt))}（固定资产，手动插）')
    if missing and args.strict_missing:
        sys.exit(3)


if __name__ == '__main__':
    main()

"""direct_build.py - 写作课仓内直出一条龙（2026-09-26 立）：截稿纸 → 课件构建脚本 → 转 PPT → 审查闸门 → 并入分组 → 注入动画 → 回读核验。

前置：详案已定稿；课件配图工具已出齐本课图并放进 `…/gpt-image/ppt配图缓存/直出/`（含 吉祥物.png）；
写作配套已渲出 `写作配套输出/<课次>/<课次>-学生用.html`；`写作课件中间稿输出/<课次>/<课次>-课件构建.py` 已写好。

    python direct_build.py <课次> [--out 路径.pptx] [--force] [--accept-edits]

**改动侦测**：出件时存 `<pptx 名>-文字快照.json`（逐页逐 shape_id 的文字，随 pptx 入库）。再次出件前先拿仓内
pptx 比快照——不一致说明有人（同事）直接在 PPT 里改过字，重出会把改动覆盖掉，于是停下列出改动清单；
把改动写回课件构建脚本后加 --accept-edits 重出。

缺省输出 `写作课件PPT输出/<课次>/<课次>-课件PPT.pptx`；该文件已存在（例如是外部件动画版）时拒绝覆盖，
须 --force 或用 --out 另存。之后照常做 COM 逐页目检与页标回注（见 laojohn-ppt/SKILL.md「仓内直出」节）。
"""
import argparse
import json
import os
import pathlib
import subprocess
import sys
from collections import Counter

HERE = pathlib.Path(__file__).resolve().parent
SKILL = HERE.parent
ROOT = HERE.parents[3]
ENV = {**os.environ, "PYTHONUTF8": "1"}


def run(*args):
    print("▶", " ".join(str(a) for a in args[1:3]), "…")
    r = subprocess.run([sys.executable, *map(str, args)], cwd=ROOT, env=ENV)
    if r.returncode:
        raise SystemExit(f"失败：{args[0]}（退出码 {r.returncode}）")


def text_map(pptx):
    """{页号: {shape_id: 文字}}——只记有字的文本框。"""
    from pptx import Presentation
    return {str(i): {str(sh.shape_id): sh.text_frame.text for sh in s.shapes
                     if sh.has_text_frame and sh.text_frame.text.strip()}
            for i, s in enumerate(Presentation(pptx).slides, 1)}


def edits_since_snapshot(pptx, snap):
    """仓内 pptx 与出件快照的文字差异：[(页, 原文, 现文)]。"""
    old = json.loads(pathlib.Path(snap).read_text(encoding="utf-8"))
    now = text_map(pptx)
    out = []
    for pg in sorted(set(old) | set(now), key=int):
        a, b = old.get(pg, {}), now.get(pg, {})
        for sid in sorted(set(a) | set(b), key=int):
            if a.get(sid) != b.get(sid):
                out.append((pg, a.get(sid, "（无）"), b.get(sid, "（删掉了）")))
    return out


def check(pptx, sheet):
    from pptx import Presentation
    prs = Presentation(pptx)
    doc = json.loads(pathlib.Path(sheet).read_text(encoding="utf-8"))
    bad = []
    for i, s in enumerate(prs.slides, 1):
        ids = [e.get("id") for e in s._element.iter() if e.tag.endswith("}cNvPr")]
        if [k for k, v in Counter(ids).items() if v > 1]:
            bad.append(f"P{i} shape id 重复")
        rec = doc["slides"][str(i)]
        if rec.get("skip"):
            continue
        tops = {sh.shape_id: sh.top / 914400 for sh in s.shapes}
        prev = -1
        for n, g in enumerate(rec["groups"], 1):
            t = min(tops[x] for x in g)
            if t < prev - 0.45:
                bad.append(f"P{i} 第{n}击由下往上跳")
            prev = max(prev, t)
    return len(prs.slides), bad


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("unit", help="课次标识，如 六上-第五单元-围绕中心意思写")
    ap.add_argument("--out")
    ap.add_argument("--force", action="store_true")
    ap.add_argument("--accept-edits", action="store_true",
                    help="确认 PPT 里的直接改动已写回课件构建脚本，允许重出覆盖")
    a = ap.parse_args()
    u = a.unit
    mid = ROOT / "写作课件中间稿输出" / u
    build = mid / f"{u}-课件构建.py"
    html = mid / f"{u}-课件.dc.html"
    student = ROOT / "写作配套输出" / u / f"{u}-学生用.html"
    img_dir = ROOT / "课件配图工具/课件产出" / u / "gpt-image/ppt配图缓存/直出"
    out = pathlib.Path(a.out) if a.out else ROOT / "写作课件PPT输出" / u / f"{u}-课件PPT.pptx"
    for p, what in ((build, "课件构建脚本"), (student, "学生用配套 html"), (img_dir / "吉祥物.png", "吉祥物透明底图")):
        if not p.exists():
            raise SystemExit(f"缺{what}：{p}")
    snap = out.with_name(out.stem + "-文字快照.json")
    if out.exists() and snap.exists():
        diff = edits_since_snapshot(out, snap)
        if diff and not a.accept_edits:
            print(f"⚠ {out.name} 在上次出件后被直接改过 {len(diff)} 处文字（多半是同事在 PPT 里改的）：")
            for pg, x, y in diff:
                print(f"   P{pg}：{x[:60]!r}")
                print(f"        → {y[:60]!r}")
            raise SystemExit("重出会覆盖这些改动。先把它们写回课件构建脚本，再加 --accept-edits 重出。")
    elif out.exists() and not a.force:
        raise SystemExit(f"{out} 已存在且没有直出快照（可能是外部件动画版）；确认要覆盖加 --force，或用 --out 另存")
    if out.parent.name != u:
        raise SystemExit("输出须放在 写作课件PPT输出/<课次>/ 下（审查闸门按目录名找详案）")
    out.parent.mkdir(parents=True, exist_ok=True)
    sheet = out.with_name(out.stem + "-anim.json")
    if sheet.exists():
        sheet.unlink()

    run(HERE / "grab_worksheet.py", student, img_dir / "稿纸.png")
    run(build)
    run(HERE / "html_to_pptx.py", html, out)
    run(SKILL / "scripts/inspect_pptx.py", out, "--quiet")
    run(SKILL / "scripts/audit_against_plan.py", out)
    run(HERE / "html_to_pptx.py", out, "--merge-anim")
    run(SKILL / "scripts/animate_pptx.py", out, sheet, "--in-place", "--effect", "appear")
    n, bad = check(out, sheet)
    snap.write_text(json.dumps(text_map(out), ensure_ascii=False, indent=1), encoding="utf-8")
    side = out.with_name(out.stem + "-动画分组.json")
    if side.exists():
        side.unlink()
    mb = out.stat().st_size / 1048576
    print(f"\n完成：{out}（{n} 页，{mb:.1f}MB）")
    print("核验：" + ("shape id 无重复、无回跳" if not bad else "；".join(bad)))
    if mb > 5:
        print("⚠ 超过 5MB（pre-commit 软拦）：入库前先跑 scripts/shrink_pptx_media.py")
    print("下一步：COM 逐页导出目检 → 页标回注 → 重渲详案 docx → 入库并推送（见 SKILL「仓内直出」节）")
    rel = lambda p: p.relative_to(ROOT).as_posix()
    print("入库清单（直出件本身就是终稿，图不入库、别处重渲不出来，所以 pptx 必须入库）：")
    for p in (out, sheet, snap, build, ROOT / "写作课详案输出" / f"{u}-写作课详案.md"):
        print("   git add", rel(p))
    print("   git add 写作课件中间稿输出/" + u + "/*页标映射.json   （页标回注后）")


if __name__ == "__main__":
    main()

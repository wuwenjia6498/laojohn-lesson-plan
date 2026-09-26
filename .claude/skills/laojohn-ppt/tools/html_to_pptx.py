"""html_to_pptx.py - 把按母版 framework.md 写的 `*.dc.html` 转成可编辑 PPTX（写作课线仓内直出，2026-09-26 立）。

为什么有它：写作课 PPT 原先由 claude.ai Design → Slides 按母版写 dc.html、再用平台「导出可编辑 PPTX」出初稿；
母版规范（assets/writing-master/framework.md）本仓照样能执行，缺的只是平台那个导出器——本脚本补这一块。

做法：Playwright 渲染每个 <section>（剥掉 deck-stage/support.js 运行时，直接放进 1920×1080 裸容器），
逐元素读几何与计算样式，映射成 python-pptx 原生对象：
  带底色/描边的块 → rect / roundRect / ellipse（四角统一圆角；局部圆角落成直角，母版已认可）
  linear-gradient → gradFill；rgba / opacity → alpha
  文字 → 无填充文本框，run 级颜色/加粗/斜体/字号（pt = px×0.75），微软雅黑，latin 在 ea 前
  data-om-raster="true"（只含图形） → 截图贴回 PNG（与平台口径一致；含文字的会报警）
  <img> → 图片；<image-slot id> → 有映射就放图（cover 裁切），否则放灰色占位图，名字写 slot:<id>
画布 1920px = 20in（96px/in），与平台导出件同尺寸，下游 inspect/animate/illustrate 直接吃。

    python html_to_pptx.py 课件.dc.html out.pptx [--slots slots.json] [--only 3 5]
slots.json：{"slot-id": "图片路径", ...}（相对 json 所在目录）

点击动画：元素上写 data-anim="N"（按详案师话顺序编号，子元素随最近的祖先标记），转换时写出
<out>-动画分组.json；跑完 inspect_pptx.py 与 audit_against_plan.py（审查闸门）后，
    python html_to_pptx.py <out.pptx> --merge-anim
并入工作单，再照常 animate_pptx.py 注入。没有标记的页一律 skip。
"""
import argparse
import io
import json
import os
import pathlib
import re
import sys
import tempfile

os.environ.setdefault(
    "PLAYWRIGHT_BROWSERS_PATH",
    os.path.join(os.path.expanduser("~"), "AppData", "Local", "ms-playwright"),
)

from lxml import etree  # noqa: E402
from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Emu  # noqa: E402

PX = 9525  # EMU / px（96 dpi）
FONT = "Microsoft YaHei"
HERE = pathlib.Path(__file__).resolve().parent

# ---------------------------------------------------------------- 浏览器侧：逐元素取几何与样式
JS_EXTRACT = r"""
(idx) => {
  const root = document.querySelectorAll('.__slide')[idx];
  const R = root.getBoundingClientRect();
  const items = [];
  let rasterSeq = 0;
  const rel = r => ({x: r.left - R.left, y: r.top - R.top, w: r.width, h: r.height});
  const transparent = c => !c || c === 'transparent' || /rgba\([^)]*,\s*0\)$/.test(c);
  function opacityOf(el) {
    let o = 1;
    for (let e = el; e && e !== root; e = e.parentElement) o *= parseFloat(getComputedStyle(e).opacity);
    return o;
  }
  function clipOf(el) {
    let c = null;
    for (let e = el.parentElement; e && e !== root; e = e.parentElement) {
      const cs = getComputedStyle(e);
      if (cs.overflow !== 'visible' || cs.overflowX !== 'visible') {
        const r = e.getBoundingClientRect();
        c = c ? {l: Math.max(c.l, r.left), t: Math.max(c.t, r.top), r: Math.min(c.r, r.right), b: Math.min(c.b, r.bottom)}
              : {l: r.left, t: r.top, r: r.right, b: r.bottom};
      }
    }
    return c;
  }
  function radii(cs) {
    return ['borderTopLeftRadius','borderTopRightRadius','borderBottomRightRadius','borderBottomLeftRadius']
      .map(k => cs[k]);
  }
  function borders(cs) {
    const o = {};
    for (const s of ['Top','Right','Bottom','Left']) {
      const w = parseFloat(cs['border'+s+'Width']);
      if (w > 0 && cs['border'+s+'Style'] !== 'none' && !transparent(cs['border'+s+'Color']))
        o[s.toLowerCase()] = {w, c: cs['border'+s+'Color']};
    }
    return o;
  }
  function styleOf(el) {
    const cs = getComputedStyle(el);
    return {c: cs.color, fs: parseFloat(cs.fontSize), fw: parseInt(cs.fontWeight) || 400,
            it: cs.fontStyle === 'italic', u: (cs.textDecorationLine || '').includes('underline')};
  }
  // 收集 el 的行内内容：文字 run、<br>、行内底色块；遇到非 inline 子元素则切段
  function collect(el, segs, bgs) {
    for (const n of el.childNodes) {
      if (n.nodeType === 3) {
        const t = n.textContent;
        if (!t) continue;
        const pe = n.parentElement;
        const ws = getComputedStyle(pe).whiteSpace;
        const txt = /pre/.test(ws) ? t : t.replace(/[ \t\r\n\f]+/g, ' ');
        segs[segs.length - 1].runs.push({t: txt, s: styleOf(pe), node: n});
      } else if (n.nodeType === 1) {
        const cs = getComputedStyle(n);
        if (cs.display === 'none') continue;
        if (n.tagName === 'BR') { segs[segs.length - 1].runs.push({br: true}); continue; }
        if (cs.display === 'inline') {
          if (!transparent(cs.backgroundColor))
            for (const r of n.getClientRects()) bgs.push({...rel(r), bg: cs.backgroundColor, rad: radii(cs)});
          collect(n, segs, bgs);
        } else {
          segs.push({runs: []});
        }
      }
    }
  }
  function textBlocks(el, cs) {
    const segs = [{runs: []}], bgs = [];
    collect(el, segs, bgs);
    const out = [];
    for (const seg of segs) {
      // 去首尾空白
      const rs = seg.runs;
      while (rs.length && !rs[0].br && !rs[0].t.trim()) rs.shift();
      while (rs.length && !rs[rs.length-1].br && !rs[rs.length-1].t.trim()) rs.pop();
      if (!rs.some(r => r.t && r.t.trim())) continue;
      if (rs[0].t) rs[0].t = rs[0].t.replace(/^\s+/, '');
      const last = rs[rs.length-1]; if (last.t) last.t = last.t.replace(/\s+$/, '');
      // 用 Range 量出这段文字实际占的框与行数
      const rg = document.createRange();
      const tn = rs.filter(r => r.node);
      rg.setStartBefore(tn[0].node); rg.setEndAfter(tn[tn.length-1].node);
      const rects = [...rg.getClientRects()].filter(r => r.width > 0.5);
      const bb = rg.getBoundingClientRect();
      const tops = new Set(rects.map(r => Math.round(r.top / 4)));
      const pl = parseFloat(cs.paddingLeft), pr = parseFloat(cs.paddingRight);
      const er = el.getBoundingClientRect();
      const bl = parseFloat(cs.borderLeftWidth) || 0, brw = parseFloat(cs.borderRightWidth) || 0;
      out.push({
        k: 'text', box: rel(bb),
        content: {x: er.left - R.left + pl + bl, w: er.width - pl - pr - bl - brw},
        lines: tops.size,
        align: cs.textAlign, indent: parseFloat(cs.textIndent) || 0,
        lh: cs.lineHeight === 'normal' ? null : parseFloat(cs.lineHeight),
        fs: parseFloat(cs.fontSize),
        runs: rs.map(r => r.br ? {br: true} : {t: r.t, s: r.s}),
        opacity: opacityOf(el),
      });
    }
    return {out, bgs};
  }
  function animOf(el) {
    const a = el.closest('[data-anim]');
    return a && root.contains(a) ? parseFloat(a.getAttribute('data-anim')) : null;
  }
  function visit(el) {
    const n0 = items.length;
    visitInner(el);
    for (let k = n0; k < items.length; k++) if (items[k].anim === undefined) items[k].anim = animOf(el);
  }
  function visitInner(el) {
    const cs = getComputedStyle(el);
    if (cs.display === 'none' || cs.visibility === 'hidden') return;
    if (['SCRIPT','STYLE','HELMET','META'].includes(el.tagName)) return;
    const r = el.getBoundingClientRect();
    if (el.tagName === 'IMAGE-SLOT') {
      const host = el.parentElement.hasAttribute('data-om-raster') ? el.parentElement : el;
      const hr = host.getBoundingClientRect();
      items.push({k: 'slot', id: el.id, ph: el.getAttribute('placeholder') || '',
                  shape: el.getAttribute('shape') || 'rect', radius: parseFloat(el.getAttribute('radius')) || 0,
                  ...rel(hr)});
      return;
    }
    if (el !== root && el.hasAttribute('data-om-raster')) {
      const slot = el.querySelector('image-slot');
      if (slot) { visit(slot); return; }
      const id = 'r' + (rasterSeq++);
      el.setAttribute('data-raster-id', id);
      items.push({k: 'raster', id, hasText: !!el.innerText.trim(), ...rel(r)});
      return;
    }
    if (el.tagName === 'IMG') {
      items.push({k: 'img', src: el.currentSrc || el.src, opacity: opacityOf(el), ...rel(r)});
      return;
    }
    if (el !== root && cs.display !== 'inline') {
      const bg = transparent(cs.backgroundColor) ? null : cs.backgroundColor;
      const grad = /gradient/.test(cs.backgroundImage) ? cs.backgroundImage : null;
      const bd = borders(cs);
      if ((bg || grad || Object.keys(bd).length) && r.width > 0 && r.height > 0) {
        const c = clipOf(el);
        const out = c && (r.left < c.l - 0.5 || r.top < c.t - 0.5 || r.right > c.r + 0.5 || r.bottom > c.b + 0.5);
        const rad = radii(cs), round = rad.some(v => parseFloat(v) > 0);
        if (out && round && !el.innerText.trim() && !el.querySelector('img,image-slot')) {
          // 超出 overflow:hidden 祖先的圆形/圆角图形：按裁切后的样子截图贴回
          const id = 'r' + (rasterSeq++);
          el.setAttribute('data-raster-id', id);
          const l = Math.max(r.left, c.l), t = Math.max(r.top, c.t), rr = Math.min(r.right, c.r), b = Math.min(r.bottom, c.b);
          if (rr > l && b > t)
            items.push({k: 'raster', id, hasText: false, clip: {x: l, y: t, width: rr - l, height: b - t},
                        x: l - R.left, y: t - R.top, w: rr - l, h: b - t});
        } else {
          let g = rel(r);
          if (out) {
            const l = Math.max(r.left, c.l), t = Math.max(r.top, c.t), rr = Math.min(r.right, c.r), b = Math.min(r.bottom, c.b);
            g = {x: l - R.left, y: t - R.top, w: Math.max(0, rr - l), h: Math.max(0, b - t)};
          }
          if (g.w > 0 && g.h > 0) items.push({k: 'box', bg, grad, bd, rad, opacity: opacityOf(el), ...g});
        }
      }
      const {out, bgs} = textBlocks(el, cs);
      for (const b of bgs) items.push({k: 'box', bg: b.bg, grad: null, bd: {}, rad: b.rad, opacity: opacityOf(el),
                                       x: b.x, y: b.y, w: b.w, h: b.h});
      items.push(...out);
    } else if (el === root) {
      const {out} = textBlocks(el, cs); items.push(...out);
    }
    for (const c of el.children) {
      if (getComputedStyle(c).display !== 'inline' || c.tagName === 'IMG' || c.tagName === 'IMAGE-SLOT') visit(c);
    }
  }
  visit(root);
  return items;
}
"""


def build_harness(src: pathlib.Path) -> str:
    """剥掉 dc 运行时：取 <helmet><style> 与每个 <section> 的内容，放进独立 1920×1080 容器。"""
    html = src.read_text(encoding="utf-8")
    styles = "\n".join(re.findall(r"<style>(.*?)</style>", html, flags=re.S))
    secs = re.findall(r"<section\b([^>]*)>(.*?)</section>", html, flags=re.S)
    if not secs:
        sys.exit("找不到 <section>：不是母版格式的 dc.html？")
    body = []
    for attrs, inner in secs:
        label = re.search(r'data-label="([^"]*)"', attrs)
        body.append(f'<div class="__slide" data-label="{label.group(1) if label else ""}">{inner}</div>')
    base = src.parent.resolve().as_uri() + "/"
    return f"""<!DOCTYPE html><html><head><meta charset="utf-8"><base href="{base}">
<style>
html,body{{margin:0;padding:0;background:transparent;}}
{styles}
.__slide{{position:relative;width:1920px;height:1080px;overflow:hidden;}}
image-slot{{display:block;width:100%;height:100%;}}
</style></head><body>{''.join(body)}</body></html>"""


# ---------------------------------------------------------------- 颜色/几何工具
def parse_color(c):
    """'rgb(1, 2, 3)' / 'rgba(1,2,3,0.5)' → ('010203', alpha)"""
    if not c:
        return None, 1.0
    m = re.match(r"rgba?\(([^)]*)\)", c)
    if not m:
        return None, 1.0
    parts = [p.strip() for p in m.group(1).replace("/", ",").split(",")]
    r, g, b = (int(float(x)) for x in parts[:3])
    a = float(parts[3]) if len(parts) > 3 else 1.0
    return f"{r:02X}{g:02X}{b:02X}", a


def parse_gradient(g):
    """computed linear-gradient → (ooxml角度, [(pos0-1, hex, alpha)])"""
    m = re.match(r"linear-gradient\((.*)\)\s*$", g.strip())
    if not m:
        return None
    body = m.group(1)
    parts = re.split(r",\s*(?![^()]*\))", body)
    ang = 180.0
    first = parts[0].strip()
    if first.endswith("deg"):
        ang = float(first[:-3]); parts = parts[1:]
    elif first.startswith("to "):
        ang = {"to top": 0, "to right": 90, "to bottom": 180, "to left": 270}.get(first, 180); parts = parts[1:]
    stops = []
    for p in parts:
        cm = re.match(r"(rgba?\([^)]*\))\s*([\d.]+%)?", p.strip())
        if not cm:
            continue
        hx, a = parse_color(cm.group(1))
        pos = float(cm.group(2)[:-1]) / 100 if cm.group(2) else None
        stops.append([pos, hx, a])
    n = len(stops)
    for i, s in enumerate(stops):
        if s[0] is None:
            s[0] = i / (n - 1) if n > 1 else 0
    return (ang - 90) % 360, stops


def px(v):
    return Emu(int(round(v * PX)))


def _srgb(hx, alpha):
    el = etree.SubElement(etree.Element("dummy"), qn("a:srgbClr"), val=hx)
    if alpha < 0.999:
        etree.SubElement(el, qn("a:alpha"), val=str(int(round(alpha * 100000))))
    return el


def set_fill(shape, hx, alpha):
    spPr = shape._element.spPr
    for tag in ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill"):
        for e in spPr.findall(qn(tag)):
            spPr.remove(e)
    sf = etree.Element(qn("a:solidFill"))
    sf.append(_srgb(hx, alpha))
    _insert_fill(spPr, sf)


def set_grad(shape, ang, stops, alpha_mul):
    spPr = shape._element.spPr
    gf = etree.Element(qn("a:gradFill"), rotWithShape="1")
    gl = etree.SubElement(gf, qn("a:gsLst"))
    for pos, hx, a in stops:
        gs = etree.SubElement(gl, qn("a:gs"), pos=str(int(round(pos * 100000))))
        gs.append(_srgb(hx, a * alpha_mul))
    etree.SubElement(gf, qn("a:lin"), ang=str(int(round(ang * 60000))), scaled="0")
    _insert_fill(spPr, gf)


def _insert_fill(spPr, fill):
    # spPr 子序：xfrm, custGeom/prstGeom, 填充, ln, ...
    geom = spPr.find(qn("a:prstGeom"))
    if geom is None:
        geom = spPr.find(qn("a:custGeom"))
    geom.addnext(fill)


def no_line(shape):
    shape.line.fill.background()


def set_line(shape, hx, alpha, w_px):
    ln = shape._element.spPr.get_or_add_ln()
    ln.set("w", str(int(round(w_px * PX))))
    for e in list(ln):
        ln.remove(e)
    sf = etree.SubElement(ln, qn("a:solidFill"))
    sf.append(_srgb(hx, alpha))


# ---------------------------------------------------------------- 各类对象写入
class Writer:
    def __init__(self, prs, slide, raster_imgs, slot_map, placeholder_cache):
        self.prs, self.slide = prs, slide
        self.raster_imgs, self.slot_map = raster_imgs, slot_map
        self.ph_cache = placeholder_cache
        self.warn = []

    def box(self, it):
        x, y, w, h = it["x"], it["y"], it["w"], it["h"]
        rad = [_rad_px(r, w, h) for r in it["rad"]]
        uniform = max(rad) - min(rad) < 0.5
        r0 = rad[0] if uniform else 0.0
        if uniform and r0 >= min(w, h) / 2 - 0.5 and abs(w - h) < 1:
            kind = MSO_SHAPE.OVAL
        elif uniform and r0 > 0.5:
            kind = MSO_SHAPE.ROUNDED_RECTANGLE
        else:
            kind = MSO_SHAPE.RECTANGLE
            if not uniform and max(rad) > 0.5:
                # 局部圆角：若两角为圆、正好是胶囊半边，也按直角；母版认可 PPT 里直角
                pass
        bd = it["bd"]
        sides = list(bd.values())
        full_border = len(bd) == 4 and len({(s["w"], s["c"]) for s in sides}) == 1
        has_body = it["bg"] or it["grad"] or full_border
        if has_body:
            shp = self.slide.shapes.add_shape(kind, px(x), px(y), px(w), px(h))
            if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
                shp.adjustments[0] = min(0.5, r0 / min(w, h))
            if it["grad"]:
                g = parse_gradient(it["grad"])
                if g:
                    set_grad(shp, g[0], g[1], it["opacity"])
            elif it["bg"]:
                hx, a = parse_color(it["bg"])
                set_fill(shp, hx, a * it["opacity"])
            else:
                shp.fill.background()
            if full_border:
                hx, a = parse_color(sides[0]["c"])
                set_line(shp, hx, a * it["opacity"], sides[0]["w"])
            else:
                no_line(shp)
            _clear_text(shp)
        if not full_border:
            # 单边描边（三档左边框、表格竖线等）→ 细条矩形
            for side, s in bd.items():
                bw = s["w"]
                geo = {"left": (x, y, bw, h), "right": (x + w - bw, y, bw, h),
                       "top": (x, y, w, bw), "bottom": (x, y + h - bw, w, bw)}[side]
                bar = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, *(px(v) for v in geo))
                hx, a = parse_color(s["c"])
                set_fill(bar, hx, a * it["opacity"])
                no_line(bar)
                _clear_text(bar)

    def text(self, it):
        box, fs = it["box"], it["fs"]
        single = it["lines"] <= 1 and not any(r.get("br") for r in it["runs"])
        lh = it["lh"] or fs * 1.32
        slack = fs * 0.35
        indent = it["indent"]
        if single:
            x, w = box["x"], box["w"] + slack
            indent = 0  # 单行按实际字位定位，首行缩进已含在 x 里
            if it["align"] in ("center", "right", "end"):
                # 单行居中/右对齐：以内容框定位，保留对齐方式
                x, w = it["content"]["x"], it["content"]["w"]
        else:
            x, w = it["content"]["x"], it["content"]["w"] + slack
        # PowerPoint 与浏览器分配行距的方式不同：行高比 1.4 时两边首行重合，偏离越多差越大（实测标定）
        ratio = lh / fs
        dy = fs * (0.56 if ratio < 1.4 else 0.73) * (ratio - 1.4)
        y, h = box["y"] - dy, max(box["h"], lh)
        tb = self.slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
        tf = tb.text_frame
        tf.word_wrap = not single
        bp = tf._txBody.find(qn("a:bodyPr"))
        for k in ("lIns", "tIns", "rIns", "bIns"):
            bp.set(k, "0")
        for e in list(bp):
            bp.remove(e)
        etree.SubElement(bp, qn("a:noAutofit"))
        # 段落：<br> 切段（母版的 <br> 就是换行意图）
        paras = [[]]
        for r in it["runs"]:
            if r.get("br"):
                paras.append([])
            else:
                paras[-1].append(r)
        txBody = tf._txBody
        for p in txBody.findall(qn("a:p")):
            txBody.remove(p)
        algn = {"center": "ctr", "right": "r", "end": "r", "justify": "just"}.get(it["align"])
        for runs in paras:
            p = etree.SubElement(txBody, qn("a:p"))
            pPr = etree.SubElement(p, qn("a:pPr"))
            if algn and not (single and algn == "just"):
                pPr.set("algn", algn)
            if indent:
                pPr.set("indent", str(int(round(indent * PX))))
            ls = etree.SubElement(pPr, qn("a:lnSpc"))
            etree.SubElement(ls, qn("a:spcPts"), val=str(int(round(lh * 0.75 * 100))))
            last_style = None
            for r in runs:
                if not r["t"]:
                    continue
                last_style = r["s"]
                _add_run(p, r["t"], r["s"], it["opacity"])
            end = etree.SubElement(p, qn("a:endParaRPr"), lang="zh-CN",
                                   sz=str(int(round((last_style or {"fs": fs})["fs"] * 0.75 * 100))))
            _font_children(end)

    def img(self, it, src_path):
        if not src_path or not src_path.exists():
            self.warn.append(f"图片找不到：{it['src']}")
            return
        data = _shrink(src_path, it["w"])
        pic = self.slide.shapes.add_picture(io.BytesIO(data), px(it["x"]), px(it["y"]), px(it["w"]), px(it["h"]))
        return pic

    def raster(self, it):
        png = self.raster_imgs.get(it["id"])
        if png is None:
            return
        if it.get("hasText"):
            self.warn.append("raster 块里含文字（母版禁止：会导致文字重影），已照截图贴入")
        pic = self.slide.shapes.add_picture(io.BytesIO(png), px(it["x"]), px(it["y"]), px(it["w"]), px(it["h"]))
        pic.name = f"raster {it['id']}"

    def slot(self, it, base_dir):
        path = self.slot_map.get(it["id"])
        x, y, w, h = it["x"], it["y"], it["w"], it["h"]
        if path:
            p = (base_dir / path) if not pathlib.Path(path).is_absolute() else pathlib.Path(path)
            im = Image.open(p)
            data = _cover(im, w, h)
            pic = self.slide.shapes.add_picture(io.BytesIO(data), px(x), px(y), px(w), px(h))
        else:
            key = (int(w), int(h))
            if key not in self.ph_cache:
                buf = io.BytesIO()
                Image.new("RGB", (max(8, int(w / 4)), max(8, int(h / 4))), (0xEE, 0xF1, 0xF5)).save(buf, "PNG")
                self.ph_cache[key] = buf.getvalue()
            pic = self.slide.shapes.add_picture(io.BytesIO(self.ph_cache[key]), px(x), px(y), px(w), px(h))
        pic.name = f"slot:{it['id']}"
        cNvPr = pic._element.nvPicPr.cNvPr
        cNvPr.set("descr", it["ph"])
        if it["shape"] in ("rounded", "circle"):
            geom = pic._element.spPr.find(qn("a:prstGeom"))
            geom.set("prst", "ellipse" if it["shape"] == "circle" else "roundRect")
            if it["shape"] == "rounded":
                av = geom.find(qn("a:avLst"))
                if av is None:
                    av = etree.SubElement(geom, qn("a:avLst"))
                etree.SubElement(av, qn("a:gd"), name="adj",
                                 fmla=f"val {int(min(50000, it['radius'] / min(w, h) * 100000))}")


def _rad_px(v, w, h):
    v = v.split()[0]
    if v.endswith("%"):
        return float(v[:-1]) / 100 * min(w, h)
    return float(v.replace("px", "") or 0)


def _clear_text(shp):
    # add_shape 自带 <p:style>（主题阴影/描边/文字色）→ 去掉；文本框设零边距
    st = shp._element.find(qn("p:style"))
    if st is not None:
        shp._element.remove(st)
    tf = shp.text_frame
    bp = tf._txBody.find(qn("a:bodyPr"))
    for k in ("lIns", "tIns", "rIns", "bIns"):
        bp.set(k, "0")


def _font_children(rPr):
    etree.SubElement(rPr, qn("a:latin"), typeface=FONT)
    etree.SubElement(rPr, qn("a:ea"), typeface=FONT)


def _add_run(p, text, s, opacity):
    r = etree.SubElement(p, qn("a:r"))
    rPr = etree.SubElement(r, qn("a:rPr"), lang="zh-CN", sz=str(int(round(s["fs"] * 0.75 * 100))),
                           b="1" if s["fw"] >= 600 else "0", i="1" if s["it"] else "0", dirty="0")
    if s["u"]:
        rPr.set("u", "sng")
    hx, a = parse_color(s["c"])
    if hx:
        sf = etree.SubElement(rPr, qn("a:solidFill"))
        sf.append(_srgb(hx, a * opacity))
    _font_children(rPr)
    t = etree.SubElement(r, qn("a:t"))
    t.text = text


_SHRINK_CACHE = {}


def _shrink(path, shown_w):
    """按显示宽度 1.5 倍缩图再编码（不透明→JPEG q85，带透明→256 色 PNG），同图同尺寸复用。
    不缩的话 30 页课件会到 18MB（源图 1200~2000px 原样塞入）。"""
    key = (str(path), int(shown_w))
    if key in _SHRINK_CACHE:
        return _SHRINK_CACHE[key]
    im = Image.open(path)
    tw = int(min(im.width, max(64, shown_w * 1.5)))
    if tw < im.width:
        im = im.resize((tw, round(im.height * tw / im.width)), Image.LANCZOS)
    buf = io.BytesIO()
    has_alpha = im.mode in ("RGBA", "LA") or (im.mode == "P" and "transparency" in im.info)
    if has_alpha:
        im = im.convert("RGBA")
        if im.getchannel("A").getextrema()[0] < 255:
            im.quantize(256, method=Image.Quantize.FASTOCTREE).save(buf, "PNG", optimize=True)
        else:
            im.convert("RGB").save(buf, "JPEG", quality=85)
    elif path.suffix.lower() == ".png" and im.width * im.height < 400 * 400:
        im.save(buf, "PNG", optimize=True)  # 小图（logo 之类）保留 PNG
    else:
        im.convert("RGB").save(buf, "JPEG", quality=85)
    _SHRINK_CACHE[key] = buf.getvalue()
    return _SHRINK_CACHE[key]


def _cover(im, w, h):
    im = im.convert("RGBA") if im.mode in ("RGBA", "LA", "P") else im.convert("RGB")
    tw, th = w / h, im.width / im.height
    if th > tw:
        nw = int(im.height * tw); off = (im.width - nw) // 2
        im = im.crop((off, 0, off + nw, im.height))
    else:
        nh = int(im.width / tw); off = (im.height - nh) // 2
        im = im.crop((0, off, im.width, off + nh))
    if im.width > 1400:
        im = im.resize((1400, int(1400 * im.height / im.width)), Image.LANCZOS)
    buf = io.BytesIO()
    if im.mode == "RGBA":
        im.save(buf, "PNG", optimize=True)
    else:
        im.save(buf, "JPEG", quality=88)
    return buf.getvalue()


# ---------------------------------------------------------------- 主流程
def convert(src, out, slot_json=None, only=None):
    from playwright.sync_api import sync_playwright

    src = pathlib.Path(src).resolve()
    slot_map, slot_base = {}, src.parent
    if slot_json:
        sj = pathlib.Path(slot_json).resolve()
        slot_map = json.loads(sj.read_text(encoding="utf-8"))
        slot_base = sj.parent
    harness = build_harness(src)
    tmp = pathlib.Path(tempfile.mkdtemp()) / "harness.html"
    tmp.write_text(harness, encoding="utf-8")

    prs = Presentation()
    prs.slide_width, prs.slide_height = px(1920), px(1080)
    blank = prs.slide_layouts[6]
    ph_cache, warns, anim_tags = {}, [], {}

    with sync_playwright() as pw:
        br = pw.chromium.launch()
        page = br.new_page(viewport={"width": 1920, "height": 1080}, device_scale_factor=2)
        page.goto(tmp.as_uri())
        page.wait_for_load_state("networkidle")
        page.evaluate("document.fonts.ready")
        n = page.evaluate("document.querySelectorAll('.__slide').length")
        for i in range(n):
            if only and (i + 1) not in only:
                continue
            items = page.evaluate(JS_EXTRACT, i)
            rasters = {}
            for it in items:
                if it["k"] == "raster":
                    rasters[it["id"]] = _shot_raster(page, i, it["id"], it.get("clip"))
            slide = prs.slides.add_slide(blank)
            w = Writer(prs, slide, rasters, slot_map, ph_cache)
            tagged = []
            for it in items:
                n0 = len(slide.shapes)
                k = it["k"]
                if k == "box":
                    w.box(it)
                elif k == "text":
                    w.text(it)
                elif k == "raster":
                    w.raster(it)
                elif k == "slot":
                    w.slot(it, slot_base)
                elif k == "img":
                    w.img(it, _local(it["src"]))
                if it.get("anim") is not None:
                    tagged += [(it["anim"], sh) for sh in list(slide.shapes)[n0:]]
            anim_tags[len(prs.slides)] = tagged
            warns += [f"P{i + 1:02d} {m}" for m in w.warn]
        br.close()

    _renumber(prs)
    prs.save(out)
    _write_anim_sidecar(out, anim_tags)
    for m in warns:
        print("  [警告]", m)
    print(f"已写出 {out}（{len(prs.slides)} 页）")


def _shot_raster(page, slide_idx, rid, clip_override=None):
    """只显示该 raster 元素，透明底截图。"""
    clip = page.evaluate(
        """([i, rid]) => {
      const s = document.querySelectorAll('.__slide')[i];
      const el = s.querySelector('[data-raster-id="'+rid+'"]');
      document.body.style.visibility = 'hidden';
      el.style.visibility = 'visible';
      const r = el.getBoundingClientRect();
      return {x: r.left + scrollX, y: r.top + scrollY, width: r.width, height: r.height};
    }""", [slide_idx, rid])
    if clip_override:
        sy = page.evaluate("scrollY")
        clip = {**clip_override, "y": clip_override["y"] + sy}
    png = b""
    if clip["width"] >= 1 and clip["height"] >= 1:
        png = page.screenshot(clip=clip, omit_background=True, full_page=True)
    page.evaluate(
        """([i, rid]) => {
      const el = document.querySelectorAll('.__slide')[i].querySelector('[data-raster-id="'+rid+'"]');
      el.style.visibility = ''; document.body.style.visibility = '';
    }""", [slide_idx, rid])
    return png or None


def _local(src):
    if src.startswith("file:"):
        from urllib.parse import unquote, urlparse
        p = unquote(urlparse(src).path)
        if re.match(r"^/[A-Za-z]:", p):
            p = p[1:]
        return pathlib.Path(p)
    return None


def _write_anim_sidecar(out, anim_tags):
    """HTML 里 data-anim="N" 标的点击分组 → <out>-动画分组.json（页号 → 按 N 排好的 shape_id 组）。
    没有任何标记的页不写（由 merge 时标 skip）。shape_id 在 _renumber 之后才定，所以放在保存后取。"""
    pages = {}
    for pg, tagged in anim_tags.items():
        groups = {}
        for n, sh in tagged:
            groups.setdefault(n, []).append(sh.shape_id)
        if groups:
            pages[str(pg)] = [groups[k] for k in sorted(groups)]
    side = pathlib.Path(out).with_name(pathlib.Path(out).stem + "-动画分组.json")
    side.write_text(json.dumps(pages, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"动画分组 → {side.name}（{len(pages)} 页有点击，合计 {sum(len(v) for v in pages.values())} 击）")


def merge_anim(pptx):
    """把 <pptx>-动画分组.json 写进 inspect/audit 产出的 <pptx>-anim.json：有标记的页用它的分组，其余页 skip。
    审查闸门字段（audit/lesson_plan/plan_digest）原样保留，所以仍须先跑 inspect_pptx 与 audit_against_plan。"""
    p = pathlib.Path(pptx)
    side = p.with_name(p.stem + "-动画分组.json")
    sheet = p.with_name(p.stem + "-anim.json")
    groups = json.loads(side.read_text(encoding="utf-8"))
    doc = json.loads(sheet.read_text(encoding="utf-8"))
    for pg, rec in doc["slides"].items():
        if pg in groups:
            rec["groups"], rec["skip"] = groups[pg], False
        else:
            rec["groups"], rec["skip"] = [], True
    sheet.write_text(json.dumps(doc, ensure_ascii=False, indent=1), encoding="utf-8")
    live = [v for v in doc["slides"].values() if not v["skip"]]
    print(f"已写入 {sheet.name}：有动画页 {len(live)}，点击合计 {sum(len(v['groups']) for v in live)}")


def _renumber(prs):
    """每页 cNvPr id 唯一（下游动画按 shape id 绑定；重复 id PowerPoint 会拒开）。"""
    for s in prs.slides:
        nid = 2
        for el in s.shapes._spTree.iter():
            if el.tag in (qn("p:cNvPr"),):
                el.set("id", str(nid))
                nid += 1


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("src", help="dc.html；--merge-anim 时为 pptx")
    ap.add_argument("out", nargs="?")
    ap.add_argument("--merge-anim", action="store_true",
                    help="把转换时写出的 -动画分组.json 并入已审查的 -anim.json（inspect+audit 之后跑）")
    ap.add_argument("--slots")
    ap.add_argument("--only", nargs="*", type=int)
    a = ap.parse_args()
    if a.merge_anim:
        merge_anim(a.src)
    else:
        convert(a.src, a.out, a.slots, a.only)


if __name__ == "__main__":
    main()

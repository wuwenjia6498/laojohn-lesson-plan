"""illustrate_pptx.py - 给已审核、已注入动画的写作课 pptx 添插图（2026-09-25 六上五试点起草）。

风格样板：外部人工添图版（记忆 ppt-illustration-style-reference-0925）——
每课一个吉祥物立在页题条右端，主图要大、常出血，正文给图让出一栏，必要时衬一条浅色底带。
本工具只动图与几何：**不改文字**（文字审查归外部 PPT 后处理链）。

用法：
    python illustrate_pptx.py <版式清单.json>
    python illustrate_pptx.py --retime <配图版.pptx> [--out 另存.pptx]   # 只改动画，不动图与版式

版式清单（路径均相对仓库根）：
{
  "src": "写作课件PPT输出/<课次>/<课次>-课件PPT.pptx",
  "out": "C:/Users/.../Desktop/<课次>-课件PPT-配图版.pptx",
  "cache": "课件配图工具/课件产出/<项目>/gpt-image/ppt配图缓存",
  "mascot": {"img": "…/角色/吉祥物-招手.jpg", "h": 2.1, "x": 14.15,
             "bar_right": 15.3, "skip": [1]},          # 自动找页题条（左1.15/顶1.56/高0.91）
  "pages": {
    "6": [ {"op": "band", "box": [0, 3.2, 20, 4.7], "color": "FFF6E5"},
           {"op": "swap", "id": 16, "img": "…/P06.jpg", "box": [x, y, w, h], "anchor": "r", "cutout": true} ],
    ...
  }
}
op 一览：
  swap  在原图位形状里换图（shape id 不变 → 动画绑定保留），并按 box/anchor 重设几何
  add   新增图片（不进动画，随页静态出现）
  geom  改已有形状几何：id + 任意 x/y/w/h（绝对值）或 dx/dy（位移）
  band  通栏浅色底带：插到最底层，不遮任何内容；可选 "round": 半径英寸（圆角底卡）、"line": 描边色
  font  只改某形状内所有 run 的字号（size＝磅），不动文字
  fill  改形状填充色（卡片表头换彩色）
swap／add 可带 "click": N（1 基）：把这张图并进本页第 N 击，与那一击的文字同时出现（0926 用户定：
  图对应某张卡／某一条的，翻页不能先露出来；会露答案的图跟答案那一击）。不写 click＝静态，开场就在，
  只用于整页情境图。按原有点击组重注入，走 helpers.add_click_reveal（单一源，禁复制时间树）。
  ⚠ 图的顶端高于上一击时会造成回跳，这种图宁可并进更早的一击。
  click 另可单独作 op：{"op": "click", "id": 66, "click": 3}——把页上已有的形状（如外部件自带的图）并进第 3 击。
  del：{"op": "del", "id": 53}——删掉一个形状（如判断页上会露答案的图，挪到结果页重新 add）；
  被删的形状若在动画里会报错退出，不静默拆动画。
分节页与结尾页一律不加图（is_section 自动跳过，0925 用户定）。
全局可选：theme（页题条色＋分节页色）、tip_icon（页脚提示句前加手指图标）、mascot、
  effect（动画效果，缺省 "appear"＝出现；"fade"＝淡入）。
动画收尾（retime，2026-09-30 用户定）：全篇点击动画一律改用 effect 指定的效果；手指图标并进它右侧那句
  页脚提示的那一击，与文字一起出现（提示句常驻则图标也常驻）。已出的配图版用 --retime 补做，
  不必重跑整份清单（重跑会冲掉 finalize_external 补的稿纸页、删的讲评页）。
图源是透明底 PNG（gen_cutouts.py 直出）时按不透明区域裁边，不走抠图。
图源可带 "crop": [x0, y0, x1, y1]（像素，先裁再处理）；"cutout": true 把与四边连通的底（白底与浅色水彩底晕）抠成透明，
"keep": "br" 表示右、下两边不起抠（出血贴页边的那侧）。
box 超出页面即出血；anchor 取 c/t/b/l/r 组合（如 br = 右下贴齐）。

⚠ 收尾只调 pptx_text_edit.prune_timing，**不调 normalize_paragraphs**：外部件常见一段多个 pPr，
  它会把文件改成 PowerPoint 打不开（记忆 pptx-text-edit-three-silent-traps-0916 的 0925 反例）。
"""
import hashlib
import importlib.util
import json
import pathlib
import sys

from PIL import Image, ImageChops, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

HERE = pathlib.Path(__file__).resolve().parent
_spec = importlib.util.spec_from_file_location("pte", HERE / "pptx_text_edit.py")
pte = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(pte)
sys.path.insert(0, str(HERE.parent / "scripts"))
from helpers import add_click_reveal  # noqa: E402  单一源，禁复制
from pptx.oxml.ns import qn  # noqa: E402

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
E = 914400


# ───────────── 图片预处理 ─────────────
def _trim(im, pad=0.02):
    """裁掉四周与左上角同色的边（白底或浅色底）。"""
    rgb = im.convert("RGB")
    bg = rgb.getpixel((2, 2))
    diff = Image.new("L", rgb.size, 0)
    px, dp = rgb.load(), diff.load()
    w, h = rgb.size
    for y in range(0, h, 2):
        for x in range(0, w, 2):
            r, g, b = px[x, y]
            if abs(r - bg[0]) + abs(g - bg[1]) + abs(b - bg[2]) > 36:
                dp[x, y] = 255
    box = diff.getbbox()
    if not box:
        return im
    p = int(max(w, h) * pad)
    return im.crop((max(0, box[0] - p), max(0, box[1] - p), min(w, box[2] + p), min(h, box[3] + p)))


def _cutout(im, keep=""):
    """与四边连通的「底」变透明。底＝高亮且低饱和的像素：白底之外，把人物身后的米色／浅绿
    水彩底晕也一并算作底（只认纯白会留下一块块晕染，0925 用户点名「右下角没抠干净」）。
    keep 里列出的边（t/b/l/r）不从该边起抠：出血贴页边的那一侧，桌子稿纸本就该被页边截断，
    从那边起抠反会把贴边的实物抠穿。只从「是底」的边缘点起泛洪，贴边的头发不会被当成底。"""
    rgb = im.convert("RGB")
    w, h = rgb.size
    mn = ImageChops.darker(ImageChops.darker(*rgb.split()[:2]), rgb.split()[2])
    mx = ImageChops.lighter(ImageChops.lighter(*rgb.split()[:2]), rgb.split()[2])
    chroma = ImageChops.subtract(mx, mn)
    bright = mn.point(lambda v: 255 if v >= 200 else 0)
    pale = chroma.point(lambda v: 255 if v <= 48 else 0)
    bgmask = ImageChops.multiply(bright, pale)          # 255＝像底
    step = max(6, min(w, h) // 80)
    seeds = []
    if "t" not in keep: seeds += [(x, 0) for x in range(0, w, step)]
    if "b" not in keep: seeds += [(x, h - 1) for x in range(0, w, step)]
    if "l" not in keep: seeds += [(0, y) for y in range(0, h, step)]
    if "r" not in keep: seeds += [(w - 1, y) for y in range(0, h, step)]
    for s in seeds:
        if bgmask.getpixel(s) == 255:
            ImageDraw.floodfill(bgmask, s, 128)
    hole = bgmask.point(lambda v: 255 if v == 128 else 0)
    hole = hole.filter(ImageFilter.MaxFilter(5))        # 向内多吃两像素，去掉浅色描边光晕
    alpha = ImageChops.invert(hole).filter(ImageFilter.GaussianBlur(1.3))
    out = rgb.convert("RGBA")
    out.putalpha(alpha)
    return out


def prep(src, cache, crop=None, cutout=False, edge=1200, keep=""):
    src = pathlib.Path(src)
    key = hashlib.md5(f"{src}|{src.stat().st_mtime}|{crop}|{cutout}|{edge}|{keep}|v3".encode()).hexdigest()[:10]
    out = pathlib.Path(cache) / f"{src.stem}-{key}.{'png' if (cutout or src.suffix.lower() == '.png') else 'jpg'}"
    if out.exists():
        return out
    im = Image.open(src)
    if crop:
        im = im.crop(tuple(crop))
    has_alpha = im.mode in ("RGBA", "LA") and im.getchannel("A").getextrema()[0] < 250
    if has_alpha:                       # 生图直出的透明底：按不透明区域裁边，不再抠
        box = im.getchannel("A").point(lambda v: 255 if v > 16 else 0).getbbox()
        im = im.crop(box)
        cutout = "native"
    else:
        im = _trim(im)
    s = edge / max(im.size)
    if s < 1:
        im = im.resize((round(im.width * s), round(im.height * s)), Image.LANCZOS)
    out.parent.mkdir(parents=True, exist_ok=True)
    if cutout == "native":
        im.convert("RGBA").quantize(256, method=Image.FASTOCTREE).save(out, optimize=True)
    elif cutout:
        _cutout(im, keep).quantize(256, method=Image.FASTOCTREE).save(out, optimize=True)
    else:
        im.convert("RGB").save(out, quality=88)
    return out


# ───────────── 版面操作 ─────────────
def fit(img, box, anchor="c"):
    x, y, w, h = box
    iw, ih = Image.open(img).size
    s = min(w / iw, h / ih)
    nw, nh = iw * s, ih * s
    nx, ny = x + (w - nw) / 2, y + (h - nh) / 2
    if "l" in anchor: nx = x
    if "r" in anchor: nx = x + w - nw
    if "t" in anchor: ny = y
    if "b" in anchor: ny = y + h - nh
    return nx, ny, nw, nh


def shape(slide, sid):
    for sh in slide.shapes:
        if sh.shape_id == sid:
            return sh
    raise KeyError(f"shape id {sid} 不在该页")


def set_geom(sh, x=None, y=None, w=None, h=None):
    if x is not None: sh.left = Inches(x)
    if y is not None: sh.top = Inches(y)
    if w is not None: sh.width = Inches(w)
    if h is not None: sh.height = Inches(h)


def op_swap(slide, o, img):
    pic = shape(slide, o["id"])
    _, rid = slide.part.get_or_add_image_part(str(img))
    blip = pic._element.blipFill.find(A + "blip")
    old = blip.get(R_EMBED)
    blip.set(R_EMBED, rid)
    src = pic._element.blipFill.find(A + "srcRect")
    if src is not None:
        pic._element.blipFill.remove(src)
    if old != rid and not slide._element.xpath('.//a:blip[@r:embed="%s"]' % old):
        slide.part.drop_rel(old)
    set_geom(pic, *fit(img, o["box"], o.get("anchor", "c")))


def op_add(slide, o, img):
    x, y, w, h = fit(img, o["box"], o.get("anchor", "c"))
    p = slide.shapes.add_picture(str(img), Inches(x), Inches(y), Inches(w), Inches(h))
    p.name = o.get("name", "插图")
    return p


def op_band(slide, o):
    x, y, w, h = o["box"]
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if o.get("round") else MSO_SHAPE.RECTANGLE
    b = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if o.get("round"):
        b.adjustments[0] = min(0.5, o["round"] / min(w, h))    # round＝圆角半径（英寸）
    b.fill.solid()
    b.fill.fore_color.rgb = RGBColor.from_string(o.get("color", "FFF6E5"))
    if o.get("line"):
        b.line.color.rgb = RGBColor.from_string(o["line"])
        b.line.width = Inches(0.02)
    else:
        b.line.fill.background()
    b.shadow.inherit = False
    b.name = "底带"
    tree = slide.shapes._spTree
    tree.remove(b._element)
    tree.insert(2, b._element)          # 紧跟 nvGrpSpPr/grpSpPr 之后＝最底层


def set_fill(sh, color):
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(color)


def footer_tips(slide):
    """页脚提示句：左 1.15、顶 ≥ 9.6、宽 ≥ 10 的文本框（图例条之类的短框不算）。"""
    return [sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()
            and abs(sh.left / E - 1.15) < 0.12 and sh.top / E >= 9.6 and sh.width / E >= 10]


def is_section(slide, no, total):
    """分节页／结尾页：无页题条、且有一块占满大半页面的底板。封面（第 1 页）不算。"""
    if no == 1 or title_bar(slide) is not None:
        return False
    return any(sh.width / E > 15 and sh.height / E > 8 for sh in slide.shapes if sh.shape_type == 1)


def title_bar(slide):
    for sh in slide.shapes:
        if sh.shape_type == 1 and abs(sh.left / E - 1.15) < 0.1 and abs(sh.top / E - 1.56) < 0.1 \
                and abs(sh.height / E - 0.91) < 0.05:
            return sh
    return None


def click_groups(slide):
    """读出本页现有点击组：[[target, ...], ...]，target＝shape_id 或 (shape_id, 段落号)。"""
    t = slide.element.find(qn("p:timing"))
    if t is None:
        return []
    gs = []
    for ctn in t.iter(qn("p:cTn")):
        nt = ctn.get("nodeType")
        if nt not in ("clickEffect", "withEffect", "afterEffect"):
            continue
        sp = ctn.find(".//" + qn("p:spTgt"))
        rg = sp.find(".//" + qn("p:pRg"))
        tgt = (int(sp.get("spid")), int(rg.get("st"))) if rg is not None else int(sp.get("spid"))
        if nt == "clickEffect" or not gs:
            gs.append([])
        if tgt not in gs[-1]:
            gs[-1].append(tgt)
    return gs


def pair_tip_icons(slide, gs):
    """手指图标并进右侧那句页脚提示所在的一击；提示句不在动画里则图标照旧常驻。返回并入个数。"""
    n = 0
    for icon in [sh for sh in slide.shapes if sh.name == "提示图标"]:
        for sh in slide.shapes:
            if not (sh.has_text_frame and 0 < (sh.left - icon.left) / E < 1.2
                    and abs(sh.top - icon.top) / E < 0.4):
                continue
            k = next((i for i, g in enumerate(gs)
                      if any((t if isinstance(t, int) else t[0]) == sh.shape_id for t in g)), None)
            if k is None:
                break
            for g in gs:
                if icon.shape_id in g:
                    g.remove(icon.shape_id)
            gs[k].append(icon.shape_id)
            n += 1
            break
    return n


def retime(prs, effect="appear"):
    """全篇重注入点击动画：手指图标随提示句同击，效果统一为 effect。分组与顺序不变。"""
    pages = icons = 0
    for slide in prs.slides:
        gs = click_groups(slide)
        if not gs:
            continue
        icons += pair_tip_icons(slide, gs)
        gs = [g for g in gs if g]
        for tag in ("p:timing", "p:bldLst"):
            for node in slide.element.findall(qn(tag)):
                slide.element.remove(node)
        add_click_reveal(slide, gs, dur=500, effect=effect)
        pages += 1
    print(f"动画：{pages} 页改为「{'出现' if effect == 'appear' else '淡入'}」，手指图标随提示句同击 {icons} 处")


def join_clicks(slide, no, joins):
    """joins：[(shape_id, 第几击)]。并进原有点击组后整页重注入。"""
    gs = click_groups(slide)
    for sid, k in joins:
        if not 1 <= k <= len(gs):
            raise SystemExit(f"P{no:02d} 只有 {len(gs)} 击，click={k} 越界")
        for g in gs:                      # 已在别的击里就是「挪」：先移出原击
            if sid in g:
                g.remove(sid)
        gs[k - 1].append(sid)
    gs = [g for g in gs if g]             # 挪空的击随之消失（击数减少，序号按挪之前算）
    for tag in ("p:timing", "p:bldLst"):
        for node in slide.element.findall(qn(tag)):
            slide.element.remove(node)
    add_click_reveal(slide, gs, dur=500)


def run(cfg_path):
    cfg = json.loads(pathlib.Path(cfg_path).read_text(encoding="utf-8"))
    cache = cfg["cache"]
    prs = Presentation(cfg["src"])
    for key, ops in cfg.get("pages", {}).items():
        slide = prs.slides[int(key) - 1]
        if is_section(slide, int(key), len(prs.slides)):
            print(f"  [跳过] P{int(key):02d} 是分节页／结尾页：不加图（0925 用户定，全线适用）")
            continue
        joins = []
        for o in ops:
            if o["op"] in ("swap", "add"):
                img = prep(o["img"], cache, o.get("crop"), o.get("cutout", False), o.get("edge", 1200), o.get("keep", ""))
                if o["op"] == "swap":
                    op_swap(slide, o, img)
                    sid = o["id"]
                else:
                    sid = op_add(slide, o, img).shape_id
                if o.get("click"):
                    joins.append((sid, o["click"]))
            elif o["op"] == "geom":
                sh = shape(slide, o["id"])
                if "dx" in o: sh.left = sh.left + Inches(o["dx"])
                if "dy" in o: sh.top = sh.top + Inches(o["dy"])
                set_geom(sh, o.get("x"), o.get("y"), o.get("w"), o.get("h"))
            elif o["op"] == "click":
                shape(slide, o["id"])
                joins.append((o["id"], o["click"]))
            elif o["op"] == "del":
                used = {t if isinstance(t, int) else t[0] for g in click_groups(slide) for t in g}
                if o["id"] in used:
                    raise SystemExit(f"P{int(key):02d} 形状 {o['id']} 在动画里，不能删")
                el = shape(slide, o["id"])._element
                el.getparent().remove(el)
            elif o["op"] == "band":
                op_band(slide, o)
            elif o["op"] == "fill":
                set_fill(shape(slide, o["id"]), o["color"])
            elif o["op"] == "font":
                # 只改字号（格式），不改文字
                for r in shape(slide, o["id"])._element.iter(A + "rPr"):
                    r.set("sz", str(int(o["size"] * 100)))
            else:
                raise ValueError(f"未知 op：{o['op']}")
        if joins:
            join_clicks(slide, int(key), joins)
    th = cfg.get("theme")
    if th:
        for i, slide in enumerate(prs.slides, 1):
            bar = title_bar(slide)
            if bar is not None and th.get("bar"):
                set_fill(bar, th["bar"])
            for sp in th.get("sections", []):
                if sp["page"] == i:
                    set_fill(shape(slide, sp["id"]), th["section"])
    tip = cfg.get("tip_icon")
    if tip:
        n = 0
        for i, slide in enumerate(prs.slides, 1):
            if i in tip.get("skip", []) or title_bar(slide) is None:
                continue
            for sh in footer_tips(slide):
                sz = tip.get("size", 0.6)
                slide.shapes.add_picture(tip["img"], sh.left, sh.top - Inches(0.04), Inches(sz), Inches(sz)).name = "提示图标"
                sh.left = sh.left + Inches(sz + 0.15)
                sh.width = sh.width - Inches(sz + 0.15)
                n += 1
        print(f"提示图标：{n} 处")
    m = cfg.get("mascot")
    if m:
        img = prep(m["img"], cache, m.get("crop"), True, 700)
        iw, ih = Image.open(img).size
        n = 0
        for i, slide in enumerate(prs.slides, 1):
            if i in m.get("skip", []):
                continue
            bar = title_bar(slide)
            if bar is None:
                continue
            right = max((bar.left + bar.width) / E, m.get("bar_right", 15.3))
            bar.width = Inches(right - bar.left / E)
            h = m.get("h", 2.1)
            w = h * iw / ih
            bottom = (bar.top + bar.height) / E
            p = slide.shapes.add_picture(str(img), Inches(m.get("x", 14.15)), Inches(bottom - h), Inches(w), Inches(h))
            p.name = "吉祥物"
            n += 1
        print(f"吉祥物：{n} 页")
    retime(prs, cfg.get("effect", "appear"))
    pte.prune_timing(prs)
    out = pathlib.Path(cfg["out"])
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print(f"已写 {out}  {out.stat().st_size / 1e6:.2f} MB")


def retime_file(src, out=None, effect="appear"):
    prs = Presentation(src)
    retime(prs, effect)
    pte.prune_timing(prs)
    prs.save(out or src)
    print(f"已写 {out or src}")


if __name__ == "__main__":
    if sys.argv[1] == "--retime":
        args = sys.argv[2:]
        out = args[args.index("--out") + 1] if "--out" in args else None
        eff = args[args.index("--effect") + 1] if "--effect" in args else "appear"
        retime_file(args[0], out, eff)
    else:
        run(sys.argv[1])

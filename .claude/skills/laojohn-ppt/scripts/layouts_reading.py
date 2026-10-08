# -*- coding: utf-8 -*-
"""读书会 profile 的 6 种页型渲染（整本书阅读读书会）。

每个函数签名：render_<type>(slide, page, ctx)
公共件（铭牌/LOGO/封面基函数/END/占位）来自 layouts_common，共享、课型无关。
本模块只放读书会专属的页型版式；写作课的在 layouts_writing。
"""
from pptx.util import Emu, Pt
from helpers import (
    add_textbox, add_rect, add_image_grid, add_table, cell_bg_color, rgb,
    _disable_shape_effects,
    add_numbered_bullets, add_click_reveal,
    add_rich_textbox, add_qa_textbox,
    set_ascii_font,
)
from theme import (
    FONT_TITLE, FONT_BODY, FONT_QUOTE,
    COLOR_TITLE, COLOR_BODY, COLOR_MUTED,
    COLOR_BG_QUOTE,
    COLOR_SECTION_BG, COLOR_SECTION_FG, COLOR_SECTION_SUB, COLOR_SECTION_NUM,
    COLOR_RED_ACCENT,
    SZ_HEADING, SZ_BODY, SZ_GUIDE_BULLET, SZ_QUOTE, SZ_SUBTITLE,
    GUIDE_TITLE_Y, GUIDE_TITLE_H,
    GUIDE_BULLETS_X, GUIDE_BULLETS_Y, GUIDE_BULLETS_W, GUIDE_BULLETS_H,
    QUOTE_BG_X, QUOTE_BG_Y, QUOTE_BG_W, QUOTE_BG_H,
    QUOTE_TITLE_X, QUOTE_TITLE_Y, QUOTE_TITLE_W, QUOTE_TITLE_H,
    QUOTE_BODY_X, QUOTE_BODY_Y, QUOTE_BODY_W, QUOTE_BODY_H,
    SUMMARY_TITLE_X, SUMMARY_TITLE_Y, SUMMARY_TITLE_W, SUMMARY_TITLE_H,
    SUMMARY_BULLETS_X, SUMMARY_BULLETS_Y, SUMMARY_BULLETS_W, SUMMARY_BULLETS_H,
    IMAGE_GRID_REGION, IMAGE_GRID_GAP,
    TABLE_TITLE_X, TABLE_TITLE_Y, TABLE_TITLE_W, TABLE_TITLE_H,
    TABLE_SUBTITLE_X, TABLE_SUBTITLE_Y, TABLE_SUBTITLE_W, TABLE_SUBTITLE_H,
    TABLE_AREA_X, TABLE_AREA_Y, TABLE_AREA_W, TABLE_AREA_H,
    LOGO_INNER_X, LOGO_INNER_Y, LOGO_INNER_W,
    pct_x, pct_y,
)
from helpers import add_logo
from theme import resolve_reading_theme
import math
import layouts_reading_v9 as v9
from layouts_common import (
    draw_anchor, draw_logo_inner, draw_page_num, maybe_placeholder,
    bullets_text, _real_suggestions, _split_title_keywords,
    _render_cover_base, render_end, legend_to_colors,
)


# ---------- 封面 ----------
def _v9(ctx):
    """声明了 `主题色` → ReadingTheme（走 v9 新样式）；未声明 → None（现行样式原样）。"""
    return resolve_reading_theme(ctx.get("theme_color"))


def render_cover(slide, page, ctx):
    th = _v9(ctx)
    if th:
        return v9.render_cover(slide, page, ctx, th)
    # 读书会：标题是书名，自动补《》
    _render_cover_base(slide, page, ctx, wrap_brackets=True)


# ---------- 环节标题 ----------
def render_section(slide, page, ctx):
    """环节标题：多环节显示巨型序号；单环节自动隐藏序号、居中标题。"""
    th = _v9(ctx)
    if th:
        return v9.render_section(slide, page, ctx, th)
    section_total = getattr(page, "_section_total", 1)
    is_solo = (section_total <= 1)

    # 深灰蓝色块（共用）
    add_rect(slide, pct_x(0.025), pct_y(0.040),
             pct_x(0.950), pct_y(0.920),
             COLOR_SECTION_BG)

    if not is_solo:
        # 多环节：保留巨型序号
        section_no = getattr(page, "_section_no", 1)
        section_box = add_textbox(
            slide, pct_x(0.08), pct_y(0.17), pct_x(0.35), pct_y(0.30),
            f"{section_no:02d}",
            font=FONT_TITLE, size=200, color=COLOR_SECTION_NUM,
            bold=False, italic=True, align="left", anchor="middle",
        )
        set_ascii_font(section_box, "Bahnschrift")
        add_rect(slide, pct_x(0.08), pct_y(0.51), pct_x(0.08), Emu(38000),
                 COLOR_RED_ACCENT)
        add_textbox(
            slide, pct_x(0.08), pct_y(0.55), pct_x(0.84), pct_y(0.16),
            page.title,
            font=FONT_TITLE, size=44, color=COLOR_SECTION_FG, bold=True,
            align="left", anchor="middle", line_spacing=1.3,
        )
        if page.body:
            add_textbox(
                slide, pct_x(0.08), pct_y(0.73), pct_x(0.84), pct_y(0.16),
                page.body,
                font=FONT_BODY, size=22, color=COLOR_SECTION_SUB,
                align="left", anchor="top", line_spacing=1.6,
            )
    else:
        # 单环节：去掉序号，居中大标题 + 红色短分隔 + 副本
        add_textbox(
            slide, pct_x(0.08), pct_y(0.34), pct_x(0.84), pct_y(0.14),
            page.title,
            font=FONT_TITLE, size=54, color=COLOR_SECTION_FG, bold=True,
            align="center", anchor="middle", line_spacing=1.3,
        )
        add_rect(slide, pct_x(0.46), pct_y(0.50), pct_x(0.08), Emu(38000),
                 COLOR_RED_ACCENT)
        if page.body:
            add_textbox(
                slide, pct_x(0.15), pct_y(0.55), pct_x(0.70), pct_y(0.20),
                page.body,
                font=FONT_BODY, size=22, color=COLOR_SECTION_SUB,
                align="center", anchor="top", line_spacing=1.6,
            )

    # LOGO（页码已取消：环节标题页也不再标注右下角页码）
    logo_white = ctx.get("logo_white_path") or ctx.get("logo_path")
    if logo_white:
        add_logo(slide, logo_white, LOGO_INNER_X, LOGO_INNER_Y, LOGO_INNER_W)


# ---------- 引导问题 ----------
def render_guide(slide, page, ctx):
    """引导问题（方案 A）：
    - 标题左侧加一根红色 ▎竖色条作为视觉引导
    - 其余排版（铭牌/要点/占位）沿用

    四图网格模式：当本页带 ≥2 条配图建议时，标题下整幅排 2×2 占位网格
    （"给你们看 N 幅画面"类页），此模式不渲染要点文本。
    """
    th = _v9(ctx)
    if th:
        return v9.render_guide(slide, page, ctx, th)
    grid_sugs = _real_suggestions(page)
    grid_mode = len(grid_sugs) >= 2

    draw_anchor(slide, page.eyebrow)
    draw_logo_inner(slide, ctx.get("logo_path"))
    draw_page_num(slide, ctx["page_index"], ctx["page_total"])

    # 标题左侧红色竖色条 ▎
    add_rect(slide, pct_x(0.070), pct_y(0.245), Emu(50000), pct_y(0.130),
             COLOR_RED_ACCENT)

    # 标题（整体右移给红色竖条让出位置）
    title_segments = []
    for txt, is_kw in _split_title_keywords(page.title):
        if is_kw:
            title_segments.append((txt, {"color": COLOR_RED_ACCENT, "bold": True}))
        else:
            title_segments.append((txt, {"color": COLOR_TITLE, "bold": True}))
    add_rich_textbox(
        slide, pct_x(0.090), GUIDE_TITLE_Y, pct_x(0.84), GUIDE_TITLE_H,
        title_segments,
        font=FONT_TITLE, size=SZ_HEADING,
        align="left", anchor="middle", line_spacing=1.3,
    )

    # 四图网格模式：标题下整幅排占位网格，不渲染要点/单图占位
    if grid_mode:
        add_image_grid(slide, *IMAGE_GRID_REGION, grid_sugs, gap=IMAGE_GRID_GAP)
        return

    # 正文(引文)与要点(追问)可同页共存：引文在上、追问在下
    # （契约 field-extraction「引文+追问」模式；旧版 if/elif 会静默丢弃正文）
    # 若要点带参考答案（page.bullet_answers），追问改走"问题+红色答案"交错版式
    # （add_qa_textbox），先出问题、点击后出红字答案。
    answers = getattr(page, "bullet_answers", None) or []
    has_ans = any(a for a in answers)
    bullets_box = None
    qa_groups = None

    def _render_bullets(bx, by, bw, bh):
        nonlocal bullets_box, qa_groups
        if has_ans:
            bullets_box, qa_groups = add_qa_textbox(
                slide, bx, by, bw, bh,
                page.bullets, answers,
                text_size=SZ_GUIDE_BULLET, font=FONT_TITLE,
            )
        else:
            bullets_box = add_textbox(
                slide, bx, by, bw, bh,
                bullets_text(page.bullets),
                font=FONT_TITLE, size=SZ_GUIDE_BULLET, color=COLOR_BODY,
                line_spacing=1.8, align="left", anchor="top",
            )

    if page.body and page.bullets:
        add_textbox(
            slide, GUIDE_BULLETS_X, GUIDE_BULLETS_Y, GUIDE_BULLETS_W, pct_y(0.20),
            page.body,
            font=FONT_BODY, size=SZ_BODY, color=COLOR_BODY,
            line_spacing=1.6, align="left", anchor="top",
        )
        _render_bullets(GUIDE_BULLETS_X, pct_y(0.66), GUIDE_BULLETS_W, pct_y(0.24))
    elif page.bullets:
        _render_bullets(GUIDE_BULLETS_X, GUIDE_BULLETS_Y, GUIDE_BULLETS_W, GUIDE_BULLETS_H)
    elif page.body:
        add_textbox(
            slide, GUIDE_BULLETS_X, GUIDE_BULLETS_Y, GUIDE_BULLETS_W, GUIDE_BULLETS_H,
            page.body,
            font=FONT_BODY, size=SZ_BODY, color=COLOR_BODY,
            line_spacing=1.8, align="left", anchor="top",
        )

    # 逐条点击：每点一次出一段（问题/答案各一击）；≥2 段才启用，单段直出
    if ctx.get("anim"):
        if qa_groups is not None and len(qa_groups) >= 2:
            add_click_reveal(slide, qa_groups)
        elif qa_groups is None and bullets_box is not None and len(page.bullets) >= 2:
            spid = bullets_box.shape_id
            add_click_reveal(slide, [[(spid, i)] for i in range(len(page.bullets))])

    maybe_placeholder(slide, page)


# ---------- 原文齐读 ----------
_EMU_PER_PT = 12700.0

# 超长引文的降档梯（字号, 行距）：默认 24pt/1.8 放得下就完全不动（向后兼容）；
# 放不下才逐档降，宁可小一点也不让引文出血到书页底外（十问全文这类整段照录页用）。
_QUOTE_FIT_LADDER = (
    (SZ_QUOTE, 1.8), (22, 1.65), (20, 1.5), (18, 1.4),
    (17, 1.35), (16, 1.3), (15, 1.3), (14, 1.25),
)


def _quote_fit(text):
    """按字符数估算引文在正文区的占高，从默认档起逐档试到放得下为止。
    保守系数与写作 profile 示范文同源（中文每字约 1.15 倍字号宽、行高略大于标称）。"""
    import math
    width_pt = QUOTE_BODY_W / _EMU_PER_PT
    avail_pt = QUOTE_BODY_H / _EMU_PER_PT
    lines = (text or "").split("\n")
    for size, spacing in _QUOTE_FIT_LADDER:
        cpl = max(1, int(width_pt / (size * 1.15)))
        total = sum(max(1, math.ceil((len(ln) + 2) / cpl)) for ln in lines)
        if total * size * spacing * 1.1 <= avail_pt:
            return size, spacing
    return _QUOTE_FIT_LADDER[-1]


def render_quote(slide, page, ctx):
    th = _v9(ctx)
    if th:
        return v9.render_quote(slide, page, ctx, th)
    draw_anchor(slide, page.eyebrow)
    draw_logo_inner(slide, ctx.get("logo_path"))
    draw_page_num(slide, ctx["page_index"], ctx["page_total"])

    # 书页底
    add_rect(slide, QUOTE_BG_X, QUOTE_BG_Y, QUOTE_BG_W, QUOTE_BG_H, COLOR_BG_QUOTE)

    if page.title:
        add_textbox(
            slide, QUOTE_TITLE_X, QUOTE_TITLE_Y, QUOTE_TITLE_W, QUOTE_TITLE_H,
            page.title,
            font=FONT_TITLE, size=SZ_HEADING - 4, color=COLOR_TITLE, bold=True,
            align="left", anchor="middle",
        )
    if page.body:
        body_size, body_spacing = _quote_fit(page.body)
        add_textbox(
            slide, QUOTE_BODY_X, QUOTE_BODY_Y, QUOTE_BODY_W, QUOTE_BODY_H,
            page.body,
            font=FONT_QUOTE, size=body_size, color=COLOR_TITLE,
            line_spacing=body_spacing, first_line_indent_chars=2,
            align="left", anchor="top",
        )


# ---------- 要点小结 ----------
def render_summary(slide, page, ctx):
    """要点小结（方案 A）：红方块编号 + 文字，永远单列。"""
    th = _v9(ctx)
    if th:
        return v9.render_summary(slide, page, ctx, th)
    draw_anchor(slide, page.eyebrow)
    draw_logo_inner(slide, ctx.get("logo_path"))
    draw_page_num(slide, ctx["page_index"], ctx["page_total"])

    add_textbox(
        slide, SUMMARY_TITLE_X, SUMMARY_TITLE_Y, SUMMARY_TITLE_W, SUMMARY_TITLE_H,
        page.title,
        font=FONT_TITLE, size=SZ_HEADING, color=COLOR_TITLE, bold=True,
        align="left", anchor="middle", line_spacing=1.3,
    )

    # 正文(引文)与要点可同页共存：引文在上、要点在下
    # （契约 field-extraction「引文+追问」模式；旧版 if/elif 会静默丢弃正文）
    # 若要点带参考答案（page.bullet_answers），改走"问题+红色答案"交错版式
    # （add_qa_textbox 圈号问题在上、红字答案在下、逐段点击），不再用红方块编号。
    answers = getattr(page, "bullet_answers", None) or []
    has_ans = any(a for a in answers)
    bullet_groups = None   # 红方块编号路径的点击分组
    qa_groups = None       # 问答交错路径的逐段点击分组

    def _render_bullets(bx, by, bw, bh):
        nonlocal bullet_groups, qa_groups
        if has_ans:
            _box, qa_groups = add_qa_textbox(
                slide, bx, by, bw, bh,
                page.bullets, answers,
                text_size=SZ_GUIDE_BULLET, font=FONT_TITLE,
            )
        else:
            bullet_groups = add_numbered_bullets(
                slide, bx, by, bw, bh,
                page.bullets, columns=1, text_size=SZ_GUIDE_BULLET, text_font=FONT_TITLE,
            )

    if page.body and page.bullets:
        add_textbox(
            slide, SUMMARY_BULLETS_X, SUMMARY_BULLETS_Y, SUMMARY_BULLETS_W, pct_y(0.18),
            page.body,
            font=FONT_BODY, size=SZ_BODY, color=COLOR_BODY,
            line_spacing=1.6, align="left", anchor="top",
        )
        _render_bullets(SUMMARY_BULLETS_X, pct_y(0.61), SUMMARY_BULLETS_W, pct_y(0.31))
    elif page.bullets:
        _render_bullets(SUMMARY_BULLETS_X, SUMMARY_BULLETS_Y, SUMMARY_BULLETS_W, SUMMARY_BULLETS_H)
    elif page.body:
        add_textbox(
            slide, SUMMARY_BULLETS_X, SUMMARY_BULLETS_Y, SUMMARY_BULLETS_W, SUMMARY_BULLETS_H,
            page.body,
            font=FONT_BODY, size=SZ_BODY, color=COLOR_BODY,
            line_spacing=1.8, align="left", anchor="top",
        )

    # 逐条点击：每点一次出一组（≥2 组才启用）
    if ctx.get("anim"):
        if qa_groups is not None and len(qa_groups) >= 2:
            add_click_reveal(slide, qa_groups)
        elif bullet_groups and len(bullet_groups) >= 2:
            add_click_reveal(slide, bullet_groups)

    maybe_placeholder(slide, page)


# ---------- 填空表格 ----------
TABLE_GRID_COLOR = "BFBFBF"     # 填写表格线：白幕投屏可见、不抢表头主题色


def render_table(slide, page, ctx):
    th = _v9(ctx)       # v9 只换色（眉标色块/表头），版式不变
    if th:
        draw_anchor(slide, page.eyebrow, bar_color=th.main)
    else:
        draw_anchor(slide, page.eyebrow)
    draw_logo_inner(slide, ctx.get("logo_path"))
    draw_page_num(slide, ctx["page_index"], ctx["page_total"])

    # 副标题折成多行时，标题与副标题整体上移、副标题框加高，表格位置不动（免得压到表头）
    sub_lines = 1
    if page.subtitle:
        cpl = max(1, int(TABLE_SUBTITLE_W / 12700 / SZ_SUBTITLE))
        sub_lines = max(1, math.ceil(len(page.subtitle) / cpl - 0.1))
    extra = int((sub_lines - 1) * SZ_SUBTITLE * 1.6 * 12700)
    add_textbox(
        slide, TABLE_TITLE_X, TABLE_TITLE_Y - extra, TABLE_TITLE_W, TABLE_TITLE_H,
        page.title,
        font=FONT_TITLE, size=SZ_HEADING - 2, color=COLOR_TITLE, bold=True,
        align="left", anchor="middle",
    )
    if page.subtitle:
        add_textbox(
            slide, TABLE_SUBTITLE_X, TABLE_SUBTITLE_Y - extra, TABLE_SUBTITLE_W,
            TABLE_SUBTITLE_H + extra,
            page.subtitle,
            font=FONT_BODY, size=SZ_SUBTITLE, color=COLOR_MUTED,
            align="left", anchor="middle",
        )
    if not page.table_headers:
        return

    # 字号/行高/列宽由 add_table 按行数与内容自适应，确保整张表塞进区域。
    # 答案格（page.table_reveals）：底表渲占位、答案做成叠层逐格点击淡入。
    # 分类上色（可选）：给了 `图例` 时，单元格 `[码:片段]` 按码上色（如五感观察表的
    # 感官列），颜色与示范文/旁批表同源；读书会表无 `图例` → 原样。
    reveals = page.table_reveals or {}
    cell_colors = legend_to_colors(page.legend) or None
    _table, geom = add_table(
        slide, TABLE_AREA_X, TABLE_AREA_Y, TABLE_AREA_W, TABLE_AREA_H,
        page.table_headers, page.table_rows, reveal_cells=reveals,
        cell_colors=cell_colors,
        # 列宽：中间稿 `列宽：` 显式比优先；留白填写表转白底+灰格线，投屏一眼看得出是表
        col_weights=page.col_weights or None, blank_grid=TABLE_GRID_COLOR,
        # v9：表头主题色；字号上限放到 22pt、行少表加高到区域 85%（投屏可读、不留大片空白）
        **({"head_bg": th.main, "body_size": 22, "fill_ratio": 0.85} if th else {}),
    )
    if not reveals:
        return

    # 为每个答案格叠"不透明底色矩形 + 完整答案文本框"，盖住底格占位
    bx, by = geom["x"], geom["y"]
    col_x, col_w = geom["col_x"], geom["col_w"]
    row_y, row_h = geom["row_y"], geom["row_h"]
    body_fs = geom["body_fs"]
    marg_l, marg_t = geom["marg_l"], geom["marg_t"]
    zebra = geom["zebra_bg"]

    groups = []
    for (d, c) in sorted(reveals.keys()):       # 阅读顺序：行优先、列内左→右
        tr = d + 1                               # 数据行 d 对应表行 d+1（表头占第 0 行）
        if tr >= len(row_y) or c >= len(col_x):
            continue
        cx = bx + col_x[c]
        cy = by + row_y[tr]
        cw, ch = col_w[c], row_h[tr]
        bg = "FFFFFF" if geom.get("form") else cell_bg_color(d, zebra)

        rect = add_rect(slide, cx, cy, cw, ch, bg)
        _disable_shape_effects(rect)
        full = str(reveals[(d, c)]["full"]).replace("<br>", "\n")
        text_box = add_textbox(
            slide, cx + marg_l, cy + marg_t, cw - 2 * marg_l, ch - 2 * marg_t,
            full,
            font=FONT_BODY, size=body_fs, color=COLOR_BODY,
            align="left", anchor="middle", line_spacing=1.1,
        )
        groups.append([rect.shape_id, text_box.shape_id])

    # 逐格点击：每格一击（矩形+文字一起淡入）；关动画/系统不支持时全部直出=答案静态全显
    if ctx.get("anim") and groups:
        add_click_reveal(slide, groups)


RENDERERS_READING = {
    "封面": render_cover,
    "环节标题": render_section,
    "引导问题": render_guide,
    "原文齐读": render_quote,
    "要点小结": render_summary,
    "填空表格": render_table,
    "_END": lambda slide, page, ctx: (
        v9.render_end(slide, page, ctx, _v9(ctx)) if _v9(ctx) else render_end(slide, page, ctx)),
}

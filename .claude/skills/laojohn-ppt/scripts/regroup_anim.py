# -*- coding: utf-8 -*-
r"""按教学节拍重建分组工作单（外部 PPT 后处理链第 2 步②「人工校正」的载体）。

inspect_pptx.py 的几何启发式只给得出「组的构成」，实测还会跨条目错位——把上一条的
正文和下一条的序号绑成一组。本脚本让你用**版式模式**重新表达每页的点击节拍，
装饰件（底框/圆点/色条）自动认领，落盘前统一转成 shape_id。

    import sys; sys.path.insert(0, r"<项目根>\.claude\skills\laojohn-ppt\scripts")
    from regroup_anim import run
    PLAN = {
        3:  ("row", 0.5),                     # 自上而下逐条
        4:  ("explicit", [[7,8],[17],[10,11]]),  # 顺序与版面不一致时显式写
        5:  ("grid", 0.8),                    # 先分行、行内再分格（2×2 网格页）
        7:  ("col", None),                    # 自左向右逐栏（并列卡片页）
    }
    run(r"...\某课-全课.pptx", PLAN)           # 改写同名 -anim.json
    # 再跑 animate_pptx.py 注入

explicit 里写的是**形状在 slide.shapes 里的位置索引**（0 起），不是 shape_id——
先用 dump_shapes() 打印对照表再写。落盘时本脚本负责转换。

## 四个踩过的坑（改本文件前必读，每条都付出过返工）

① **工作单 JSON 里的数字是 OOXML shape_id，不是位置索引。** 两套编号数值范围重叠
   （实测 shape_id = 位置索引 + 2），animate_pptx 只校验「该 id 在本页是否存在」，
   绑错了照样放行——表现为标题被当正文藏起来、末尾金句一开机就亮着，19 页全中。

② **装饰的吸附目标只能是文字形状。** 装饰一旦入组就参与后续比距离的话，一串小圆点
   会互相吸引、全聚到第一组（实测：三个圆点＋标题短横线全挤进第 1 次点击）。

③ **判断"这个装饰罩住了谁"必须二维。** 只看纵向重叠的话，左栏一个高底框会跟右栏
   每张卡片都"重叠"，被误判成整表外框而漏掉动画——表现是「框先摆在台上，字才一条条
   往里填」。同理，别拿固定线卡页顶：大底纹的 top 常常压着标题区（实测 2.49 英寸差
   一丝就被误判成页顶元素）。

④ **认领小装饰要用边缘间隙，不是中心距离。** 项目符号圆点紧贴它那条文字的左边缘，
   中心距离会被同行一个高大的标签抢走（实测：圆点被高 3.1 英寸的大标签吸走，
   跟着蓝框一起跳出来）。

阈值一律按页面尺寸取相对值，换成 13.33×7.5 的常规尺寸也能直接用。
"""
import json

from pptx import Presentation
from pptx.util import Emu


class Sheet:
    """一页的几何上下文：把绝对英寸换算成相对页面尺寸的比例。"""

    def __init__(self, prs):
        self.W = Emu(prs.slide_width).inches
        self.H = Emu(prs.slide_height).inches

    @property
    def head_top(self):        # 此线以上＝眉标/标题区
        return self.H * 0.22

    @property
    def foot_top(self):        # col 模式下，此线以下的通栏条改按 row 收尾
        return self.H * 0.70

    def is_backdrop(self, w, h):
        return w > self.W * 0.92 and h > self.H * 0.45

    def is_artwork(self, h):   # 配图/占位框（相对的：窄色条不算）
        return h > self.H * 0.18

    @property
    def foot_gap(self):
        return self.H * 0.04


def geo(sh):
    return (Emu(sh.left).inches, Emu(sh.top).inches,
            Emu(sh.width).inches, Emu(sh.height).inches)


def dump_shapes(pptx, pages):
    """打印位置索引 → shape_id → 几何 → 文字，写 explicit 前对着它挑索引。"""
    prs = Presentation(pptx)
    slides = list(prs.slides)
    for pg in pages:
        print(f"=== P{pg:02d}")
        for j, sh in enumerate(slides[pg - 1].shapes):
            l, t, w, h = [round(v, 1) for v in geo(sh)]
            txt = (sh.text_frame.text.replace("\n", " ")[:40]
                   if sh.has_text_frame and sh.text_frame.text.strip() else "«无字»")
            print(f"  idx{j:2d} → id{sh.shape_id:3d}  L{l:5.1f} T{t:5.1f} W{w:5.1f} H{h:4.1f}  {txt}")


def cluster(vals, gap):
    out = []
    for v in sorted(vals):
        if out and v - out[-1][1] <= gap:
            out[-1][1] = v
        else:
            out.append([v, v])
    return out


def build(shapes, mode, arg, sheet):
    if mode == "explicit":
        return [list(g) for g in arg]

    def rows(items, gap):
        bands = cluster([t for _, _, t, _, _, _ in items], gap)
        return [[it for it in items if lo <= it[2] <= hi] for lo, hi in bands]

    def cols(items):
        gap = sheet.W * 0.15
        bands = cluster([l for _, l, _, _, _, _ in items], gap)
        return [[it for it in items if lo <= it[1] <= hi] for lo, hi in bands]

    if mode == "row":
        return [[it[0] for it in r] for r in rows(shapes, arg)]
    if mode == "grid":
        return [[it[0] for it in c] for r in rows(shapes, arg) for c in cols(r)]
    if mode == "col":
        body = [it for it in shapes if it[2] < sheet.foot_top]
        foot = [it for it in shapes if it[2] >= sheet.foot_top]
        out = [[it[0] for it in c] for c in cols(body)]
        out += [[it[0] for it in r] for r in rows(foot, sheet.foot_gap)]
        return out
    raise ValueError("未知版式模式：%s" % mode)


def attach_decor(slide, groups, sheet):
    """底框/圆点/色条认领到它所衬的那一条；页顶装饰、整页背景、配图保持常驻。"""
    anchors = [list(g) for g in groups]            # 文字成员快照（坑②）
    assigned = {i for g in groups for i in g}
    boxes = {j: geo(sh) for j, sh in enumerate(slide.shapes)}
    body_top = min((boxes[i][1] for i in assigned), default=sheet.head_top)

    def yov(a, b):
        return min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])

    for j, sh in enumerate(slide.shapes):
        if j in assigned:
            continue
        l, t, w, h = boxes[j]
        if sheet.is_backdrop(w, h):
            continue
        if sh.__class__.__name__ == "Picture" and sheet.is_artwork(h):
            continue
        if t + h <= body_top:                      # 整体在正文之上＝标题装饰（坑③）
            continue

        covered = [g_i for g_i, g in enumerate(anchors)                     # 坑③
                   if any(l <= boxes[k][0] + boxes[k][2] / 2 <= l + w and
                          t <= boxes[k][1] + boxes[k][3] / 2 <= t + h for k in g)]
        if covered:
            groups[covered[0]].append(j)           # 跨多组的容器跟它罩住的首条一起出
            continue

        best, score = None, None                   # 罩不住任何文字的小件（坑④）
        for g_i, g in enumerate(anchors):
            for k in g:
                kb = boxes[k]
                gx = max(0.0, kb[0] - (l + w), l - (kb[0] + kb[2]))
                gy = max(0.0, kb[1] - (t + h), t - (kb[1] + kb[3]))
                s = (0 if yov(boxes[j], kb) > 0 else 1, round(gx + gy, 2))
                if score is None or s < score:
                    best, score = g_i, s
        if best is not None:
            groups[best].append(j)
    return groups


def run(pptx, plan, only=None):
    prs = Presentation(pptx)
    sheet = Sheet(prs)
    slides = list(prs.slides)
    jpath = pptx.replace(".pptx", "-anim.json")
    with open(jpath, encoding="utf-8") as f:
        doc = json.load(f)

    for pg, (mode, arg) in plan.items():
        sl = slides[pg - 1]
        body = []
        for j, sh in enumerate(sl.shapes):
            if not sh.has_text_frame or not sh.text_frame.text.strip():
                continue
            l, t, w, h = geo(sh)
            if t < sheet.head_top:
                continue
            body.append((j, l, t, w, h, sh.text_frame.text.replace("\n", " ")))
        groups = [g for g in build(body, mode, arg, sheet) if g]
        groups = attach_decor(sl, groups, sheet)

        sid = [sh.shape_id for sh in sl.shapes]     # 位置索引 → shape_id（坑①）
        doc["slides"][str(pg)]["groups"] = [[sid[i] for i in g] for g in groups]
        doc["slides"][str(pg)]["skip"] = False

        if only is None or pg in only:
            texts = {j: t for j, _, _, _, _, t in body}
            print(f"P{pg:02d} · {len(groups)} 次点击 · {doc['slides'][str(pg)]['title'][:26]}")
            for n, g in enumerate(groups, 1):
                print(f"   {n}. " + " ／ ".join(texts[i][:34] for i in sorted(g) if i in texts))

    with open(jpath, "w", encoding="utf-8") as f:
        json.dump(doc, f, ensure_ascii=False, indent=1)
    live = [v for v in doc["slides"].values() if not v.get("skip")]
    print("-> %s｜有动画页 %d，点击合计 %d"
          % (jpath, len(live), sum(len(v["groups"]) for v in live)))


def audit(pptx):
    """落盘前自检：装饰认领得离谱、某组只有装饰没有文字，都在这儿抓。"""
    prs = Presentation(pptx)
    sheet = Sheet(prs)
    with open(pptx.replace(".pptx", "-anim.json"), encoding="utf-8") as f:
        doc = json.load(f)
    bad = []
    for pg, sl in enumerate(prs.slides, 1):
        v = doc["slides"].get(str(pg))
        if not v or v.get("skip"):
            continue
        box = {sh.shape_id: geo(sh) for sh in sl.shapes}
        has = {sh.shape_id: (sh.has_text_frame and bool(sh.text_frame.text.strip()))
               for sh in sl.shapes}
        for n, g in enumerate(v["groups"], 1):
            texts = [s for s in g if has.get(s)]
            if not texts:
                bad.append("P%02d 第%d组全是装饰、没有文字" % (pg, n))
                continue
            for s in g:
                if has.get(s):
                    continue
                l, t, w, h = box[s]
                gap = min(max(0, box[k][0] - (l + w), l - (box[k][0] + box[k][2])) +
                          max(0, box[k][1] - (t + h), t - (box[k][1] + box[k][3]))
                          for k in texts)
                if gap > sheet.W * 0.06:
                    bad.append("P%02d 第%d组：装饰离本组文字 %.1f 英寸，疑似认领错"
                               % (pg, n, gap))
    print("自检：", "\n  ".join(bad) if bad else "通过")
    return bad

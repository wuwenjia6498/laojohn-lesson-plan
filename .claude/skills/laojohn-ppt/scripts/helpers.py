# -*- coding: utf-8 -*-
"""通用绘制工具：文本框、矩形、占位框、表格。

【单一共享源 · 课型无关原语库】本模块是读书会与写作课两 profile 共用的底层渲染原语，
是有 bug 史的横切层（字体槽顺序、表格自适应等修一处即重烘焙全部）。
铁律：**永不 fork、禁出现 `doc_kind` / 课型分支**——课型差异在 layouts_reading /
layouts_writing 各自的 renderer 与 theme_writing 变体里表达，不下沉到本层。
"""
import os
import math
import re
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

from theme import (
    FONT_TITLE, FONT_BODY, FONT_ASCII,
    COLOR_DASH, COLOR_BG_PLACE, COLOR_MUTED, COLOR_BODY,
    SZ_PLACEHOLDER,
    COLOR_NUMBOX_BG, COLOR_NUMBOX_FG, COLOR_BULLET_SEP,
    COLOR_RED_ACCENT, COLOR_CARD_BG, COLOR_CARD_BORDER,
)


def rgb(hex_str):
    return RGBColor.from_string(hex_str.lstrip("#"))


def add_textbox(slide, x, y, w, h, text, *,
                font=FONT_BODY, size=20, color="404040",
                bold=False, italic=False, align="left", anchor="top",
                line_spacing=1.2, first_line_indent_chars=0):
    """添加一个文本框。

    align: "left" | "center" | "right"
    anchor: "top" | "middle" | "bottom"
    """
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    anchor_map = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
    tf.vertical_anchor = anchor_map.get(anchor, MSO_ANCHOR.TOP)

    align_map = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}

    lines = text.split("\n") if isinstance(text, str) else list(text)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align_map.get(align, PP_ALIGN.LEFT)
        p.line_spacing = line_spacing
        if first_line_indent_chars:
            # 用全角空格模拟首行缩进
            line = "\u3000" * first_line_indent_chars + line
        run = p.add_run()
        run.text = line
        f = run.font
        f.name = font
        f.size = Pt(size)
        f.bold = bold
        f.italic = italic
        f.color.rgb = rgb(color)
        # 显式设置东亚字体
        _set_east_asia_font(run, font)
    return box


def add_rich_textbox(slide, x, y, w, h, segments, *,
                    font=FONT_BODY, size=20, default_color="404040",
                    default_bold=False, default_italic=False,
                    align="left", anchor="top", line_spacing=1.2):
    """富文本框：segments = [(text, {color,bold,italic,size}), ...]。
    text 内 \n 自动换段；style 缺省回退到 default_*。"""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(0); tf.margin_right = Emu(0)
    tf.margin_top = Emu(0);  tf.margin_bottom = Emu(0)
    anchor_map = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
    tf.vertical_anchor = anchor_map.get(anchor, MSO_ANCHOR.TOP)
    align_map = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}

    p = tf.paragraphs[0]
    p.alignment = align_map.get(align, PP_ALIGN.LEFT)
    p.line_spacing = line_spacing

    for text, style in segments:
        for j, part in enumerate(str(text).split("\n")):
            if j > 0:
                p = tf.add_paragraph()
                p.alignment = align_map.get(align, PP_ALIGN.LEFT)
                p.line_spacing = line_spacing
            run = p.add_run()
            run.text = part
            f = run.font
            f.name = font
            f.size = Pt(style.get("size", size))
            f.bold = style.get("bold", default_bold)
            f.italic = style.get("italic", default_italic)
            f.color.rgb = rgb(style.get("color", default_color))
            _set_east_asia_font(run, font)
    return box


def _set_east_asia_font(run, font_name):
    """python-pptx 默认不设 eastAsia 字体，中文会回退到主题字体。手动注入。
    注意：OOXML(CT_TextCharacterProperties) 规定 rPr 子元素须 latin 在前、
    ea 在后。顺序写反会被 PowerPoint 判为非法而忽略 ea，导致中文也套用
    latin 槽的字体(Calibri)，本地打开“所有字体显示 Calibri”。"""
    rPr = run._r.get_or_add_rPr()
    # 先 latin（须排在 ea 之前）：设为 Calibri，避免中文字体应用到英文/数字时显丑
    for latin in rPr.findall(qn("a:latin")):
        rPr.remove(latin)
    latin = etree.SubElement(rPr, qn("a:latin"))
    latin.set("typeface", FONT_ASCII)
    # 再 ea：东亚字体
    for ea in rPr.findall(qn("a:ea")):
        rPr.remove(ea)
    ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", font_name)


def add_rect(slide, x, y, w, h, fill_color, line=False, line_color=None,
             shadow=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill_color)
    if line:
        shape.line.color.rgb = rgb(line_color or fill_color)
        shape.line.width = Pt(0.5)
    else:
        shape.line.fill.background()
    if not shadow:
        # 平色条/分隔线/铭牌默认不要主题阴影（否则细线也拖一道灰影，AI 感重）
        _disable_shape_effects(shape)
    return shape


def add_gradient_rect(slide, x, y, w, h, color_from, color_to, angle_deg=135):
    """添加线性渐变矩形（无边框）。

    angle_deg：渐变方向角度。0=从左到右；90=从上到下；135=从左上到右下。
    OpenXML 用 60000 单位 = 1 度。
    """
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    spPr = shape.fill._xPr
    # 移除现有 fill 元素（必须放在 ln 之前；先清空避免冲突）
    for tag in ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill"):
        for el in spPr.findall(qn(tag)):
            spPr.remove(el)

    # 渐变填充：在 ln 元素之前插入
    ln_elem = spPr.find(qn("a:ln"))
    gradFill = etree.Element(qn("a:gradFill"))
    gradFill.set("flip", "none")
    gradFill.set("rotWithShape", "1")

    gsLst = etree.SubElement(gradFill, qn("a:gsLst"))
    for pos, color in [("0", color_from), ("100000", color_to)]:
        gs = etree.SubElement(gsLst, qn("a:gs"))
        gs.set("pos", pos)
        srgb = etree.SubElement(gs, qn("a:srgbClr"))
        srgb.set("val", color.lstrip("#"))

    lin = etree.SubElement(gradFill, qn("a:lin"))
    lin.set("ang", str(int(angle_deg * 60000)))
    lin.set("scaled", "0")
    etree.SubElement(gradFill, qn("a:tileRect"))

    if ln_elem is not None:
        ln_elem.addprevious(gradFill)
    else:
        spPr.append(gradFill)

    # 去边框
    shape.line.fill.background()
    _disable_shape_effects(shape)
    return shape


def add_dashed_rect(slide, x, y, w, h, fill_color=COLOR_BG_PLACE, line_color=COLOR_DASH):
    """虚线矩形（用于配图占位框）。"""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill_color)
    shape.line.color.rgb = rgb(line_color)
    shape.line.width = Pt(1.25)
    # 设置虚线：操作底层 XML
    spPr = shape.line._get_or_add_ln()
    # 移除已有 prstDash
    for d in spPr.findall(qn("a:prstDash")):
        spPr.remove(d)
    dash = etree.SubElement(spPr, qn("a:prstDash"))
    dash.set("val", "dash")
    return shape


def add_image_placeholder(slide, x, y, w, h, suggestion: str):
    """配图占位框：浅灰底 + 虚线 + 📷 + 多行说明。

    suggestion 形如：苏七块·正骨场景｜书内P17–20｜建议元素：正骨大夫、银元、病患
    """
    add_dashed_rect(slide, x, y, w, h)

    # 占位框内文本
    parts = [p.strip() for p in suggestion.split("｜") if p.strip()]
    lines = ["📷 建议配图"]
    if parts:
        lines.append("")
        lines.extend(parts)

    add_textbox(
        slide, x, y, w, h,
        "\n".join(lines),
        font=FONT_BODY, size=SZ_PLACEHOLDER, color=COLOR_MUTED,
        align="center", anchor="middle", line_spacing=1.5,
    )


def add_image_grid(slide, x, y, w, h, suggestions, *, gap=None):
    """四图网格占位：把 2–4 条配图建议排成 2 列网格，每条一个占位框。

    用于"给你们看 N 幅画面"这类引导问题页（每幅一图，逐格各配一题）。
    suggestions：已过滤掉"无"的配图建议字符串列表（每条仍是三段式）。
    """
    sugs = [s for s in suggestions if s and s.strip() not in
            {"无", "无（页面已满）", "—", "无配图"}]
    n = len(sugs)
    if n == 0:
        return
    if gap is None:
        gap = Emu(int(w * 0.02))
    cols = 1 if n == 1 else 2
    rows = (n + cols - 1) // cols
    cell_w = (w - gap * (cols - 1)) // cols
    cell_h = (h - gap * (rows - 1)) // rows
    for i, sug in enumerate(sugs):
        r, c = divmod(i, cols)
        # 末行不足 cols 个时整行居中
        in_row = cols if (r + 1) * cols <= n else (n - r * cols)
        row_w = cell_w * in_row + gap * (in_row - 1)
        x_off = x + (w - row_w) // 2
        cx = x_off + c * (cell_w + gap)
        cy = y + r * (cell_h + gap)
        add_image_placeholder(slide, cx, cy, cell_w, cell_h, sug)


def add_logo(slide, logo_path, x, y, w, h=None):
    """添加 LOGO。h=None 时按图片原比例自动缩放（推荐，避免拉伸变形）。"""
    if logo_path and os.path.isfile(logo_path):
        if h is None:
            slide.shapes.add_picture(logo_path, x, y, width=w)
        else:
            slide.shapes.add_picture(logo_path, x, y, w, h)


def add_image_or_skip(slide, image_path, x, y, w, h):
    """安全添加图片：文件不存在时静默跳过。"""
    if image_path and os.path.isfile(image_path):
        slide.shapes.add_picture(image_path, x, y, w, h)


def _crop_to_fill(pic, target_w, target_h):
    """把 add_picture 已按原比例放好的图裁成正好铺满 target 框（cover 语义）。

    add_picture 给了 width→高按比例、或给了 width+height→拉伸。这里两边都给再用
    crop 修掉溢出：先按较大缩放比铺满、算出溢出百分比、对称裁掉，杜绝拉伸变形。
    pic.crop_* 取值是"裁掉的比例"(0~1)。图内在尺寸用 pic.image 原始像素算。
    """
    try:
        iw, ih = pic.image.size  # 原始像素
    except Exception:
        return
    if iw <= 0 or ih <= 0:
        return
    tw = float(target_w); th = float(target_h)
    scale = max(tw / iw, th / ih)      # cover：取较大比，保证铺满
    disp_w = iw * scale; disp_h = ih * scale
    crop_x = (disp_w - tw) / disp_w / 2 if disp_w > tw else 0
    crop_y = (disp_h - th) / disp_h / 2 if disp_h > th else 0
    pic.crop_left = crop_x
    pic.crop_right = crop_x
    pic.crop_top = crop_y
    pic.crop_bottom = crop_y


def _round_picture(pic, radius_frac=0.06):
    """把图形状换成圆角矩形（prstGeom = roundRect），radius_frac 控制圆角半径。"""
    try:
        spPr = pic._element.spPr
        for tag in ("a:prstGeom", "a:custGeom"):
            for el in spPr.findall(qn(tag)):
                spPr.remove(el)
        geom = etree.SubElement(spPr, qn("a:prstGeom"))
        geom.set("prst", "roundRect")
        av = etree.SubElement(geom, qn("a:avLst"))
        gd = etree.SubElement(av, qn("a:gd"))
        gd.set("name", "adj")
        gd.set("fmla", f"val {int(radius_frac * 100000)}")
    except Exception:
        pass


def add_image_cover(slide, x, y, w, h, image_path, *,
                    rounded=True, radius_frac=0.05,
                    caption=None, caption_band_color="6F9F5B",
                    caption_fg="FFFFFF", caption_size=15,
                    border_color=None):
    """等比裁切铺满给定框放真图（cover），可选圆角 + 底部半透明说明条。

    - image_path 不存在时：退回浅灰占位块 + 📷（不留白洞、不报错）。
    - caption 非空时：在图底沿压一条 caption_band_color 说明条 + 白字场景名
      （压的是纯色条不是文字浮层，不影响图面可读性；对齐文件4实景观察页样式）。
    - border_color 非空时给图描一圈细边（浅灰卡感）。
    返回 (pic_or_none, caption_shapes)——caption_shapes 供需要时纳入点击分组。
    """
    if not (image_path and os.path.isfile(image_path)):
        # 占位兜底：浅灰圆角块 + 相机字
        shp = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
            x, y, w, h)
        if rounded:
            try:
                shp.adjustments[0] = radius_frac
            except Exception:
                pass
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(COLOR_BG_PLACE)
        shp.line.color.rgb = rgb(COLOR_DASH)
        shp.line.width = Pt(1)
        _disable_shape_effects(shp)
        add_textbox(slide, x, y, w, h, "📷 待贴图" + (f"\n{caption}" if caption else ""),
                    font=FONT_BODY, size=SZ_PLACEHOLDER, color=COLOR_MUTED,
                    align="center", anchor="middle", line_spacing=1.4)
        return None, []

    pic = slide.shapes.add_picture(image_path, x, y, width=w, height=h)
    _crop_to_fill(pic, w, h)
    if rounded:
        _round_picture(pic, radius_frac)
    if border_color:
        pic.line.color.rgb = rgb(border_color)
        pic.line.width = Pt(0.75)

    cap_shapes = []
    if caption:
        band_h = Pt(30)
        band_y = y + h - band_h
        band = add_rect(slide, x, band_y, w, band_h, caption_band_color)
        _set_shape_alpha(band, 78000)        # 半透明说明条（78%），压图底沿
        _disable_shape_effects(band)
        cap = add_textbox(slide, x + Pt(12), band_y, w - Pt(24), band_h, caption,
                          font=FONT_TITLE, size=caption_size, color=caption_fg,
                          bold=True, align="left", anchor="middle")
        cap_shapes = [band.shape_id, cap.shape_id]
    return pic, cap_shapes


def _round_rect_shape(shape, radius_frac=0.05):
    """把一个已存在的矩形 shape 的几何换成圆角矩形。"""
    try:
        shape.adjustments[0] = radius_frac
    except Exception:
        pass


def _set_shape_alpha(shape, alpha=78000):
    """给实心填充加透明度（alpha 单位：0~100000，100000=不透明）。"""
    try:
        srgb = shape.fill.fore_color._xFill.find(qn("a:srgbClr"))
        if srgb is not None:
            a = etree.SubElement(srgb, qn("a:alpha"))
            a.set("val", str(alpha))
    except Exception:
        pass


def _disable_shape_effects(shape):
    """禁用形状默认阴影等效果。

    两处都要清才真正无阴影（历史 bug：只清 spPr 的 effectLst 时，<p:style> 里的
    <a:effectRef idx="2"> 仍从主题拉回阴影，LibreOffice/PowerPoint 都会渲染出来）：
      1) spPr 内注入空 <a:effectLst/> 覆盖局部效果；
      2) 把 <p:style><a:effectRef> 的 idx 归 0（不引用主题效果矩阵）。
    """
    spPr = shape.fill._xPr
    for tag in ("a:effectLst", "a:effectDag"):
        for el in spPr.findall(qn(tag)):
            spPr.remove(el)
    etree.SubElement(spPr, qn("a:effectLst"))
    # 2) 归零 styleRef 的 effectRef（主题阴影的真正来源）
    try:
        sp = shape._element
        style = sp.find(qn("p:style"))
        if style is not None:
            eff_ref = style.find(qn("a:effectRef"))
            if eff_ref is not None:
                eff_ref.set("idx", "0")
    except Exception:
        pass


def add_conclusion_card(slide, x, y, w, h, text, *,
                        bg_color=COLOR_CARD_BG, border_color=COLOR_CARD_BORDER,
                        text_color=COLOR_RED_ACCENT, text_size=22):
    """要点小结首条结论卡：浅奶油圆角矩形 + 红字加粗。"""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    try:
        shape.adjustments[0] = 0.18
    except Exception:
        pass
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(bg_color)
    shape.line.color.rgb = rgb(border_color)
    shape.line.width = Pt(1)
    _disable_shape_effects(shape)

    pad = Pt(20)
    add_textbox(
        slide, x + pad, y, w - 2 * pad, h, text,
        font=FONT_TITLE, size=text_size, color=text_color,
        bold=True, align="left", anchor="middle", line_spacing=1.4,
    )


def add_numbered_bullets(slide, x, y, w, h, bullets, *,
                          text_size=20, columns=1,
                          num_bg=COLOR_NUMBOX_BG, num_fg=COLOR_NUMBOX_FG,
                          text_color=COLOR_BODY, text_font=FONT_BODY,
                          line_spacing=1.45, start_index=1, num_style="box"):
    """要点列表：序号标记 + 文字。

    num_style 决定标记视觉（纯样式参数，非课型分支——调用方按需选）：
      - "box"（默认，读书会要点小结）：实心 num_bg 方块 + num_fg 数字。
      - "outline"（写作课讲解页·有序要点）：白底 + num_bg 细描边 + num_bg 数字。
      - "plain"：无方块，仅 num_bg 数字。
      - "dot"（写作课讲解页·无序要点）：num_bg 实心小圆点，**不显示数字**——
        并列要点不硬套 1/2/3 的次序暗示。
      - "bar"（写作课讲解页·有序要点）：大号 num_bg 数字 + 右侧红色短竖线分隔正文，
        **无方框**——比描边小方框更精致。
    标记去除主题阴影；要点之间不画分隔线。
    columns=1 默认；调用方可传 columns=2 改双列。

    返回：每条要点对应的 shape_id 分组列表（成组供逐条点击动画同时淡入，
    见 add_click_reveal）：含标记形状时 [标记id, (数字id,) 文字id]，plain 时 [数字id, 文字id]。
    """
    n = len(bullets)
    if n == 0:
        return []
    cols = max(1, columns)
    rows_per_col = (n + cols - 1) // cols
    col_w = w // cols
    row_h = h // rows_per_col

    # 方块尺寸：随字号缩放
    box_side = Pt(text_size + 8)
    text_gap = Pt(14)   # 方块到文字的横向间隙
    num_text_color = num_fg if num_style == "box" else num_bg

    groups = []
    for i, b in enumerate(bullets):
        col = i // rows_per_col
        row = i % rows_per_col
        cx = x + col * col_w
        cy = y + row * row_h
        box_top = cy + Pt(2)

        # 标记（去阴影）：box=实心方块 / outline=白底描边方块 / dot=实心小圆点 / plain=不画
        box_shape = None
        num_box = None
        if num_style == "dot":
            dot_d = Pt(int(text_size * 0.42))
            dot_x = cx + (box_side - dot_d) // 2      # 在 box_side 缩进区居中，文字左界与其他模式对齐
            dot_y = box_top + (box_side - dot_d) // 2
            dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, dot_x, dot_y, dot_d, dot_d)
            dot.fill.solid()
            dot.fill.fore_color.rgb = rgb(num_bg)
            dot.line.fill.background()
            _disable_shape_effects(dot)
            box_shape = dot
        elif num_style == "bar":
            # 无框：大号 num_bg 数字 + 右侧红色短竖线分隔正文
            num_box = add_textbox(
                slide, cx, box_top, box_side, Pt(text_size + 8),
                str(i + start_index),
                font=FONT_TITLE, size=text_size + 4, color=num_bg,
                bold=True, align="center", anchor="top",
            )
            bar_h = Pt(int(text_size * 1.4))
            box_shape = add_rect(slide, cx + box_side, box_top, Pt(3), bar_h,
                                 COLOR_RED_ACCENT)
            _disable_shape_effects(box_shape)
        else:
            if num_style == "box":
                box_shape = add_rect(slide, cx, box_top, box_side, box_side, num_bg)
                _disable_shape_effects(box_shape)
            elif num_style == "outline":
                box_shape = add_rect(slide, cx, box_top, box_side, box_side,
                                     "FFFFFF", line=True, line_color=num_bg)
                _disable_shape_effects(box_shape)
            # 数字
            num_box = add_textbox(
                slide, cx, box_top, box_side, box_side,
                str(i + start_index),
                font=FONT_TITLE, size=text_size - 2, color=num_text_color,
                bold=True, align="center", anchor="middle",
            )
        # 内容文字
        text_x = cx + box_side + text_gap
        text_w = col_w - box_side - text_gap - Pt(8)
        content_box = add_textbox(
            slide, text_x, cy, text_w, row_h,
            b,
            font=text_font, size=text_size, color=text_color,
            align="left", anchor="top", line_spacing=line_spacing,
        )
        grp = [content_box.shape_id]
        if num_box is not None:
            grp.insert(0, num_box.shape_id)
        if box_shape is not None:
            grp.insert(0, box_shape.shape_id)
        groups.append(grp)
    return groups


_QA_CIRCLES = ["①", "②", "③", "④", "⑤", "⑥", "⑦", "⑧", "⑨"]


def add_qa_textbox(slide, x, y, w, h, bullets, answers, *,
                   text_size=20, body_color=COLOR_BODY,
                   answer_color=COLOR_RED_ACCENT, font=FONT_BODY,
                   line_spacing=1.35):
    """问题(圈号·正文色) + 参考答案(红字) 交错渲染为单个文本框。

    bullets：问题列表；answers：与之平行对齐的参考答案列表（""=该题无答案）。
    每条问题占一段（①②③…），其下若有答案再占一段（缩进、红字）。

    返回 (box, reveal_groups)：
        reveal_groups —— 按揭示顺序排列的逐段分组 [[(spid, pidx)], ...]，
        问题段与答案段各自成组（各一次点击）。调用方据此调 add_click_reveal
        做"先出问题、点击后出红色答案"的逐段淡入；动画关闭时全部静态显示
        （问题正文色、答案红色），不影响线上系统。
    """
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(0); tf.margin_right = Emu(0)
    tf.margin_top = Emu(0);  tf.margin_bottom = Emu(0)
    tf.vertical_anchor = MSO_ANCHOR.TOP

    reveal_order = []   # 段落 index，按出现(=揭示)顺序
    pidx = 0
    first = True

    def _new_para(space_before_pt=0):
        nonlocal first
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = line_spacing
        if space_before_pt:
            p.space_before = Pt(space_before_pt)
        return p

    def _run(p, text, color, size):
        run = p.add_run()
        run.text = text
        f = run.font
        f.name = font
        f.size = Pt(size)
        f.color.rgb = rgb(color)
        _set_east_asia_font(run, font)

    for i, q in enumerate(bullets):
        marker = _QA_CIRCLES[i] if i < len(_QA_CIRCLES) else f"{i + 1}."
        p = _new_para(space_before_pt=0 if i == 0 else 10)
        _run(p, f"{marker}  {q}", body_color, text_size)
        reveal_order.append(pidx); pidx += 1

        ans = answers[i] if i < len(answers) else ""
        if ans:
            pa = _new_para(space_before_pt=2)
            # 缩进两个全角空格，与问题正文错开；红字
            _run(pa, "　　" + ans, answer_color, text_size)
            reveal_order.append(pidx); pidx += 1

    groups = [[(box.shape_id, pi)] for pi in reveal_order]
    return box, groups


# ---------- 点击逐条呈现动画（注入 OpenXML 时间树）----------
def _sptgt_xml(target):
    """构造 <p:spTgt>：int=整形动画；(spid, 段落号)=该形状第 N 段落动画。"""
    if isinstance(target, tuple):
        spid, pidx = target
        return (f'<p:spTgt spid="{spid}">'
                f'<p:txEl><p:pRg st="{pidx}" end="{pidx}"/></p:txEl>'
                f'</p:spTgt>')
    return f'<p:spTgt spid="{target}"/>'


def add_click_reveal(slide, groups, *, dur=500, effect="fade"):
    """为 slide 注入"逐组点击淡入"动画时间树（PowerPoint 原生 mainSeq）。

    groups：列表，每元素是一组 target（同一次鼠标点击一起淡入）。
        target = int(shape_id)            —— 整个形状淡入
              或 (shape_id, paragraph_idx) —— 该形状第 paragraph_idx 段落淡入（按段落构建）
    每次点击推进一组；未点击的组初始隐藏（不支持动画的渲染器一般退化为全部直出）。
    effect："fade"＝淡入（缺省，存量产物都是它）；"appear"＝出现（PowerPoint 预设 1，
        只切可见性、无过渡，dur 不起作用）。写作课配图版用 appear（2026-09-30 用户定）。
    """
    if effect not in ("fade", "appear"):
        raise ValueError(f"未知动画效果：{effect}")
    groups = [g for g in groups if g]
    if not groups:
        return
    P = "http://schemas.openxmlformats.org/presentationml/2006/main"
    A = "http://schemas.openxmlformats.org/drawingml/2006/main"

    cid = [3]   # id 1=tmRoot、2=mainSeq；其余节点从 3 起，保证全树唯一
    def nid():
        v = cid[0]; cid[0] += 1; return v

    steps_xml = []
    para_spids = []
    for group in groups:
        eff_xml = []
        for k, target in enumerate(group):
            node_type = "clickEffect" if k == 0 else "withEffect"
            tgt = _sptgt_xml(target)
            if isinstance(target, tuple) and target[0] not in para_spids:
                para_spids.append(target[0])
            set_id, anim_id, eff_id = nid(), nid(), nid()
            fade_xml = (
                f'<p:animEffect transition="in" filter="fade"><p:cBhvr>'
                f'<p:cTn id="{anim_id}" dur="{dur}"/><p:tgtEl>{tgt}</p:tgtEl>'
                f'</p:cBhvr></p:animEffect>') if effect == "fade" else ""
            preset = "10" if effect == "fade" else "1"
            eff_xml.append(
                f'<p:par><p:cTn id="{eff_id}" presetID="{preset}" presetClass="entr"'
                f' presetSubtype="0" fill="hold" grpId="0" nodeType="{node_type}">'
                f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
                f'<p:set><p:cBhvr><p:cTn id="{set_id}" dur="1" fill="hold">'
                f'<p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
                f'<p:tgtEl>{tgt}</p:tgtEl>'
                f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>'
                f'</p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
                f'{fade_xml}</p:childTnLst></p:cTn></p:par>'
            )
        outer_id, inner_id = nid(), nid()
        steps_xml.append(
            f'<p:par><p:cTn id="{outer_id}" fill="hold">'
            f'<p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
            f'<p:par><p:cTn id="{inner_id}" fill="hold">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst>'
            f'<p:childTnLst>{"".join(eff_xml)}</p:childTnLst></p:cTn></p:par>'
            f'</p:childTnLst></p:cTn></p:par>'
        )

    bld_xml = ""
    if para_spids:
        bld_items = "".join(
            f'<p:bldP spid="{spid}" grpId="0" build="byParagraph"/>'
            for spid in para_spids)
        bld_xml = f'<p:bldLst>{bld_items}</p:bldLst>'

    timing_xml = (
        f'<p:timing xmlns:p="{P}" xmlns:a="{A}"><p:tnLst><p:par>'
        f'<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">'
        f'<p:childTnLst><p:seq concurrent="1" nextAc="seek">'
        f'<p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
        f'{"".join(steps_xml)}</p:childTnLst></p:cTn>'
        f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
        f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
        f'</p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>{bld_xml}</p:timing>'
    )

    timing = etree.fromstring(timing_xml.encode("utf-8"))
    sld = slide.element
    ext = sld.find(qn("p:extLst"))
    if ext is not None:
        ext.addprevious(timing)
    else:
        sld.append(timing)


def _remove_cell_borders(cell):
    """移除单元格四边边框（设为 noFill）。"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for side in ("lnL", "lnR", "lnT", "lnB"):
        # 移除已有
        for ln in tcPr.findall(qn(f"a:{side}")):
            tcPr.remove(ln)
        # 添加 noFill
        ln = etree.SubElement(tcPr, qn(f"a:{side}"))
        ln.set("w", "0")
        etree.SubElement(ln, qn("a:noFill"))


# 单元格分类上色标注：`[码:片段]`（与示范文分句上色同语法）。传 add_table(cell_colors=)
# 时，标注片段渲染成该码对应的颜色；几何计算用剥掉标注后的纯文本。
_CELL_ANNOT_RE = re.compile(r"\[([^:：\]]+)[:：]([^\]]+)\]")


def _strip_cell_annot(text):
    """剥掉单元格里的 `[码:片段]` 标注、只留片段文字（供列宽/字号几何估算）。"""
    return _CELL_ANNOT_RE.sub(lambda m: m.group(2), str(text))


def _cell_segments(text):
    """单元格文本 → 行段列表（`<br>` 与 `\n` 均视为换行）。"""
    return str(text).replace("<br>", "\n").split("\n")


def _cell_maxlen(text):
    """单元格最长一行的字符数（按换行拆分后取最长段）。"""
    return max((len(s) for s in _cell_segments(text)), default=0)


def _content_col_widths(headers, rows, total_w):
    """按各列"最长一行内容"自适应列宽：短列窄、长列宽。

    用 len^0.7 软化极端差异，避免某一长列吃掉几乎全部宽度；每列设字符数下限 3。
    返回 EMU 整数列表，合计 == total_w。
    """
    ncols = len(headers)
    maxlen = [max(3, _cell_maxlen(headers[c])) for c in range(ncols)]
    for row in rows:
        for c in range(ncols):
            if c < len(row):
                maxlen[c] = max(maxlen[c], _cell_maxlen(row[c]))
    weights = [m ** 0.7 for m in maxlen]
    tot = sum(weights) or 1.0
    widths = [int(int(total_w) * wt / tot) for wt in weights]
    widths[-1] = int(total_w) - sum(widths[:-1])   # 余数归末列，保证合计精确
    return widths


def _fit_table_font(headers, rows, widths_emu, area_h_emu, marg_pt,
                    size_hi=16, size_lo=10):
    """自适应字号：从 size_hi 往下试，找到"全表估算高度 ≤ 区域高度"的最大字号。

    估算逐行：每个单元格按列宽换算每行可容字符数（中文按 1 em≈字号宽），
    对每个换行段做 ceil(段长/每行字符数) 累加得该格行数，取本行各格最大行数。
    返回 (字号, [每行高度EMU]).
    """
    widths_pt = [wd / 914400 * 72 for wd in widths_emu]
    area_h_pt = area_h_emu / 914400 * 72
    all_rows = [headers] + list(rows)
    ncols = len(headers)

    best = None
    for fs in range(size_hi, size_lo - 1, -1):
        line_h = fs * 1.32           # 行距（含 1.1 行间）
        row_heights_pt = []
        total = 0.0
        for row in all_rows:
            row_lines = 1
            for c in range(ncols):
                txt = row[c] if c < len(row) else ""
                avail_pt = max(8.0, widths_pt[c] - 2 * marg_pt)
                cpl = max(1, int(avail_pt / (fs * 1.02)))   # 每行可容字符数（中文为主）
                cell_lines = 0
                for seg in _cell_segments(txt):
                    cell_lines += max(1, math.ceil(len(seg) / cpl))
                row_lines = max(row_lines, cell_lines)
            rh = row_lines * line_h + 2 * marg_pt
            row_heights_pt.append(rh)
            total += rh
        if total <= area_h_pt:
            best = (fs, row_heights_pt)
            break
    if best is None:
        # 连最小字号也放不下：用最小字号，并把行高等比压进区域（宁可挤，不可出界）
        fs = size_lo
        line_h = fs * 1.32
        row_heights_pt = []
        total = 0.0
        for row in all_rows:
            row_lines = 1
            for c in range(ncols):
                txt = row[c] if c < len(row) else ""
                avail_pt = max(8.0, widths_pt[c] - 2 * marg_pt)
                cpl = max(1, int(avail_pt / (fs * 1.02)))
                cell_lines = sum(max(1, math.ceil(len(s) / cpl)) for s in _cell_segments(txt))
                row_lines = max(row_lines, cell_lines)
            rh = row_lines * line_h + 2 * marg_pt
            row_heights_pt.append(rh); total += rh
        scale = area_h_pt / total if total else 1.0
        row_heights_pt = [rh * scale for rh in row_heights_pt]
        best = (fs, row_heights_pt)

    fs, row_heights_pt = best
    row_heights_emu = [Emu(int(rh / 72 * 914400)) for rh in row_heights_pt]
    return fs, row_heights_emu


def cell_bg_color(data_row_idx, zebra_bg="F5F5F5"):
    """数据行底色：与 add_table 斑马纹一致（d=0 白、d=1 浅灰、d=2 白……）。

    add_table 内表头=第 0 行，数据行从第 1 行起（`r % 2 == 1` 为白）；
    数据行 0 基索引 d 对应表行 r=d+1，故 d 偶=白、d 奇=斑马色。
    """
    return "FFFFFF" if data_row_idx % 2 == 0 else zebra_bg


def add_table(slide, x, y, w, h, headers, rows,
              head_bg="44546A", head_fg="FFFFFF",
              head_size=None, body_size=None,
              zebra_bg="F5F5F5", reveal_cells=None, cell_colors=None):
    """添加表格：无边框 + 表头深色 + 斑马纹，并**自适应塞进给定区域**。

    - `<br>` / `\n` 渲染为单元格内真实换行（多段落），不再印出字面量；
    - 列宽按内容长短分配（短列窄、长列宽）；
    - 字号按行数/内容自动下调（投屏可读下限 10pt），行高均分以不超出区域 h；
    - head_size/body_size 传 None 时自动定档；显式传入则作为上限基准。
    样式：表头深灰蓝粗体白字居中；正文行白/浅灰交替、左对齐、垂直居中；无可见边框。

    reveal_cells：`{(数据行0基, 列0基): {"blank": 占位版, "full": 完整版}}`。
        - 几何/行高仍用传入的 `rows`（完整答案版）算 → 叠层放得下；
        - 但底表里这些格只渲 `blank`（占位）；答案叠层由调用方（render_table）按
          返回的 `geom` 定位 + 逐格点击。

    cell_colors：`{码: 颜色}`。给定时，单元格里 `[码:片段]` 标注的片段渲染成对应色并加粗
        （与示范文分句上色同语法、同一套码 → 全课"颜色=手法/感官"一致）；几何用剥标注纯文本算。
        不给（默认 None）时表格文字原样渲染——纯参数化加法，读书会调用不受影响。

    返回：`(table, geom)`，geom = {x, y, col_x[相对左偏移], col_w, row_y[相对顶偏移],
        row_h, body_fs, marg_l, marg_t, zebra_bg}，供调用方放置答案叠层。
    """
    reveal_cells = reveal_cells or {}
    n_cols = len(headers)
    n_rows = len(rows) + 1
    dense = n_rows >= 7

    # 边距（密集表收紧）
    marg_l = Emu(45720) if dense else Emu(82296)   # 0.05" / 0.09"
    marg_t = Emu(13716) if dense else Emu(27432)   # 0.015" / 0.03"
    marg_pt = (marg_l / 914400 * 72)

    # 几何估算用文本：有分类上色标注时剥掉 `[码:]` 只留片段，避免列宽/字号被标记撑偏
    if cell_colors:
        calc_headers = [_strip_cell_annot(hh) for hh in headers]
        calc_rows = [[_strip_cell_annot(c) for c in row] for row in rows]
    else:
        calc_headers, calc_rows = headers, rows

    # 列宽（内容自适应，按完整答案版 rows）
    widths = _content_col_widths(calc_headers, calc_rows, w)

    # 自适应字号 + 逐行行高（确保合计 ≤ 区域高度 h，按完整答案版 rows）
    size_hi = body_size if body_size else 16
    auto_fs, row_heights = _fit_table_font(calc_headers, calc_rows, widths, h, marg_pt, size_hi=size_hi)
    body_fs = auto_fs
    head_fs = head_size if head_size else min(auto_fs + 2, 18)

    shape = slide.shapes.add_table(n_rows, n_cols, x, y, w, h)
    table = shape.table

    for c in range(n_cols):
        table.columns[c].width = widths[c]
    for r in range(n_rows):
        table.rows[r].height = row_heights[r]

    def style_cell(cell, text, *, fill_color, font_color, size, bold, align):
        _remove_cell_borders(cell)
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(fill_color)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.text = ""
        tf = cell.text_frame
        tf.word_wrap = True
        tf.margin_left = marg_l
        tf.margin_right = marg_l
        tf.margin_top = marg_t
        tf.margin_bottom = marg_t
        align_map = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER}

        def _mk_run(p, txt, color, run_bold):
            run = p.add_run()
            run.text = txt
            font_name = FONT_TITLE if run_bold else FONT_BODY
            run.font.name = font_name
            run.font.size = Pt(size)
            run.font.bold = run_bold
            run.font.color.rgb = rgb(color)
            _set_east_asia_font(run, font_name)

        for j, seg in enumerate(_cell_segments(text)):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.alignment = align_map.get(align, PP_ALIGN.LEFT)
            p.line_spacing = 1.1
            if cell_colors and _CELL_ANNOT_RE.search(seg):
                # 分类上色：`[码:片段]` 片段用对应色+加粗，其余用默认色
                pos = 0
                for m in _CELL_ANNOT_RE.finditer(seg):
                    if m.start() > pos:
                        _mk_run(p, seg[pos:m.start()], font_color, bold)
                    code, txt = m.group(1).strip(), m.group(2)
                    if code in cell_colors:
                        _mk_run(p, txt, cell_colors[code], True)
                    else:
                        _mk_run(p, txt, font_color, bold)
                    pos = m.end()
                if pos < len(seg):
                    _mk_run(p, seg[pos:], font_color, bold)
            else:
                _mk_run(p, seg, font_color, bold)

    # 表头
    for c, header in enumerate(headers):
        style_cell(
            table.cell(0, c), header,
            fill_color=head_bg, font_color=head_fg,
            size=head_fs, bold=True, align="center",
        )

    # 数据行（斑马纹：奇数行白，偶数行浅灰）
    for r, row in enumerate(rows, start=1):
        bg = "FFFFFF" if r % 2 == 1 else zebra_bg
        d = r - 1                       # 数据行 0 基索引
        for c in range(n_cols):
            text = row[c] if c < len(row) else ""
            # 答案格底表只渲占位（blank），完整答案由调用方做点击叠层
            rv = reveal_cells.get((d, c))
            if rv is not None:
                text = rv["blank"]
            style_cell(
                table.cell(r, c), text,
                fill_color=bg, font_color=COLOR_BODY,
                size=body_fs, bold=False, align="left",
            )

    # 几何：相对偏移（调用方加 x/y 得绝对坐标），用于放置答案叠层
    col_x, acc = [], 0
    for wd in widths:
        col_x.append(acc); acc += int(wd)
    row_y, acc = [], 0
    for rh in row_heights:
        row_y.append(acc); acc += int(rh)
    geom = {
        "x": x, "y": y,
        "col_x": col_x, "col_w": [int(wd) for wd in widths],
        "row_y": row_y, "row_h": [int(rh) for rh in row_heights],
        "body_fs": body_fs, "marg_l": marg_l, "marg_t": marg_t,
        "zebra_bg": zebra_bg,
    }
    return table, geom


def set_ascii_font(textbox, font_name):
    """覆盖 textbox 内所有 run 的 latin 字体。
    用于纯数字/英文场景换装饰字体（如章节序号用 Bahnschrift）。"""
    for p in textbox.text_frame.paragraphs:
        for run in p.runs:
            rPr = run._r.get_or_add_rPr()
            for latin in rPr.findall(qn("a:latin")):
                rPr.remove(latin)
            latin = rPr.makeelement(qn("a:latin"), {})
            latin.set("typeface", font_name)
            # OOXML 要求 latin 在 ea 之前；若已有 ea，插到它前面，否则追加
            ea = rPr.find(qn("a:ea"))
            if ea is not None:
                ea.addprevious(latin)
            else:
                rPr.append(latin)

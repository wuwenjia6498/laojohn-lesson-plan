# -*- coding: utf-8 -*-
"""外部 pptx 勘查器：导出形状清单 + 建议的「逐条点击」分组工作单。

用途：写作课 PPT 改由外部平台生成后，本仓只做动画后处理。本脚本是后处理的第一段——
只读勘查，产出一份分组工作单 JSON，供人工校正后交给 animate_pptx.py 注入动画。

为什么拆两段：**分组是教学判断**（哪几样东西该一起出现、按什么顺序出现），
几何启发式只能给出「组的构成」，给不出教学顺序（实测样本里，页面最底部那句结论
是第 3 次点击、而不是最后一次）。所以脚本只建议，顺序与取舍由人在 JSON 里定。

    PYTHONUTF8=1 python inspect_pptx.py <foo.pptx> [-o foo-anim.json]

建议分组的启发式：
  1. 排除固定骨架——整页背景、页顶 25% 以内的眉标/标题/装饰/LOGO；
  2. 余下形状按 shape_id 升序扫描：**无文本形状**作簇头（卡片底框），向后吸收
     id 相邻、且中心落在簇头矩形内的形状（卡内图标与文字）；
  3. 未被吸收的有文本形状自成一组；
  4. 组按阅读顺序排（先上后下、同行先左后右）。

自动标 skip 的页：可分组数 <= 1，或整页只有一张表格（表格整体淡入无意义）。
"""
import argparse
import json
import os
import re
import sys
import zipfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx import Presentation                                        # noqa: E402

from plan_link import PlanNotFound, digest, locate, project_root, rel  # noqa: E402

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

# 页顶这个比例以内的形状视为眉标/标题/装饰，不参与动画
HEAD_ZONE = 0.25
# 形状宽高都超过这个比例即视为整页背景
FULL_BG = 0.95
# 判「是否在卡片内」时给簇头矩形的外扩量（占页宽/页高的比例）。
# 实测样本里卡内小图标是骑在底框边线上的，中心点严格落在框内会判失败、吸收随即中断。
PAD = 0.02
# 吸收时允许的连续失配次数（PptxGenJS 偶尔在卡内穿插非本卡形状）
MISS_TOLERANCE = 1
# 「同行相邻」判定允许的水平间距（占页宽比例）
GAP = 0.03
# 「同行相邻」只对不宽于此的簇头生效——超过就是卡片而非项目符号
BULLET_MAX_W = 0.04


def _rect(shape, W, H):
    """归一化矩形 (x, y, w, h)，取值 0~1；缺字段的形状返回 None。"""
    if shape.left is None or shape.top is None:
        return None
    return (shape.left / W, shape.top / H, (shape.width or 0) / W, (shape.height or 0) / H)


def _text(shape):
    if not shape.has_text_frame:
        return ""
    return shape.text_frame.text.strip()


def _center_inside(inner, outer, pad=PAD):
    """inner 的中心点是否落在 outer 矩形（外扩 pad 后）内。"""
    cx = inner[0] + inner[2] / 2.0
    cy = inner[1] + inner[3] / 2.0
    return (outer[0] - pad <= cx <= outer[0] + outer[2] + pad) and \
           (outer[1] - pad <= cy <= outer[1] + outer[3] + pad)


def _same_row(a, b, gap=GAP):
    """a、b 是否属于同一行——垂直区间相交，且水平间距在 gap 以内。

    对应第二种簇头形态：文字左侧的小圆点/项目符号（它包不住文字，只和文字同行相邻）。
    **仅对窄簇头生效**：并排的大卡片彼此也是「同行相邻」，放开会把整排卡片并成一次点击。
    """
    if a[2] > BULLET_MAX_W:
        return False
    if min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1]) <= 0:
        return False
    return max(b[0] - (a[0] + a[2]), a[0] - (b[0] + b[2])) <= gap


def existing_clicks(path):
    """逐页读出已有动画的点击次数（slide 顺序 -> 次数）。"""
    z = zipfile.ZipFile(path)
    out = {}
    for name in z.namelist():
        m = re.match(r"ppt/slides/slide(\d+)\.xml$", name)
        if m:
            xml = z.read(name).decode("utf-8", "ignore")
            out[int(m.group(1))] = xml.count('nodeType="clickEffect"')
    z.close()
    return out


def scan_slide(slide, W, H):
    """返回 (建议分组, 骨架形状 id 列表, 表格数)。"""
    cands = []      # 参与动画的候选：(shape_id, rect, text)
    skeleton = []   # 被排除的骨架形状
    tables = 0

    for sh in slide.shapes:
        if getattr(sh, "has_table", False) and sh.has_table:
            tables += 1
        r = _rect(sh, W, H)
        if r is None:
            skeleton.append(sh.shape_id)
            continue
        if r[2] >= FULL_BG and r[3] >= FULL_BG:      # 整页背景
            skeleton.append(sh.shape_id)
            continue
        if r[1] < HEAD_ZONE:                          # 眉标/标题/装饰/LOGO
            skeleton.append(sh.shape_id)
            continue
        cands.append((sh.shape_id, r, _text(sh)))

    cands.sort(key=lambda c: c[0])
    used = set()
    groups = []
    for i, (sid, rect, txt) in enumerate(cands):
        if sid in used:
            continue
        if txt:
            groups.append(([sid], rect))              # 有文本且没被卡片吸收 -> 独立一组
            used.add(sid)
            continue
        members = [sid]                               # 无文本 -> 卡片底框，向后吸收
        used.add(sid)
        miss = 0
        for sid2, rect2, _t2 in cands[i + 1:]:
            if sid2 in used:
                continue
            if _center_inside(rect2, rect) or _same_row(rect, rect2):
                members.append(sid2)
                used.add(sid2)
                miss = 0
            else:
                miss += 1                             # 只吸收 id 相邻的形状，容忍少量穿插
                if miss > MISS_TOLERANCE:
                    break
        if len(members) == 1:
            # 吸收不到任何东西的空形状 = 孤立装饰或整幅配图，单独淡入没有教学意义
            continue
        groups.append((members, rect))

    # 阅读顺序：先上后下，同一横带内先左后右（10% 高度算同一带）
    groups.sort(key=lambda g: (round(g[1][1] / 0.10), g[1][0]))
    return [g[0] for g in groups], skeleton, tables


def main():
    ap = argparse.ArgumentParser(description="外部 pptx 勘查 + 生成点击分组工作单")
    ap.add_argument("pptx", help="待勘查的 .pptx")
    ap.add_argument("-o", "--out", default=None, help="工作单输出路径（默认 <同名>-anim.json）")
    ap.add_argument("--quiet", action="store_true", help="不打印人读版概览")
    ap.add_argument("--no-lesson-plan", action="store_true",
                    help="该课确实没有详案时显式放行（分组顺序将无据可依，交付时须讲明）")
    ap.add_argument("--overwrite", action="store_true",
                    help="覆盖已经过审查/人工校正的工作单（默认拒绝，免得冲掉排好的分组）")
    args = ap.parse_args()

    out = args.out or os.path.splitext(args.pptx)[0] + "-anim.json"
    if os.path.isfile(out) and not args.overwrite:
        try:
            with open(out, encoding="utf-8") as f:
                old = json.load(f)
        except Exception:
            old = {}
        if old.get("audit") or old.get("lesson_plan"):
            sys.exit("%s 已经过审查/人工校正——重跑会把排好的分组冲掉（一次要重排几十页）。\n"
                     "  只想重看概览：加 --quiet 跑 audit_against_plan.py，或直接读该 json。\n"
                     "  确实要重来：加 --overwrite。" % out)

    # 详案绑定：分组顺序以详案为准，没有详案根本排不出顺序（后处理链第 2 步）
    md = None
    if not args.no_lesson_plan:
        try:
            md = locate(args.pptx)
        except PlanNotFound as e:
            sys.exit(str(e))

    prs = Presentation(args.pptx)
    W, H = prs.slide_width, prs.slide_height
    have = existing_clicks(args.pptx)

    slides = {}
    rows = []
    for idx, slide in enumerate(prs.slides, 1):
        groups, skeleton, tables = scan_slide(slide, W, H)
        texts = [_text(sh) for sh in slide.shapes if _text(sh)]
        kicker = texts[0][:10] if texts else ""
        title = texts[1][:34] if len(texts) > 1 else ""

        skip, note = False, ""
        if idx == 1:
            skip, note = True, "封面页"
        elif idx == len(prs.slides._sldIdLst):
            skip, note = True, "末页（收束/END）"
        elif re.match(r"^\d+$", kicker.strip()) or "END" in kicker.upper():
            skip, note = True, "课时分隔页"
        elif len(groups) <= 1:
            skip, note = True, "可分组数不足，整页直出"
        elif tables and len(groups) <= tables:
            skip, note = True, "表格页，整表淡入无意义"

        entry = {"skip": skip, "groups": groups, "kicker": kicker, "title": title}
        if note:
            entry["note"] = note
        if have.get(idx):
            entry["existing_clicks"] = have[idx]
        slides[str(idx)] = entry
        rows.append((idx, kicker, title, len(groups), have.get(idx, 0), skip, tables))

    root = project_root(args.pptx)
    doc = {"file": os.path.basename(args.pptx),
           "slide_count": len(prs.slides._sldIdLst),
           "dur": 500,
           "lesson_plan": rel(root, md) if md else None,
           "slides": slides}
    if md:
        d = digest(md)
        doc["plan_digest"] = {"md5": d["md5"],
                              "lessons": [{"title": L["title"],
                                           "steps": [{"title": s["title"],
                                                      "minutes": s["minutes"]}
                                                     for s in L["steps"]]}
                                          for L in d["lessons"] if L.get("is_lesson")]}
    with open(out, "w", encoding="utf-8") as f:
        json.dump(doc, f, ensure_ascii=False, indent=1)

    if not args.quiet:
        print("页码 | 建议点击 | 现有点击 | 表 | 眉标 | 标题")
        print("-" * 78)
        for idx, kicker, title, n, cur, skip, tables in rows:
            flag = "skip" if skip else "%4d" % n
            print("P%02d  | %8s | %8d | %2d | %-10s | %s" % (idx, flag, cur, tables, kicker, title))
        total = sum(r[3] for r in rows if not r[5])
        print("-" * 78)
        print("建议点击合计 %d（现有 %d）" % (total, sum(have.values())))

        if md:      # 排 PLAN 时必然看见——分组顺序按这张表跟着详案师话走
            print("\n详案：%s" % rel(root, md))
            for L in digest(md)["lessons"]:
                if not L.get("is_lesson"):
                    continue
                print("  %s" % L["title"])
                for s in L["steps"]:
                    print("     %-42s %s 分钟" % (s["title"][:42], s["minutes"]))
        else:
            print("\n! 未绑定详案（--no-lesson-plan）：分组顺序无据可依，交付时须向用户讲明。")
    print("工作单 -> %s" % out)
    if md:
        print("下一步：跑 audit_against_plan.py 出机检报告（不跑的话 animate_pptx.py 会拒绝注入）")


if __name__ == "__main__":
    main()

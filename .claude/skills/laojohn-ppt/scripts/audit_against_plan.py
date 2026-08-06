# -*- coding: utf-8 -*-
r"""外部 PPT 后处理链第 2 步的机检部分：拿详案审 pptx，结果回写工作单。

    PYTHONUTF8=1 python audit_against_plan.py <课次目录>\<题目>-全课.pptx

**这道闸门守的是「审查做过」，不是「审查全绿」**：查出的问题多半要另行改详案或补页，
不该卡住投屏件产出。所以本脚本只把结论写进工作单的 `audit` 字段，`animate_pptx.py`
校验该字段存在、且其中记的详案 md5 与当前详案一致（详案改了必须重审），issues 非空
照样放行。

**机器只做机器判得准的五类**，教学顺序、环节该不该合并、判语该不该延后，仍归人：
  1 时长加总    课时各环节标称是否等于 45；环节内方括号子环节加总是否超环节标称
  2 重复环节    同一课时里 `[XXX（约 N 分钟）` 前缀重复出现
  3 屏幕指令    详案 `[投屏/发放/出示…]` 条数 ↔ pptx 正文页数，并排列出供人对
  4 逐字比对    详案 `>` 引用块各段、markdown 表格各单元格，能否在 pptx 全文里找到
  5 引号        pptx 侧 ASCII 直引号（须读 run.text，外部件存成 &quot; 实体、数 XML 会误报 0）
                与详案侧弯引号配对

第 1、2 类是《小小“动物园”》那次「第 2 课时环节三整段重复写了两遍」的抓手——两版措辞
不同，逐行比对抓不到，但子环节加总 38 分钟对标称 20 分钟、环节标注前缀各出现两次，
两条都硬。
"""
import argparse
import io
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx import Presentation                                  # noqa: E402
from pptx.util import Emu                                      # noqa: E402

from plan_link import PlanNotFound, digest, locate, project_root, rel   # noqa: E402

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

LESSON_MINUTES = 45
# 逐字比对时忽略的排版差异：空白、软换行、强调星号
NOISE = re.compile(r"[\s*_`]+")
# markdown 表格分隔行 |---|---|
TABLE_SEP = re.compile(r"^\|[\s:|-]+\|$")
# 单元格太短就没有比对价值（“妈妈”“开头”这类到处都是）
MIN_CELL = 6


def norm(s):
    return NOISE.sub("", s or "").replace("“", "\"").replace("”", "\"") \
                                .replace("‘", "'").replace("’", "'")


def slide_texts(prs):
    """每页的文本块列表（按纵向位置排序）+ 全 pptx 归一化后的大字符串。"""
    pages = []
    for sl in prs.slides:
        blocks = sorted([sh for sh in sl.shapes
                         if sh.has_text_frame and sh.text_frame.text.strip()],
                        key=lambda s: Emu(s.top).inches)
        pages.append([b.text_frame.text.strip() for b in blocks])
    return pages, norm("\n".join("\n".join(p) for p in pages))


def check_minutes(d):
    out = []
    for L in d["lessons"]:
        if not L.get("is_lesson"):
            continue
        total = sum(s["minutes"] or 0 for s in L["steps"])
        if total != LESSON_MINUTES:
            out.append("【时长】%s 各环节标称加总 %d 分钟，不等于 %d"
                       % (L["title"], total, LESSON_MINUTES))
        for s in L["steps"]:
            sub = sum(x["minutes"] for x in s["subs"])
            if s["subs"] and s["minutes"] and sub > s["minutes"]:
                out.append("【时长】%s → %s：方括号子环节加总 %d 分钟，超过环节标称 %d "
                           "（子环节可能重复写了两遍）"
                           % (L["title"], s["title"], sub, s["minutes"]))
    return out


def check_duplicate_steps(d):
    out = []
    for L in d["lessons"]:
        seen = {}
        for s in L["steps"]:
            for x in s["subs"]:
                seen.setdefault(x["label"], []).append(x["line"])
        for label, lines in seen.items():
            if len(lines) > 1:
                out.append("【重复】%s 里 [%s 出现 %d 次（行 %s）——整段环节可能重复写了两遍"
                           % (L["title"], label, len(lines),
                              "、".join(str(n) for n in lines)))
    return out


def check_screens(d, body_pages):
    """只在「指令条数多于正文页数」时判定必然缺页；否则清单交人对。

    别拿条数不等当 issue——大量正文页是师话驱动的，本来就没有对应的投屏指令
    （实测 7 条指令 ↔ 21 页正文页，全报出来只是噪声）。
    """
    if len(d["screens"]) > len(body_pages):
        return ["【屏幕】详案有 %d 条投屏/发放指令，pptx 只有 %d 页正文页——必有缺页"
                % (len(d["screens"]), len(body_pages))]
    return []


def in_lesson(no, d):
    """该行是否落在课时区块内（教案提纲表、附录讲评说明不算——它们本就不上屏）。"""
    cur = None
    for sec in d["lessons"]:
        if sec["line"] <= no:
            cur = sec
        else:
            break
    return bool(cur and cur.get("is_lesson"))


def check_model_text(md_path, ppt_text, d):
    """课时区块内 `>` 引用块＝教师示范文，老师要照着念，必须逐字一致。"""
    out = []
    for no, line in enumerate(io.open(md_path, encoding="utf-8").read().split("\n"), 1):
        s = line.strip()
        if not s.startswith(">") or not in_lesson(no, d):
            continue
        body = s.lstrip("> ").strip()
        if len(body) < 20 or body.startswith("【"):
            continue
        if norm(body) not in ppt_text:
            out.append("【示范文】L%d 这一段在 pptx 里找不到原文：%s…" % (no, body[:30]))
    return out


def table_notes(md_path, ppt_text, d):
    """课时区块内表格的数据单元格 ↔ pptx。

    **只作提示、不算 issue**：PPT 常有意压缩表头与说明列（「家里的这一位」→「这一位」），
    压缩得对不对是人的判断；机器只负责把差异摆出来，免得人逐格核。
    表头行（分隔线上面那一行）跳过，差异几乎全在那儿。
    """
    out = []
    lines = io.open(md_path, encoding="utf-8").read().split("\n")
    for i, line in enumerate(lines):
        no, s = i + 1, line.strip()
        if not s.startswith("|") or TABLE_SEP.match(s) or not in_lesson(no, d):
            continue
        if i + 1 < len(lines) and TABLE_SEP.match(lines[i + 1].strip()):
            continue                                    # 表头行
        for cell in [c.strip() for c in s.strip("|").split("|")]:
            if len(cell) < MIN_CELL or cell.startswith("--"):
                continue
            if norm(cell) not in ppt_text:
                out.append("L%d 表格格「%s」在 pptx 里是另一种写法（多半是有意压缩，请人判）"
                           % (no, cell[:30]))
    return out


def check_quotes(md_path, prs):
    out = []
    straight = 0
    for sl in prs.slides:                       # 必须读 run.text：外部件存 &quot; 实体
        for sh in sl.shapes:
            if sh.has_text_frame:
                t = sh.text_frame.text
                straight += t.count("\"") + t.count("'")
    if straight:
        out.append("【引号】pptx 里有 %d 处 ASCII 直引号，与全仓弯引号铁律冲突，投屏前须改"
                   % straight)
    src = io.open(md_path, encoding="utf-8").read()
    if src.count("\"") or src.count("'"):
        out.append("【引号】详案里有 %d 处 ASCII 直引号" % (src.count("\"") + src.count("'")))
    if src.count("“") != src.count("”"):
        out.append("【引号】详案弯双引号不配对：%d 左 / %d 右"
                   % (src.count("“"), src.count("”")))
    return out


def env_page_table(d, body_pages):
    """「详案环节 → 页码」待填对照表。不强制填——强制会让它沦为走过场。"""
    rows = []
    for L in d["lessons"]:
        for s in L["steps"]:
            rows.append({"lesson": L["title"], "step": s["title"],
                         "minutes": s["minutes"], "pages": []})
    return rows


def main():
    ap = argparse.ArgumentParser(description="拿详案机检 pptx，结论回写工作单 audit 字段")
    ap.add_argument("pptx")
    ap.add_argument("--plan", default=None, help="工作单路径（默认 <同名>-anim.json）")
    ap.add_argument("--lesson-plan", default=None, help="详案路径（默认按课次目录名定位）")
    ap.add_argument("--quiet", action="store_true")
    args = ap.parse_args()

    root = project_root(args.pptx)
    jpath = args.plan or os.path.splitext(args.pptx)[0] + "-anim.json"
    if not os.path.isfile(jpath):
        sys.exit("找不到工作单 %s——先跑 inspect_pptx.py" % jpath)
    with io.open(jpath, encoding="utf-8") as f:
        doc = json.load(f)

    md = args.lesson_plan or doc.get("lesson_plan")
    md = os.path.join(root, md) if md and not os.path.isabs(md) else md
    if not md or not os.path.isfile(md):
        try:                                    # 存量工作单没有该字段，自己补，免得重跑 inspect
            md = locate(args.pptx)
        except PlanNotFound as e:
            sys.exit(str(e))
        doc["lesson_plan"] = rel(root, md)
        print("工作单原先没有 lesson_plan 字段，已按课次目录补上 -> %s" % doc["lesson_plan"])

    d = digest(md)
    prs = Presentation(args.pptx)
    pages, ppt_text = slide_texts(prs)
    skips = {int(k) for k, v in doc.get("slides", {}).items() if v.get("skip")}
    body_pages = [i for i in range(1, len(pages) + 1) if i not in skips]

    issues = (check_minutes(d) + check_duplicate_steps(d) + check_screens(d, body_pages)
              + check_model_text(md, ppt_text, d) + check_quotes(md, prs))
    notes = table_notes(md, ppt_text, d)

    doc["lesson_plan"] = rel(root, md)
    doc["plan_digest"] = {"md5": d["md5"],
                          "lessons": [{"title": L["title"],
                                       "steps": [{"title": s["title"], "minutes": s["minutes"]}
                                                 for s in L["steps"]]}
                                      for L in d["lessons"] if L.get("is_lesson")]}
    doc["audit"] = {"md5": d["md5"], "issues": issues, "notes": notes,
                    "step_page_map": env_page_table(d, body_pages)}
    with io.open(jpath, "w", encoding="utf-8", newline="") as f:
        json.dump(doc, f, ensure_ascii=False, indent=1)

    if not args.quiet:
        print("\n详案：%s" % rel(root, md))
        print("pptx：%d 页（正文 %d、跳过 %d）" % (len(pages), len(body_pages), len(skips)))
        print("\n—— 详案环节 → 待人工核对页码 ——")
        for L in d["lessons"]:
            if not L.get("is_lesson"):
                continue
            print("  %s" % L["title"])
            for s in L["steps"]:
                print("     %-40s %s 分钟" % (s["title"][:40], s["minutes"]))
        print("\n—— 详案投屏/发放指令 %d 条 ——" % len(d["screens"]))
        for s in d["screens"]:
            print("  L%-4d %s" % (s["line"], s["text"][:64]))
        print("\n—— 机检 ——")
        if issues:
            for i in issues:
                print("  ! " + i)
        else:
            print("  五类全过")
        if notes:
            print("\n—— 提示（机器判不准，人过一眼）——")
            for n in notes:
                print("  · " + n)

    print("\n审查已记入工作单（%d 条待办、%d 条提示）-> %s" % (len(issues), len(notes), jpath))
    print("机器只查得了这五类；教学顺序、缺页与归类判断仍须读详案人工过一遍，"
          "结论报用户后再排 PLAN。")


if __name__ == "__main__":
    main()

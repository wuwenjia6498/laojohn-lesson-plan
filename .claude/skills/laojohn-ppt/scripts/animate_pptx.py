# -*- coding: utf-8 -*-
"""外部 pptx 动画注入器：按分组工作单，为每页注入「逐条点击淡入」。

写作课 PPT 改由外部平台生成后，本仓只做动画后处理，本脚本是第二段：
    inspect_pptx.py  勘查 -> 分组工作单 JSON（人工校正顺序与取舍）
    animate_pptx.py  按工作单注入动画 -> 可投屏的 pptx

    PYTHONUTF8=1 python animate_pptx.py <foo.pptx> <foo-anim.json> [-o out.pptx] [--in-place]

铁律：动画 XML 一律走 helpers.add_click_reveal，**禁在本文件复制那段时间树**——
helpers.py 是读书会与写作课共用、有 bug 史的横切层，必须保持单一源。

注入前会先清除目标页已有的 <p:timing>，因此可以反复跑（幂等），
也可以拿一份已带动画的 pptx 重新分组、重新注入。
"""
import argparse
import json
import os
import re
import sys
import zipfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx import Presentation                      # noqa: E402
from pptx.util import Emu                          # noqa: E402
from pptx.oxml.ns import qn                        # noqa: E402

from helpers import add_click_reveal               # noqa: E402  单一源，禁复制
from plan_link import digest, project_root         # noqa: E402

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")


def check_audit(plan, pptx, allow_no_plan):
    """闸门：审查过没有。**守的是「做过」，不是「全绿」**——issues 非空照样放行。

    2026-08-06 立。此前两个脚本都不知道详案的存在：那次搜详案的关键词不带弯引号、
    搜不到就判「详案不存在」，整道审查被跳过，26 页里 8 页分组按版面猜错顺序、全部返工。
    文档改成四步治不住这个——脚本照跑不误、不报一个错，所以在这里加硬拦。
    """
    md_rel = plan.get("lesson_plan")
    if md_rel is None:
        if allow_no_plan:
            print("! 未绑定详案（--no-lesson-plan）：分组顺序无据可依，交付时须讲明。")
            return
        sys.exit("工作单没有绑定详案。\n"
                 "  重跑 inspect_pptx.py 让它按课次目录名定位详案，或跑 audit_against_plan.py 补上；\n"
                 "  该课确实没有详案，加 --no-lesson-plan 显式放行。")

    audit = plan.get("audit")
    if not audit:
        sys.exit("工作单没有 audit 字段——审查还没跑。\n"
                 "  先跑：PYTHONUTF8=1 python audit_against_plan.py %r\n"
                 "  （分组顺序以详案为准，审查那一步的产出就是排 PLAN 的依据）" % pptx)

    md = os.path.join(project_root(pptx), md_rel)
    if not os.path.isfile(md):
        sys.exit("工作单记的详案不存在：%s" % md_rel)
    now = digest(md)["md5"]
    if audit.get("md5") != now:
        sys.exit("详案在审查之后改过（md5 对不上）——必须重审，否则分组可能跟着旧稿走。\n"
                 "  重跑：PYTHONUTF8=1 python audit_against_plan.py %r" % pptx)

    issues, notes = audit.get("issues") or [], audit.get("notes") or []
    if issues:
        print("审查记了 %d 条待办（不阻断注入，但交付前要有说法）：" % len(issues))
        for i in issues[:6]:
            print("   ! " + i)
        if len(issues) > 6:
            print("   …另有 %d 条，见工作单 audit.issues" % (len(issues) - 6))
    else:
        print("审查已跑，机检五类全过%s。" % ("（另有 %d 条待人判的提示）" % len(notes) if notes else ""))


def clear_timing(slide):
    """移除该页已有的动画时间树，返回是否移除过。"""
    sld = slide.element
    removed = False
    for tag in ("p:timing", "p:bldLst"):
        for node in sld.findall(qn(tag)):
            sld.remove(node)
            removed = True
    return removed


def check_bindings(slide, groups, idx, page_h):
    """防呆：工作单里的数字是 **shape_id**，不是形状在 slide.shapes 里的位置索引。

    两套编号数值范围重叠（实测 shape_id = 位置索引 + 2），下面那道
    「id 是否存在于本页」的校验拦不住，绑错了照样放行——一整份 19 页的动画
    全部错位、标题被当正文藏起来，是靠投屏才发现的。

    这里查两个「绝不可能是有意为之」的信号：页顶的眉标/标题被卷进动画、
    整页背景被卷进动画。命中任一就停下，别再往下写坏文件。
    """
    picked = {s for g in groups for s in g}
    head, backdrop = [], []
    for sh in slide.shapes:
        if sh.shape_id not in picked:
            continue
        top = Emu(sh.top).inches
        w, h = Emu(sh.width).inches, Emu(sh.height).inches
        if sh.has_text_frame and sh.text_frame.text.strip() and top < page_h * 0.22:
            head.append("id%d「%s」" % (sh.shape_id, sh.text_frame.text.strip()[:16]))
        if h > page_h * 0.85:
            backdrop.append("id%d" % sh.shape_id)
    if head or backdrop:
        msg = ["P%02d 的分组疑似绑错形状：" % idx]
        if head:
            msg.append("  页顶的眉标/标题被卷进了动画 —— " + "、".join(head))
        if backdrop:
            msg.append("  整页背景被卷进了动画 —— " + "、".join(backdrop))
        msg.append("  最常见的原因：工作单里填的是 slide.shapes 的位置索引，"
                   "而这里需要的是 shape_id（两者数值范围会重叠，查不出来）。")
        msg.append("  确属有意为之，加 --allow-header 跳过本检查。")
        sys.exit("\n".join(msg))


def count_clicks(path):
    """从落盘的 pptx 里逐页数 clickEffect，用于回读核验。"""
    z = zipfile.ZipFile(path)
    out = {}
    for name in z.namelist():
        m = re.match(r"ppt/slides/slide(\d+)\.xml$", name)
        if m:
            xml = z.read(name).decode("utf-8", "ignore")
            out[int(m.group(1))] = xml.count('nodeType="clickEffect"')
    z.close()
    return out


def main():
    ap = argparse.ArgumentParser(description="按分组工作单为 pptx 注入逐条点击动画")
    ap.add_argument("pptx", help="待处理的 .pptx")
    ap.add_argument("plan", help="inspect_pptx.py 产出并经人工校正的工作单 .json")
    ap.add_argument("-o", "--out", default=None, help="输出路径（默认 <同名>-anim.pptx）")
    ap.add_argument("--in-place", action="store_true", help="直接覆盖原文件")
    ap.add_argument("--allow-header", action="store_true",
                    help="放行「页顶元素被卷入动画」的防呆检查（默认拦截）")
    ap.add_argument("--dur", type=int, default=None, help="单个淡入时长 ms（默认取工作单的 dur，否则 500）")
    ap.add_argument("--effect", choices=("fade", "appear"), default="fade",
                    help="动画效果：fade 淡入（缺省，外部件链沿用）；appear 出现（仓内直出件，2026-09-30 用户定）")
    ap.add_argument("--no-lesson-plan", action="store_true",
                    help="该课确实没有详案时显式放行审查闸门")
    args = ap.parse_args()

    with open(args.plan, encoding="utf-8") as f:
        plan = json.load(f)
    check_audit(plan, args.pptx, args.no_lesson_plan)
    dur = args.dur or plan.get("dur") or 500

    prs = Presentation(args.pptx)
    page_h = Emu(prs.slide_height).inches
    total_slides = len(prs.slides._sldIdLst)
    if plan.get("slide_count") and plan["slide_count"] != total_slides:
        sys.exit("工作单是按 %d 页做的，当前 pptx 有 %d 页——请重跑 inspect_pptx.py"
                 % (plan["slide_count"], total_slides))

    expect = {}
    done = skipped = 0
    for idx, slide in enumerate(prs.slides, 1):
        entry = plan["slides"].get(str(idx))
        if entry is None:
            expect[idx] = 0
            continue
        clear_timing(slide)
        groups = [[int(s) for s in g] for g in entry.get("groups", [])]
        if entry.get("skip") or not groups:
            expect[idx] = 0
            skipped += 1
            continue
        ids = {sh.shape_id for sh in slide.shapes}
        missing = [s for g in groups for s in g if s not in ids]
        if missing:
            sys.exit("P%02d 工作单里的 shape_id %s 在该页不存在——工作单与 pptx 不匹配"
                     % (idx, missing))
        if not args.allow_header:
            check_bindings(slide, groups, idx, page_h)
        add_click_reveal(slide, groups, dur=dur, effect=args.effect)
        expect[idx] = len(groups)
        done += 1

    out = args.pptx if args.in_place else (args.out or
                                           os.path.splitext(args.pptx)[0] + "-anim.pptx")
    prs.save(out)

    # 回读核验：落盘后重新数一遍，与工作单逐页对账
    actual = count_clicks(out)
    bad = [(i, expect[i], actual.get(i, 0)) for i in expect if actual.get(i, 0) != expect[i]]
    if bad:
        for i, e, a in bad:
            print("P%02d 期望 %d 次点击，实际 %d 次" % (i, e, a))
        sys.exit("回读核验不通过")

    print("注入完成：%d 页加动画、%d 页跳过，共 %d 次点击（%s）"
          % (done, skipped, sum(expect.values()),
             "效果：出现" if args.effect == "appear" else "每次淡入 %dms" % dur))
    print("回读核验通过 -> %s" % out)


if __name__ == "__main__":
    main()

# -*- coding: utf-8 -*-
"""build_ppt.py - 把中间稿 .md 编译为 .pptx

用法：
    python build_ppt.py --input 中间稿.md --output 课件.pptx [--course 导读课] [--book 俗世奇人]
    python build_ppt.py --input 中间稿.md --output 课件.pptx --logo path/to/logo.png --banner path/to/banner.jpg

如果 --logo / --banner 未指定，会自动从 ../assets/logo/ 和 ../assets/decorations/ 取默认资源。
"""
import argparse
import os
import re

from pptx import Presentation

import theme
from parser import parse_md, Page
from layouts_reading import RENDERERS_READING
from layouts_writing import RENDERERS_WRITING
from layouts_promo import RENDERERS_PROMO


# profile 分派表（唯一接缝）：中间稿元信息 `文体：` → 该 profile 的 renderer 字典。
# 键 "" = 缺省读书会（向后兼容，元信息不写文体时走它）。
# 新增课型只在这里加一行 + 新建 layouts_<profile>.py / theme_<profile>.py，
# 底层 helpers / parser 不得出现任何课型分支（见 references/architecture.md）。
PROFILES = {
    "": RENDERERS_READING,
    "写作": RENDERERS_WRITING,
    "宣讲": RENDERERS_PROMO,
}


def find_default_asset(filename_candidates, asset_subdir):
    here = os.path.dirname(os.path.abspath(__file__))
    base = os.path.normpath(os.path.join(here, "..", "assets", asset_subdir))
    if not os.path.isdir(base):
        return None
    for name in filename_candidates:
        p = os.path.join(base, name)
        if os.path.isfile(p):
            return p
    # 兜底：返回该目录第一个 png/jpg
    for fn in sorted(os.listdir(base)):
        if fn.lower().endswith((".png", ".jpg", ".jpeg")):
            return os.path.join(base, fn)
    return None


def jump_back_pages(prs, tol_emu=274320):
    """逐页查点击动画有没有「回跳」：后一击出现的内容顶边高于此前各击（超出 0.3 英寸容差）。
    只读检查、课型无关，返回 [(第几张幻灯片, 第几击)]。"""
    from pptx.oxml.ns import qn
    out = []
    for idx, slide in enumerate(prs.slides, 1):
        seq = slide._element.find(".//" + qn("p:cTn") + "[@nodeType='mainSeq']")
        if seq is None:
            continue
        tops = {sh.shape_id: sh.top for sh in slide.shapes if sh.top is not None}
        steps = seq.find(qn("p:childTnLst"))
        prev = None
        for k, step in enumerate(steps if steps is not None else [], 1):
            ids = {int(t.get("spid")) for t in step.iter(qn("p:spTgt"))}
            ys = [tops[i] for i in ids if i in tops]
            if not ys:
                continue
            top = min(ys)
            if prev is not None and top < prev - tol_emu:
                out.append((idx, k))
            prev = top if prev is None else max(prev, top)
    return out


def build(input_md: str, output_pptx: str, *,
          course_name: str = "", book_title: str = "",
          logo_path: str = None, banner_path: str = None,
          anim: bool = True, strict_images: bool = False) -> dict:
    with open(input_md, encoding="utf-8") as f:
        md = f.read()

    deck = parse_md(md)
    # 命令行参数覆盖元信息
    if course_name:
        deck.course = course_name
    if book_title:
        deck.title = book_title

    # 按文体选 profile（查表，见模块级 PROFILES）。
    # 刻意不用「非写作即读书会」的三元式：那样 `文体：写作课`（多打一个字）会**静默**
    # 落回读书会 renderer，整份稿按错版式烘出来还不报错。
    renderers = PROFILES.get(deck.doc_kind)
    if renderers is None:
        known = "、".join(k or "（省略=读书会）" for k in PROFILES)
        raise ValueError(f"未知文体：{deck.doc_kind!r}（应为：{known}）")

    # 默认资源
    if logo_path is None:
        logo_path = find_default_asset(["logo-red.png", "logo.png"], "logo")
    if banner_path is None:
        banner_path = find_default_asset(["cover-banner.jpg", "banner.jpg", "banner.png"], "decorations")
    # END 页用白色 LOGO（找不到时回退红色）
    logo_white_path = find_default_asset(["logo-white.png"], "logo") or logo_path

    prs = Presentation()
    prs.slide_width = theme.SLIDE_W
    prs.slide_height = theme.SLIDE_H
    blank_layout = prs.slide_layouts[6]

    # ===== 自动注入封面 + END =====
    # 中间稿 P 编号由 ppt-draft 工具自排（每课时 P01=封面、正文从 P02 起），仅作内部引用；PPT 不标注页码。
    pages = list(deck.pages)

    # 1) 若首页不是封面，按"和导读课起始页一样"的规则注入一页封面
    if not pages or pages[0].page_type != "封面":
        synthetic_cover = Page(num=0, page_type="封面")
        # 标题留空：render_cover 会回退到 ctx["book_title"]
        # 副标题留空：会回退到 ctx["course_name"]
        # body 留空：会回退到 ctx["meta"]（作者 · 年级）
        pages.insert(0, synthetic_cover)

    # 2) 始终追加一页 END
    pages.append(Page(num=999, page_type="_END"))

    # 3) 给每个"环节标题"页打章节序号（课时内从 1 递增）
    section_counter = 0
    for p in pages:
        if p.page_type == "环节标题":
            section_counter += 1
            p._section_no = section_counter
    section_total = sum(1 for p in pages if p.page_type == "环节标题")
    for p in pages:
        if p.page_type == "环节标题":
            p._section_total = section_total

    # 页码：仅对中间内容页计数（封面、END 不纳入）
    inner_total = sum(1 for p in pages if p.page_type not in {"封面", "_END"})
    placeholder_report = []
    missing_images = []
    inner_idx = 0

    for page in pages:
        slide = prs.slides.add_slide(blank_layout)
        is_inner = page.page_type not in {"封面", "_END"}
        if is_inner:
            inner_idx += 1

        ctx = {
            "course_name": deck.course or course_name,
            "book_title": deck.title or book_title,
            "page_index": inner_idx,
            "page_total": inner_total,
            "logo_path": logo_path,
            "logo_white_path": logo_white_path,
            "banner_path": banner_path,
            "meta": "　·　".join(p for p in (deck.author, deck.grade) if p),
            "doc_kind": deck.doc_kind,
            # 主题色：课型无关透传；读书会 v9 新样式据此选色板（缺省空＝现行样式）
            "theme_color": deck.theme_color,
            "anim": anim,
            "input_dir": os.path.dirname(os.path.abspath(input_md)),
        }
        renderer = renderers.get(page.page_type)
        if renderer is None:
            raise ValueError(f"P{page.num} 无渲染器：{page.page_type}")
        renderer(slide, page, ctx)

        # 缺图清单（课型无关）：页上有有效配图建议、但没有可用真图文件
        if page.num not in {0, 999}:
            _sugs = [s for s in (page.image_suggestions or []) if s.strip() not in
                     {"无", "无（页面已满）", "—", "无配图"}]
            _img = page.image_path
            if _img and not os.path.isabs(_img):
                _img = os.path.join(ctx["input_dir"], _img)
            if _sugs and not (_img and os.path.isfile(_img)):
                missing_images.append((page.num, page.page_type, page.image_path or _sugs[0]))

        # 演讲者备注 → pptx 自带备注页（放映时只有讲者看得到）。
        # 课型无关：中间稿不写 `备注：` 就是空字符串，此处直接跳过。
        if getattr(page, "notes", ""):
            slide.notes_slide.notes_text_frame.text = page.notes

        # 占位清单（仅对中间稿原生页报告，跳过 synthetic）
        if page.page_type in {"_END", "封面"} and page.num in {0, 999}:
            continue
        if page.page_type == "填空表格" and page.table_reveals:
            placeholder_report.append(
                (page.num, page.page_type,
                 f"[填空表格] {len(page.table_reveals)} 格答案逐格点击"))
            continue
        raw_sugs = page.image_suggestions or ([page.image_suggestion] if page.image_suggestion else [])
        sugs = [s.strip() for s in raw_sugs
                if s and s.strip() not in {"无", "无（页面已满）", "—", "无配图"}]
        if sugs and page.page_type in {"引导问题", "要点小结", "环节标题"}:
            if len(sugs) >= 2:
                placeholder_report.append(
                    (page.num, page.page_type, f"[四图网格×{len(sugs)}] " + " ／ ".join(sugs)))
            else:
                placeholder_report.append((page.num, page.page_type, sugs[0]))
        else:
            placeholder_report.append((page.num, page.page_type, "—"))

    total = len(pages)
    if strict_images and missing_images:
        lines = "\n".join(f"  P{n:02d} {t}：{s}" for n, t, s in missing_images)
        raise SystemExit(f"--strict-images：{len(missing_images)} 页缺图，未出件\n{lines}")

    os.makedirs(os.path.dirname(os.path.abspath(output_pptx)) or ".", exist_ok=True)
    prs.save(output_pptx)
    jumps = jump_back_pages(prs) if deck.theme_color else []
    lint = draft_lint(deck)

    return {
        "pages": total,
        "output": os.path.abspath(output_pptx),
        "placeholders": placeholder_report,
        "missing_images": missing_images,
        "theme_color": deck.theme_color,
        "jump_back": jumps,
        "lint": lint,
        "logo": logo_path,
        "banner": banner_path,
    }


# 上屏答案里的「判定口径」：教学生怎么答、老师怎么判对错的话，不是答案本身（2026-10-08 用户反馈）
_GRADING_RE = re.compile(r"就算数|都算对|都对[，,]|就好[。！]?$|就很好|就是好的|都行|关键说清|重要的是有自己的理由")


def draft_lint(deck) -> list:
    """中间稿两项机检（只报警、不拦）：参考答案带判定口径；留白填写表没写 `列宽：`。"""
    out = []
    for pg in deck.pages:
        for ans in pg.bullet_answers:
            if ans and _GRADING_RE.search(ans):
                out.append((pg.num, "参考答案带判定口径（只上答案本身，口径进讲稿）", ans[-24:]))
        if pg.table_headers and not pg.col_weights:
            has_blank = any(not str(row[c] if c < len(row) else "").strip()
                            and (d, c) not in (pg.table_reveals or {})
                            for d, row in enumerate(pg.table_rows) for c in range(len(pg.table_headers)))
            if has_blank:
                out.append((pg.num, "留白填写表未写 `列宽：`（按预估填写量分宽）", " | ".join(pg.table_headers)))
    return out


def main():
    ap = argparse.ArgumentParser(description="老约翰投屏 PPT 编译器")
    ap.add_argument("--input", "-i", required=True, help="中间稿 .md 路径")
    ap.add_argument("--output", "-o", required=True, help="输出 .pptx 路径")
    ap.add_argument("--course", default="", help="课时名（覆盖中间稿元信息）")
    ap.add_argument("--book", default="", help="书名（覆盖中间稿元信息）")
    ap.add_argument("--logo", default=None, help="LOGO 图片路径")
    ap.add_argument("--banner", default=None, help="封面横幅图片路径")
    ap.add_argument("--no-anim", dest="anim", action="store_false",
                    help="关闭逐条点击动画（默认开启：要点小结/引导问题追问逐条淡入）")
    ap.add_argument("--strict-images", action="store_true",
                    help="有配图建议却没有可用真图的页一律报错、不出件（新样式定稿用）")
    args = ap.parse_args()

    info = build(
        args.input, args.output,
        course_name=args.course,
        book_title=args.book,
        logo_path=args.logo,
        banner_path=args.banner,
        anim=args.anim,
        strict_images=args.strict_images,
    )

    print(f"[OK] 已生成 {info['pages']} 页 -> {info['output']}")
    print(f"     LOGO: {info['logo']}")
    print(f"     横幅: {info['banner']}")
    print()
    print("配图占位清单：")
    for num, ptype, sug in info["placeholders"]:
        print(f"  P{num:02d}  [{ptype}]  {sug}")
    if info["theme_color"]:     # 新样式下缺图＝该页按无图版式出，须补图
        print()
        print(f"缺图清单（新样式·主题色{info['theme_color']}）：{len(info['missing_images'])} 页")
        for num, ptype, sug in info["missing_images"]:
            print(f"  P{num:02d}  [{ptype}]  {sug}")
        jb = info["jump_back"]
        print(f"动画回跳：{len(jb)} 处" + ("" if not jb else "（须改分组）：" +
              "、".join(f"第{i}张幻灯片第{k}击" for i, k in jb)))
    lint = info.get("lint") or []
    print()
    print(f"中间稿机检：{len(lint)} 处" + ("" if not lint else "（须回中间稿改）"))
    for num, what, frag in lint:
        print(f"  P{num:02d}  {what}：…{frag}")


if __name__ == "__main__":
    main()

# -*- coding: utf-8 -*-
"""看图写话详案 docx 首页版式化（样式源＝用户提供《详案首页示例.docx》，2026-07-21 拍板）。

用法：
    PYTHONUTF8=1 python style_front_page.py <详案.md> <目标.docx> [<目标2.docx> ...]

职责：
  1) 从详案 .md 解析首页数据（H1 标题、「教案提纲表」表、「这一课在整个课程里的位置」表、校内基线注）；
  2) 打开共享引擎已产出的 docx（无图版或配图版均可），删除其开头的朴素首页块
     （H1 至第一个「第N课时」分页标题之前的全部元素）；
  3) 重建示例版式首页：居中三行标题区 → 「本课提纲表」节标题 → 10×2 分组表
     （「本课要点」＋「这一课在整个课程里的位置」两个跨列组头）→ 校内基线左色条提示框。

只做首页重排，不碰正文、页眉、图位；不修改共享 docx 引擎。样式常量以本文件为单一源。
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

# ---- 样式常量（照抄示例 docx 实测值 + 2026-07-21 截图定稿） ----
NAVY = "1F3A5F"      # 深蓝：标题/标签
GRAY = "5A5A5A"      # 副题灰 / 「这一期」括号注
BODY = "222222"      # 表格正文
RED = "C00000"       # 关键词强调红（课型/本期训练点/增量词）
NOTE_BODY = "333333"  # 基线框正文
BORDER = "4E8FA6"    # 表格边框 / 基线框左色条
GROUP_BG = "E4EEF2"  # 组头底纹
LABEL_BG = "EAF1F4"  # 标签列底纹
NOTE_BG = "F4F7F8"   # 基线框底纹
NOTE_BORDER = "DDE6EA"  # 基线框细边
RULE = "D9D9D9"      # 标题区细分隔线
FONT = "宋体"
BRAND_LINE = "老约翰 · 看图写话 24 期方法主线"
TABLE_W = 9386       # dxa
COL_LABEL_W = 1760
COL_VALUE_W = TABLE_W - COL_LABEL_W
NOTE_W = 9072


# ---------- md 解析 ----------
def parse_md(md_path: Path):
    text = md_path.read_text(encoding="utf-8")
    lines = text.splitlines()

    m = re.search(r"^#\s+(.+?)\s*·\s*看图写话详案\s*$", text, re.M)
    if not m:
        raise SystemExit("未找到 H1「# 期N《…》· 看图写话详案」")
    title = m.group(1).strip()

    def read_table(after_pat):
        """返回 after_pat 之后第一个 md 两列表的 [(label, value), ...]（含表头行）。"""
        idx = None
        for i, l in enumerate(lines):
            if re.search(after_pat, l):
                idx = i
                break
        if idx is None:
            raise SystemExit(f"未找到锚：{after_pat}")
        rows = []
        started = False
        for l in lines[idx + 1:]:
            s = l.strip()
            if s.startswith("|"):
                started = True
                cells = [c.strip() for c in s.strip("|").split("|")]
                if len(cells) >= 2 and not set(cells[0]) <= {"-", " ", ":"}:
                    rows.append((cells[0], cells[1]))
            elif started:
                break
        if not rows:
            raise SystemExit(f"锚 {after_pat} 后未解析到表格")
        return rows

    tbl_main = read_table(r"^##\s*教案提纲表")
    tbl_pos = read_table(r"^\*\*这一课在整个课程里的位置\*\*")

    m2 = re.search(r"^>\s*校内基线：(.+)$", text, re.M)
    if not m2:
        raise SystemExit("未找到「> 校内基线：」注")
    baseline = m2.group(1).strip()

    # 副题 = 课次·课型 值的前两段（│ 分隔）
    kv = dict(tbl_main)
    course = kv.get("课次·课型", "")
    parts = [p.strip() for p in re.split(r"[│|]", course) if p.strip()]
    subtitle = " · ".join(parts[:2]) if parts else course
    return title, subtitle, tbl_main, tbl_pos, baseline


# ---------- docx 低层小件 ----------
def _set_font(run, size, bold=False, color=BODY):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), FONT)
    rfonts.set(qn("w:hAnsi"), FONT)
    rfonts.set(qn("w:eastAsia"), FONT)


def _cell_bg(cell, hexval):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexval)
    tcpr.append(shd)


def _cell_width(cell, dxa):
    tcpr = cell._tc.get_or_add_tcPr()
    tcw = OxmlElement("w:tcW")
    tcw.set(qn("w:w"), str(dxa))
    tcw.set(qn("w:type"), "dxa")
    tcpr.append(tcw)


def _tbl_borders(tbl, spec):
    """spec: dict side -> (val, color, sz)；side ∈ top/left/bottom/right/insideH/insideV"""
    tblpr = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side, (val, color, sz) in spec.items():
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), val)
        el.set(qn("w:color"), color)
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        borders.append(el)
    tblpr.append(borders)


def _tbl_fixed_width(tbl, dxa):
    tblpr = tbl._tbl.tblPr
    tblw = OxmlElement("w:tblW")
    tblw.set(qn("w:w"), str(dxa))
    tblw.set(qn("w:type"), "dxa")
    tblpr.append(tblw)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)


def _cell_text(cell, text, size, bold, color, space=3, center=False):
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(space)
    p.paragraph_format.space_after = Pt(space)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    _set_font(run, size, bold, color)
    return p


def _cell_valign_center(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    va = OxmlElement("w:vAlign")
    va.set(qn("w:val"), "center")
    tcpr.append(va)


def _p_rule(p, side="bottom", color=RULE, sz=4):
    """段落边框横线（标题区分隔线）：side ∈ top/bottom。"""
    ppr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:color"), color)
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), "6")
    pbdr.append(el)
    ppr.append(pbdr)


def _p_left_bar(p, color=NAVY, sz=24):
    """段落左侧粗竖条（节标题装饰）。"""
    ppr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    el = OxmlElement("w:left")
    el.set(qn("w:val"), "single")
    el.set(qn("w:color"), color)
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), "4")
    pbdr.append(el)
    ppr.append(pbdr)


# ---------- 值列富文本渲染（截图定稿的红词/灰注规则） ----------
SEP = "　│　"  # 全角分隔


def _render_value(cell, label, value, stage_name, space=3):
    """按行名定制值列样式：课型红、期名红、增量词红、括号注灰、主题/统编行加粗。"""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(space)
    p.paragraph_format.space_after = Pt(space)

    def run(text, size=11.5, bold=False, color=BODY):
        _set_font(p.add_run(text), size, bold, color)

    if label == "课次·课型":
        parts = [s.strip() for s in re.split(r"[│|]", value) if s.strip()]
        for i, seg in enumerate(parts):
            if i:
                run("　│　", 12, False, GRAY)
            # 第二段=课型 → 红
            run(seg, 12, True, RED if i == 1 else NAVY)
    elif label == "这一期":
        m = re.match(r"^(.*?)(（.*）)?\s*$", value)
        head = (m.group(1) if m else value).strip()
        note = (m.group(2) or "") if m else ""
        run(head, 11.5, True, NAVY)
        if note:
            run("　" + note.strip(), 10.5, False, GRAY)
    elif label == "本学期主题":
        run(value, 11.5, True, BODY)
    elif label == "能力线":
        key = stage_name
        if key and key not in value and "：" in key:
            key = key.split("：", 1)[1]  # 期名含副题时退化到冒号后半段
        if key and key in value:
            pre, post = value.split(key, 1)
            run(pre, 11, False, BODY)
            run(key, 11, True, RED)
            run(post, 11, False, BODY)
        else:
            run(value, 11, False, BODY)
    elif label == "对齐统编":
        parts = [s.strip() for s in re.split(r"[│|]", value) if s.strip()]
        for i, seg in enumerate(parts):
            if i:
                run("　│　", 11.5, False, GRAY)
            m = re.match(r"^(增量：)([^（(]+)(.*)$", seg)
            if m:
                run(m.group(1), 11.5, True, BODY)
                run(m.group(2).strip(), 11.5, True, RED)
                if m.group(3):
                    run(m.group(3), 11, False, BODY)
            else:
                run(seg, 11.5, True, BODY)
    else:  # 教什么 / 学生带走 / 用什么
        run(value, 11.5, False, BODY)
    return p


# ---------- 首页构建 ----------
def build_front(doc, data):
    title, subtitle, tbl_main, tbl_pos, baseline = data
    blocks = []

    def para(text, size, bold, color, align=WD_ALIGN_PARAGRAPH.CENTER,
             before=0, after=6):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        _set_font(p.add_run(text), size, bold, color)
        blocks.append(p._p)
        return p

    stage_name = ""
    m = re.search(r"《(.+?)》", title)
    if m:
        stage_name = m.group(1)

    para(title, 20, True, NAVY, before=6, after=4)
    para(subtitle, 13, False, GRAY, after=8)
    p_brand = para(BRAND_LINE, 12, True, NAVY, after=14)
    _p_rule(p_brand, side="top", color=BORDER, sz=12)  # 品牌行上方一条粗线，浅蓝灰（2026-07-21 用户定稿）
    p_sec = para("本课提纲表", 14, True, NAVY, align=WD_ALIGN_PARAGRAPH.LEFT, after=6)
    _p_left_bar(p_sec)

    # 主表：组头 + 4 行 ×2 组
    n_rows = 2 + len(tbl_main) + len(tbl_pos)
    tbl = doc.add_table(rows=n_rows, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    _tbl_fixed_width(tbl, TABLE_W)
    _tbl_borders(tbl, {s: ("single", BORDER, 4) for s in
                       ("top", "left", "bottom", "right", "insideH", "insideV")})
    def group_row(ri, text):
        row = tbl.rows[ri]
        row.cells[0].merge(row.cells[1])
        cell = row.cells[0]
        _cell_bg(cell, GROUP_BG)
        _cell_text(cell, text, 11, True, NAVY)

    def data_row(ri, label, value):
        c0, c1 = tbl.rows[ri].cells
        _cell_width(c0, COL_LABEL_W)
        _cell_width(c1, COL_VALUE_W)
        _cell_bg(c0, LABEL_BG)
        _cell_bg(c1, "FFFFFF")
        _cell_text(c0, label, 11.5, True, NAVY, center=True)
        _cell_valign_center(c0)
        _cell_valign_center(c1)
        _render_value(c1, label, value, stage_name)

    ri = 0
    group_row(ri, "本课要点"); ri += 1
    for label, value in tbl_main:
        data_row(ri, label, value); ri += 1
    group_row(ri, "这一课在整个课程里的位置"); ri += 1
    for label, value in tbl_pos:
        data_row(ri, label, value); ri += 1
    blocks.append(tbl._tbl)

    # 间隔小段
    gap = doc.add_paragraph()
    gap.paragraph_format.space_before = Pt(2)
    gap.paragraph_format.space_after = Pt(2)
    blocks.append(gap._p)

    # 基线框（左粗色条）
    note = doc.add_table(rows=1, cols=1)
    _tbl_fixed_width(note, NOTE_W)
    _tbl_borders(note, {
        "top": ("single", NOTE_BORDER, 2),
        "bottom": ("single", NOTE_BORDER, 2),
        "right": ("single", NOTE_BORDER, 2),
        "left": ("single", BORDER, 24),
        "insideH": ("none", "auto", 0),
        "insideV": ("none", "auto", 0),
    })
    cell = note.rows[0].cells[0]
    _cell_bg(cell, NOTE_BG)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(5)
    _set_font(p.add_run("校内基线　"), 11, True, NAVY)
    _set_font(p.add_run(baseline), 11, False, NOTE_BODY)
    blocks.append(note._tbl)
    return blocks


def restyle(md_path: Path, docx_path: Path):
    data = parse_md(md_path)
    doc = Document(str(docx_path))
    body = doc.element.body

    # 找 anchor：第一个以「第N课时」开头的段
    anchor = None
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            texts = child.findall(qn("w:r") + "/" + qn("w:t"))
            txt = "".join(t.text or "" for t in texts).strip()
            if re.match(r"^第\s*\d+\s*课时", txt) or txt.startswith("第1课时"):
                anchor = child
                break
    if anchor is None:
        raise SystemExit(f"{docx_path.name}: 未找到「第N课时」锚段，放弃重排")

    # 删除 anchor 之前的全部块（旧首页）
    removed = 0
    for child in list(body.iterchildren()):
        if child is anchor:
            break
        if child.tag in (qn("w:p"), qn("w:tbl")):
            body.remove(child)
            removed += 1

    # 构建新首页（先 append 到文末，再整体搬到 anchor 前）
    blocks = build_front(doc, data)
    for el in blocks:
        anchor.addprevious(el)

    # 首页独占一页：给「第N课时」锚段补段前分页（重排可能吞掉引擎原有的分页衔接）
    ppr = anchor.find(qn("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        anchor.insert(0, ppr)
    if ppr.find(qn("w:pageBreakBefore")) is None:
        ppr.append(OxmlElement("w:pageBreakBefore"))

    doc.save(str(docx_path))
    print(f"首页版式化完成：{docx_path.name}（删旧块 {removed}，插新块 {len(blocks)}）")


def main():
    if len(sys.argv) < 3:
        raise SystemExit(__doc__)
    md_path = Path(sys.argv[1])
    for target in sys.argv[2:]:
        restyle(md_path, Path(target))


if __name__ == "__main__":
    main()

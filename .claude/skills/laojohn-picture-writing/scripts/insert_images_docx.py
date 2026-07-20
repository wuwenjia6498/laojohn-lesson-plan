# -*- coding: utf-8 -*-
"""
insert_images_docx —— 看图写话「回插后处理」。不碰 CLAUDE.md §3 跨技能共享 docx 引擎：
先用共享引擎把详案 .md 出基础 docx（引擎把 `【图位:编号】` 原样印为文字），
再用 python-docx 打开，把每个 `【图位:编号】` 标记换成 图位/<编号>.png 真图。
段内只有标记则就地换图；标记与正文（师话等）同段时**只剥标记、正文留在原段**，图另起一段——
整段清空会把老师要念的问题句一起吞掉。缺图则留灰字「（图位 <编号> 待补）」，不抛错。

用法：
  PYTHONUTF8=1 python insert_images_docx.py <详案.md> [--base-docx 已有基础.docx] [--width-cm 12]
输出：<详案名>-配图.docx（不覆盖无图版，便于对照）。
"""
import os
import re
import sys
import tempfile
import argparse
import subprocess

import imgspec_parser as ip

_ENGINE = os.path.normpath(os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    '..', '..', 'laojohn-lesson-plan', 'assets', 'md_to_laojohn_docx.py'))

_PLACE_RE = re.compile(r'【图位[：:]\s*([锚例格]-\d+)[^】]*】')

# 看图写话的 docx 页眉（覆盖引擎默认的「老约翰深度阅读 / 阅读·思辨·表达」）。
# 右侧标语＝本课程三阶主张，与六阶总轴第一/二阶对应。
HEADER_LEFT  = '老约翰·看图写话'
HEADER_RIGHT = '说清楚·写完整·连成段'


def _ensure_docx():
    try:
        import docx  # noqa
        return
    except ImportError:
        pass
    for args in (
        [sys.executable, '-m', 'pip', 'install', '-q', 'python-docx'],
        [sys.executable, '-m', 'pip', 'install', '-q', '--user', 'python-docx'],
    ):
        try:
            subprocess.check_call(args)
            import docx  # noqa
            return
        except Exception:
            continue
    print('[错误] 无法自动安装 python-docx，请手动安装后重试。')
    sys.exit(1)


def build_base_docx(md_path):
    """调共享引擎产基础 docx（落临时文件，绝不碰用户已导出的同名 docx）。"""
    fd, base = tempfile.mkstemp(prefix='看图写话基础_', suffix='.docx')
    os.close(fd)
    subprocess.check_call([sys.executable, _ENGINE, md_path, base,
                           '--header-left', HEADER_LEFT,
                           '--header-right', HEADER_RIGHT])
    return base


def _clear_paragraph(p):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)


def _strip_placeholders(p):
    """就地剥除段内占位标记，保留同段其余正文（师话等），返回剥后文本。

    曾因整段清空导致与占位同段的师话被静默吞掉（老师拿到的配图版少了问题句）。
    优先逐 run 替换以保住加粗等格式；标记被引擎拆到多个 run 时回退为纯文本重建。
    """
    for r in p.runs:
        if _PLACE_RE.search(r.text):
            r.text = _PLACE_RE.sub('', r.text)
    if _PLACE_RE.search(p.text):  # 跨 run 拆分，逐 run 剥不掉
        residual = _PLACE_RE.sub('', p.text)
        _clear_paragraph(p)
        if residual.strip():
            p.add_run(residual.strip())
    return p.text.strip()


def _new_paragraph_after(p):
    """在 p 之后插入一个同级空段落，用于承载图片。"""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph

    new_el = OxmlElement('w:p')
    p._element.addnext(new_el)
    np = Paragraph(new_el, p._parent)
    # 继承样式，避免图片段落掉回 Normal 造成间距突变
    try:
        np.style = p.style
    except Exception:
        pass
    return np


def insert(md_path, base_docx, width_cm):
    from docx import Document
    from docx.shared import Cm, RGBColor

    doc = Document(base_docx)
    inserted, missing, fmt_skipped = [], [], []
    for p in doc.paragraphs:
        codes = [m.group(1) for m in _PLACE_RE.finditer(p.text)]
        if not codes:
            continue
        # 同段可含多个占位（如一行里连列例-02/03/04）：逐个回插，不得只取首个
        fmt_skipped.extend(c for c in codes if c.startswith('格-'))
        codes = [c for c in codes if not c.startswith('格-')]
        if not codes:
            continue
        # 剥标记而非清整段：段内若还有正文就把正文留在原段，图另起一段
        residual = _strip_placeholders(p)
        target = _new_paragraph_after(p) if residual else p
        if not residual:
            _clear_paragraph(target)
        p = target
        for code in codes:
            png = ip.image_path(md_path, code)
            if os.path.isfile(png):
                run = p.add_run()
                run.add_picture(png, width=Cm(width_cm))
                inserted.append(code)
            else:
                run = p.add_run(f'【图位 {code} 待补：未找到 {os.path.basename(png)}】')
                try:
                    run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
                except Exception:
                    pass
                missing.append(code)
        try:
            p.alignment = 1  # 居中
        except Exception:
            pass

    out = os.path.splitext(md_path)[0] + '-配图.docx'
    doc.save(out)
    return out, inserted, missing, fmt_skipped


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('md_path')
    ap.add_argument('--base-docx', default='', help='已有基础 docx；不给则现调引擎生成')
    ap.add_argument('--width-cm', type=float, default=12.0)
    args = ap.parse_args()

    _ensure_docx()
    temp_base = not args.base_docx
    base = args.base_docx or build_base_docx(args.md_path)
    try:
        out, inserted, missing, fmt = insert(args.md_path, base, args.width_cm)
    finally:
        if temp_base and os.path.exists(base):
            try:
                os.remove(base)
            except OSError:
                pass
    print(f'回插完成：{out}')
    print(f'  插入 {len(inserted)} 张：{inserted}')
    if missing:
        print(f'  缺图待补 {len(missing)}：{missing}（已在文中留灰字提示）')
    if fmt:
        print(f'  跳过格式图 {len(set(fmt))}：{sorted(set(fmt))}（固定资产，手动插）')


if __name__ == '__main__':
    main()

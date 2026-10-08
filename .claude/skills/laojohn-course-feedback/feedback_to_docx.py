#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
laojohn-course-feedback : 渲染引擎
将一份「课程反馈话术」结构化 JSON 渲染进老约翰品牌 .docx 模板。

模板 assets/template.docx 已预置页眉(老约翰深度阅读 / 阅读·思辨·表达 + logo
+ 分隔线)与正文样式,本脚本只负责往空白正文里按固定排版写入三段内容。

用法:
    python3 feedback_to_docx.py <content.json> <output.docx>

content.json 结构(由 SKILL.md 约定的写作流程产出):
{
  "book_title": "俗世奇人",
  "segments": [
    {
      "heading": "一、导读课后:",
      "salutation": true,                 # 是否插入「亲爱的家长:您好!」
      "intro": "今天我们开始了一本新书《俗世奇人》的阅读。……谢谢。",
      "tasks": [                          # 编号家庭任务,可为空列表
        {"title": "天天爱读书", "body": "请督促孩子保持每日阅读……每天至少需要读 XX 页。"},
        {"title": "……", "body": "……"}
      ],
      "closing": "祝小朋友大朋友阅读愉快!"   # 收尾祝语,可省略
    },
    ...
  ]
}
"""
import sys
import json
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

SLOGAN_1 = "更好的协助者,"
SLOGAN_2 = "成就更好的阅读者!"
BRAND = RGBColor(0xE8, 0x55, 0x3A)   # 老约翰品牌橙(与其他 skill 一致)

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, "assets", "template.docx")


def _set_font(run, *, bold=False, size=None, color=None):
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    rf.set(__import__("docx").oxml.ns.qn("w:eastAsia"), "宋体")
    run.font.name = "宋体"
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color


def _add_para(doc, align=None, space_after=6):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.5
    return p


def add_slogan(doc):
    """两行标语,与正文一致:左对齐、黑色、不加粗,作段间分隔。"""
    for line in (SLOGAN_1, SLOGAN_2):
        p = _add_para(doc, space_after=0)
        r = p.add_run(line)
        _set_font(r, size=11)


def add_title(doc, book_title):
    p = _add_para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    r = p.add_run("课后服务内容 —— 社群反馈")
    _set_font(r, bold=True, size=15)


def add_card_image(doc, path):
    """书目卡图片：居中、宽 9cm（长图按比例），供家长群转发时一眼看到书。"""
    p = _add_para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    p.add_run().add_picture(path, width=Cm(9))


def add_segment(doc, seg, card=None):
    # 小标题前空一行
    _add_para(doc, space_after=0)
    # 小标题
    p = _add_para(doc, space_after=4)
    r = p.add_run(seg["heading"])
    _set_font(r, bold=True, size=12)

    # 称呼
    if seg.get("salutation", True):
        for line in ("亲爱的家长:", "您好!"):
            p = _add_para(doc, space_after=0)
            r = p.add_run(line)
            _set_font(r, size=11)

    # 正文导入
    if seg.get("intro"):
        p = _add_para(doc)
        r = p.add_run(seg["intro"])
        _set_font(r, size=11)

    # 书目卡图片（仅第一段、插在第一条任务「天天爱读书」之前）
    if card:
        add_card_image(doc, card)

    # 编号任务
    for i, task in enumerate(seg.get("tasks", []), start=1):
        p = _add_para(doc, space_after=4)
        r = p.add_run(f"{i}、{task['title']}:")
        _set_font(r, bold=True, size=11)
        r2 = p.add_run(task["body"])
        _set_font(r2, size=11)

    # 收尾祝语
    if seg.get("closing"):
        p = _add_para(doc, space_after=0)
        r = p.add_run(seg["closing"])
        _set_font(r, size=11)


def find_card(content, output):
    """同目录下的 `<书名>_书目卡.jpg`（书目卡物料的固定产物名）；没有就返回 None、照旧不插图。"""
    path = os.path.join(os.path.dirname(os.path.abspath(output)),
                        f"{content.get('book_title', '')}_书目卡.jpg")
    return path if os.path.isfile(path) else None


def render(content, output):
    doc = Document(TEMPLATE)
    card = find_card(content, output)

    add_title(doc, content.get("book_title", ""))
    add_slogan(doc)

    segments = content["segments"]
    for idx, seg in enumerate(segments):
        add_segment(doc, seg, card=card if idx == 0 else None)
        # 段间(及末尾)用标语分隔
        add_slogan(doc)

    doc.save(output)
    return output


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    content_path, output = sys.argv[1], sys.argv[2]
    with open(content_path, encoding="utf-8") as f:
        content = json.load(f)
    out = render(content, output)
    print(f"已生成: {out}")


if __name__ == "__main__":
    main()

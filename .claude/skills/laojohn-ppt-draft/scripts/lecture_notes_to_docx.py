# -*- coding: utf-8 -*-
"""
逐页讲稿 .md -> .docx 转换器（讲稿专用干净版式）

用法：
    python lecture_notes_to_docx.py 讲稿.md [输出.docx]
    省略输出路径时，默认与输入同名同目录、改后缀 .docx。

讲稿 Markdown 契约（由 laojohn-ppt-draft 产出）：
    # 《书名》课型 · 逐页讲稿（…）        —— 文档标题（H1，单 #）
    > 用法说明……                          —— 引导块（多行），渲染成顶部说明
    ---                                    —— 页分隔线（忽略）
    ## 第N页 · [页型] 标题                  —— 页头（封面/END 页标题可空）
    ## END 页                              —— 结束页
    **回溯**：…｜**节奏**：…                —— 同行合并的元信息（灰色小字）
    **口播**（可选舞台提示）：…             —— 主念稿（可跨多行；续行无标签）
    **兜底参考**：…                         —— 学生答不出时的收口（橙色，仅老师看）
行内 **加粗** 会渲染成真加粗；未配对的残余 ** 自动剥除、不外泄。

设计：不带详案课案封面页；每页一条横分隔线 + 「第N页·页型」色标 + 加粗标题，
口播为主体可念字号，回溯/节奏为灰色元信息，兜底参考为橙色提示。
"""
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

# ---- 调色板（与 PPT 表头深蓝灰 / 详案页标橙 呼应）----
C_TITLE = RGBColor(0x44, 0x54, 0x6A)   # 深蓝灰：文档标题、页头色标
C_BODY = RGBColor(0x22, 0x22, 0x22)    # 近黑：口播正文
C_META = RGBColor(0x90, 0x90, 0x90)    # 灰：回溯/节奏元信息、用法说明
C_FALLBACK = RGBColor(0xC0, 0x6A, 0x12)  # 橙棕：兜底参考（仅老师看）
C_LABEL_KOU = RGBColor(0x2E, 0x7D, 0x6B)  # 青绿：口播标签

FONT = "微软雅黑"

FIELD_SPLIT = re.compile(r"\*\*(口播|回溯|节奏|兜底参考)\*\*")
PAGE_RE = re.compile(r"^第\s*(\d+)\s*页\s*·\s*\[([^\]]+)\]\s*(.*)$")


def set_font(run, size=None, color=None, bold=False):
    run.font.name = FONT
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), FONT)
    rfonts.set(qn("w:ascii"), FONT)
    rfonts.set(qn("w:hAnsi"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color


def add_top_rule(paragraph):
    """给段落加一条顶部细横线，作为页与页的分隔。"""
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), "6")
    top.set(qn("w:space"), "6")
    top.set(qn("w:color"), "C9D1DB")
    borders.append(top)
    p_pr.append(borders)


def add_inline(paragraph, text, size, color, base_bold=False):
    """渲染含 **加粗** 的行内文本；未配对的 ** 自动剥除。"""
    parts = text.split("**")
    for i, seg in enumerate(parts):
        if seg == "":
            continue
        run = paragraph.add_run(seg)
        set_font(run, size=size, color=color, bold=base_bold or (i % 2 == 1))


def strip_lead(seg):
    """去掉字段值前导的「（舞台提示）」与冒号，返回 (提示, 正文)。"""
    note = ""
    m = re.match(r"^\s*（([^）]*)）", seg)
    if m:
        note = m.group(1)
        seg = seg[m.end():]
    seg = re.sub(r"^\s*[：:]\s*", "", seg)
    return note, seg.strip()


def parse_fields(line):
    """把一行拆成 [(label, value), ...]；非字段行返回 []。"""
    if not FIELD_SPLIT.search(line):
        return []
    tokens = FIELD_SPLIT.split(line)  # [pre, label, seg, label, seg, ...]
    out = []
    for i in range(1, len(tokens), 2):
        label = tokens[i]
        seg = tokens[i + 1] if i + 1 < len(tokens) else ""
        seg = seg.rstrip("｜| ").strip()
        out.append((label, seg))
    return out


def parse(md_text):
    lines = md_text.splitlines()
    title = ""
    intro = []
    pages = []
    cur = None
    last_field = None  # 用于续行归并

    def new_page(num, ptype, ptitle):
        nonlocal cur, last_field
        cur = {"num": num, "type": ptype, "title": ptitle, "fields": []}
        pages.append(cur)
        last_field = None

    for raw in lines:
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("# ") and not line.startswith("## "):
            title = line[2:].strip()
            continue
        if line.startswith(">"):
            intro.append(line.lstrip(">").strip())
            continue
        if line.strip() == "---":
            continue
        if line.startswith("## "):
            head = line[3:].strip()
            m = PAGE_RE.match(head)
            if m:
                new_page(int(m.group(1)), m.group(2).strip(), m.group(3).strip())
            elif head.startswith("END"):
                new_page(None, "END", "")
            else:
                new_page(None, head, "")
            continue
        # 内容行
        fields = parse_fields(line)
        if fields:
            for label, seg in fields:
                note, val = strip_lead(seg)
                if cur is None:
                    continue
                cur["fields"].append({"label": label, "note": note, "lines": [val] if val else []})
                last_field = cur["fields"][-1]
        else:
            # 续行：并入最近字段
            if last_field is not None:
                last_field["lines"].append(line.strip())
    return title, intro, pages


def render(title, intro, pages, out_path):
    doc = Document()
    doc.styles["Normal"].font.name = FONT
    doc.styles["Normal"].font.size = Pt(11)

    # 文档标题
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(p, title, size=17, color=C_TITLE, base_bold=True)

    # 用法说明
    for t in intro:
        ip = doc.add_paragraph()
        add_inline(ip, t, size=8.5, color=C_META)

    for pg in pages:
        # 页头：横线 + 「第N页 · 页型」色标
        hp = doc.add_paragraph()
        hp.paragraph_format.space_before = Pt(10)
        add_top_rule(hp)
        if pg["num"] is not None:
            chip = f"第{pg['num']}页 · {pg['type']}"
        else:
            chip = "结束" if pg["type"] == "END" else pg["type"]
        run = hp.add_run(chip)
        set_font(run, size=10, color=C_TITLE, bold=True)
        # 标题（封面/END 可空）
        if pg["title"]:
            tp = doc.add_paragraph()
            tp.paragraph_format.space_before = Pt(0)
            add_inline(tp, pg["title"], size=13, color=C_BODY, base_bold=True)

        # 先合并回溯/节奏为一条元信息
        meta = []
        body_fields = []
        for f in pg["fields"]:
            if f["label"] in ("回溯", "节奏"):
                val = " ".join(f["lines"]).strip()
                if val:
                    meta.append((f["label"], val))
            else:
                body_fields.append(f)
        if meta:
            mp = doc.add_paragraph()
            for idx, (lab, val) in enumerate(meta):
                if idx > 0:
                    sep = mp.add_run("   ·   ")
                    set_font(sep, size=8.5, color=C_META)
                lr = mp.add_run(f"{lab} ")
                set_font(lr, size=8.5, color=C_META, bold=True)
                add_inline(mp, val, size=8.5, color=C_META)

        # 口播 / 兜底参考
        for f in body_fields:
            label = f["label"]
            lab_color = C_LABEL_KOU if label == "口播" else C_FALLBACK
            txt_color = C_BODY if label == "口播" else C_FALLBACK
            size = 11.5 if label == "口播" else 10
            first = True
            for ln in f["lines"]:
                bp = doc.add_paragraph()
                bp.paragraph_format.line_spacing = 1.25
                if first:
                    lr = bp.add_run(f"{label}　")
                    set_font(lr, size=size, color=lab_color, bold=True)
                    if f["note"]:
                        nr = bp.add_run(f"（{f['note']}）")
                        set_font(nr, size=size - 1, color=C_META)
                    first = False
                else:
                    bp.paragraph_format.left_indent = Pt(28)
                add_inline(bp, ln, size=size, color=txt_color)
            if first and f["note"]:
                bp = doc.add_paragraph()
                lr = bp.add_run(f"{label}　")
                set_font(lr, size=size, color=lab_color, bold=True)
                nr = bp.add_run(f"（{f['note']}）")
                set_font(nr, size=size - 1, color=C_META)

    doc.save(out_path)
    return len(pages)


def main():
    if len(sys.argv) < 2:
        print("用法: python lecture_notes_to_docx.py 讲稿.md [输出.docx]")
        sys.exit(1)
    md_path = sys.argv[1]
    out_path = sys.argv[2] if len(sys.argv) >= 3 else os.path.splitext(md_path)[0] + ".docx"
    with open(md_path, encoding="utf-8") as fh:
        md_text = fh.read()
    title, intro, pages = parse(md_text)
    n = render(title, intro, pages, out_path)
    print(f"saved -> {out_path}  （{n} 页）")


if __name__ == "__main__":
    main()

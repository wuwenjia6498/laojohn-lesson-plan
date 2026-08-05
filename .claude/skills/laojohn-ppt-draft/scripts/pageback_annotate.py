# -*- coding: utf-8 -*-
"""详案页标回注：把 `〖PPT 第N页 · 眉标 · 标题〗` 内联插进详案，供老师备课交叉对照。

规则见 `references/detail-pageback-annotation.md`（含「写作课线适配」节）。
本脚本只干机械活——**哪一页锚到哪句话是教学判断，由人填进映射表**。

两步用法：

    # ① 从已定稿的 pptx 生成待填骨架（page/kicker/title 已填好，anchor 留空）
    PYTHONUTF8=1 python pageback_annotate.py <详案.md> --from-pptx <课件.pptx> -o <映射.json>
    # ② 人工把每条的 anchor 填成详案里该页取材处的行首原文，然后回注
    PYTHONUTF8=1 python pageback_annotate.py <详案.md> <映射.json>

回注是**幂等的**：每次先删掉详案里所有既有 `〖PPT …〗` 独立行再重插，
所以分页变了、锚点改了都可以直接重跑。
"""
import argparse
import json
import os
import re
import sys

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

TAG_RE = re.compile(r"^〖PPT\s")


def slide_texts(slide):
    items = []
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            items.append(((sh.top or 0) // 500000, sh.left or 0, sh.text_frame.text.strip()))
    items.sort()
    return [t for _, _, t in items]


def _smart_quotes(root):
    """复用根目录 fix_quotes_md.py 的 smart_quotes_line（弯引号铁律的单一源，禁重写）。"""
    import importlib.util
    path = os.path.join(root, "fix_quotes_md.py")
    if not os.path.isfile(path):
        return lambda s: s
    spec = importlib.util.spec_from_file_location("fix_quotes_md", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod.smart_quotes_line


def tidy(s, curly=None):
    """压掉排版用的空格；中文之间的空格是 PPT 排版强调，进页标要去掉。

    外部平台产的文本常带 ASCII 直引号，而详案是弯引号铁律——落进页标前必须转。
    """
    s = s.replace("\n", " ")
    # 先转弯引号再压空格：直引号是 ASCII，压「中文间空格」的规则认不出 `把 "他很急"` 里那个空格
    if curly:
        s = curly(s)
    s = re.sub(r"(?<=[^\x00-\x7f])\s+(?=[^\x00-\x7f])", "", s)
    return re.sub(r"\s{2,}", " ", s).strip()


def short_kicker(kicker):
    """眉标只取分隔号后的末段——`附 · 习作讲评` 整个塞进去会把页标撑成四段。

    注意要在 tidy 之后按 `·` 判（tidy 会把中文间的空格压掉，` · ` 变 `·`）。
    """
    if re.match(r"^\d+$", kicker):
        return "课时分隔"                      # 眉标是纯数字 = 课时分隔页
    if "·" in kicker:
        return kicker.split("·")[-1].strip()
    return kicker


def build_skeleton(pptx_path, root):
    """按本线口径出待填骨架：封面与 END 不标，课时分隔页标 `课时分隔`。"""
    from pptx import Presentation
    curly = _smart_quotes(root)
    prs = Presentation(pptx_path)
    total = len(prs.slides._sldIdLst)
    rows = []
    for idx, slide in enumerate(prs.slides, 1):
        if idx == 1 or idx == total:                 # 封面 / END 无详案取材源
            continue
        t = slide_texts(slide)
        kicker = short_kicker(tidy(t[0], curly)) if t else ""
        title = tidy(t[1], curly) if len(t) > 1 else ""
        rows.append({"page": idx, "kicker": kicker, "title": title, "anchor": ""})
    return rows


def annotate(md_path, rows):
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().split("\n")

    before = len(lines)
    lines = [ln for ln in lines if not TAG_RE.match(ln.strip())]
    cleared = before - len(lines)
    # 清标会留下连续空行，压回单个
    squeezed = []
    for ln in lines:
        if ln.strip() == "" and squeezed and squeezed[-1].strip() == "":
            continue
        squeezed.append(ln)
    lines = squeezed

    plan = []
    for r in rows:
        anchor = (r.get("anchor") or "").strip()
        if not anchor:
            sys.exit("第 %s 页的 anchor 还没填——先把映射表补完再回注" % r.get("page"))
        hits = [i for i, ln in enumerate(lines) if ln.startswith(anchor)]
        if len(hits) != 1:
            sys.exit("第 %s 页锚点命中 %d 次（应为 1）：%s\n"
                     "  —— 把 anchor 写长一点直到唯一，或检查详案是否改过"
                     % (r["page"], len(hits), anchor[:60]))
        kicker = r.get("kicker", "").replace(" · ", "·")
        tag = "〖PPT 第%d页 · %s · %s〗" % (r["page"], kicker, r.get("title", ""))
        plan.append((hits[0], tag, r["page"]))

    rows_idx = [i for i, _, _ in plan]
    if rows_idx != sorted(rows_idx):
        bad = [p for (i, _, p), (j, _, q) in zip(plan, plan[1:]) if i > j]
        sys.exit("锚点在详案里的先后顺序与页序不一致（第 %s 页附近）——检查映射表" % bad)

    for idx, tag, _ in sorted(plan, key=lambda x: -x[0]):
        lines.insert(idx, tag)
        lines.insert(idx + 1, "")

    with open(md_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    return cleared, len(plan)


def main():
    ap = argparse.ArgumentParser(description="详案页标回注（幂等）")
    ap.add_argument("detail", help="详案 .md")
    ap.add_argument("mapping", nargs="?", help="映射 .json（回注时必给）")
    ap.add_argument("--from-pptx", default=None, help="据此 pptx 生成待填骨架，不回注")
    ap.add_argument("-o", "--out", default=None, help="骨架输出路径")
    ap.add_argument("--root", default=".", help="项目根目录（用于取 fix_quotes_md.py）")
    args = ap.parse_args()

    if args.from_pptx:
        rows = build_skeleton(args.from_pptx, os.path.abspath(args.root))
        out = args.out or (os.path.splitext(args.detail)[0] + "-页标映射.json")
        with open(out, "w", encoding="utf-8") as f:
            json.dump(rows, f, ensure_ascii=False, indent=1)
        print("骨架 %d 条（已跳过封面与 END）-> %s" % (len(rows), out))
        print("下一步：逐条填 anchor＝详案里该页取材处的行首原文，再跑回注。")
        print("  锚点按「翻页发生在开始讲这段时」定位——锚到那句师话，别等到表格或引文。")
        return

    if not args.mapping:
        sys.exit("要回注就给映射 .json；要生成骨架就加 --from-pptx")
    with open(args.mapping, encoding="utf-8") as f:
        rows = json.load(f)
    if isinstance(rows, dict):
        rows = rows.get("pages", [])

    cleared, n = annotate(args.detail, rows)
    print("清除旧页标 %d 行；回注 %d 个页标 -> %s" % (cleared, n, args.detail))
    print("详案已改动，记得重渲 docx（写作课须带 --header-left/--header-right）。")


if __name__ == "__main__":
    main()

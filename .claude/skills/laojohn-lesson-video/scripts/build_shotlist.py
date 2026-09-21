# -*- coding: utf-8 -*-
r"""备课视频链第 1 步：详案 + 教师用 json + 动画工作单 → 分镜单 JSON。**零模型，纯解析。**

    PYTHONUTF8=1 python build_shotlist.py <课次> [--dry-run] [--fallback-repo-pptx]

产出 `写作课备课视频输出\<课次>\<课次>-分镜单.json`，是下游全链的骨架。

三条不可推断的解析口径（踩过才知道，改这个脚本前先读）：

1. **页码权威源是详案内联页标，不是页标映射 json。**
   映射 json 在带跨页标记的课次会整体偏移：三上一实测映射记 page 2–24（23 条），
   而详案是 `2,…,18,19-20,21,…,25`、pptx 25 页——它把 `19-20` 记成一页，此后每页少 1。
   17 个课次里 9 个带跨页标记，且五上一展开了、三上一没展开，映射生成本身就不一致。
   所以 kicker/title 一律取详案页标里那两段（人工写的眉标，本就最准），映射 json
   只用来做交叉校验、产出偏移报告。
   ⚠ 但要配对着看：**详案页标对的是外部终稿**，`写作课件PPT输出\` 里那份不一定是终稿
   （外部会重新导入）。所以画面源要拿终稿；用 `--fallback-repo-pptx` 时
   `check_title_match` 报一片失配是**预期**，不是详案错了，别去改详案页标。

2. **末页区间遇 `## 附：` 必须截断。** 五上一实测：不截断的话，最后一个页标会把整个
   「附：习作讲评指导环节」1135 字全吞进来，那一页的旁白预算会炸掉。

3. **环节与 delivery.envs 不是严格 1:1。** 五上一详案 8 个环节、envs 只有 6 条，且末条
   写作 `⑥–⑦` 合并。对不上时标 match=null 列进报告交人工填，**不许静默猜**——猜错的
   后果是整页旁白讲的是别的环节的设计意图，而这事机器再也发现不了。
"""
import argparse
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from video_link import Sources, project_root  # noqa: E402

# 〖PPT 第3页 · 创设情境 · 先做一个游戏——猜猜他是谁〗 / 第19-20页 跨页
PAGE_TAG = re.compile(r"^〖PPT\s*第\s*(\d+)(?:\s*[-–—]\s*(\d+))?\s*页\s*·\s*(.+?)\s*·\s*(.+?)〗\s*$")
LESSON_H = re.compile(r"^##\s+(?!#)(第.课时.*)$")
APPENDIX_H = re.compile(r"^##\s+(?!#)(附：.*)$")
STEP_H = re.compile(r"^###\s+(.+?)\s*$")
MINUTES = re.compile(r"（约\s*(\d+)\s*分钟）")
# envs 的 heading 形如「① 情境导入（旧包的痕迹）」「⑥–⑦ 当堂写作与同伴评改」
ENV_NUM = re.compile(r"^([①-⑳])(?:\s*[-–—]\s*([①-⑳]))?\s*(.*)$")
CIRCLED = "①②③④⑤⑥⑦⑧⑨⑩⑪⑫"
BOLD_MARK = re.compile(r"\{/?b\}")

# 每页旁白字数预算：按页型定基准，再按信息量微调（见 tune_budget）
# 0921 按两个课次的实写量回调。原版 model=290 是拍的，实测下来：
# 模型在 model 页自然写 150–250 字，把预算抬到 290 只会逐页报「偏离预算」、
# 机检跟着麻木。下调后六上四 raw 总量落进区间，归一化不再触发。
# ⚠ cover/lesson_split 要与 write_narration 里告诉模型的字数对得上：
# prompt 说「全页 90–160 字」而预算表给 90，模型写 144 就被报「偏离预算」——
# 两处口径不一致，机检就在报自己。
BUDGET = {"cover": 80, "lesson_split": 130, "teach": 175, "model": 230,
          "write": 195, "review": 205, "appendix": 150, "outro": 250}
# kicker（详案页标里的页型词）→ 分镜 kind
KIND_BY_KICKER = [
    ("课时分隔", "lesson_split"),
    ("学写法", "model"), ("判断对比", "model"),
    ("教师示范文", "model"), ("读示范文", "model"),
    ("旁批表", "model"), ("方法", "model"),
    ("自由写作", "write"), ("写作热身", "write"),
    ("定下起点", "write"), ("口头练笔", "write"),
    ("交流评议", "review"), ("修改收束", "review"),
    ("收束", "review"),
    ("习作讲评", "appendix"), ("讲评", "appendix"),
]


def kind_of(kicker):
    for key, kind in KIND_BY_KICKER:
        if key in kicker:
            return kind
    return "teach"


def strip_bold(s):
    return BOLD_MARK.sub("", s or "").strip()


def parse_plan(path):
    """切详案：返回 (shots, meta)。每个 shot 带 pages / kicker / title / 四桶原文 / 环节归属。"""
    with open(path, encoding="utf-8") as f:
        lines = f.read().split("\n")

    # 先定位所有结构行，附录起始行之后的页标仍要（第25页就是讲评页），
    # 但正文归集必须在附录标题处停——否则末页吞掉整个附录章节。
    appendix_at = None
    for i, ln in enumerate(lines):
        if APPENDIX_H.match(ln):
            appendix_at = i
            break

    marks = []
    for i, ln in enumerate(lines):
        m = PAGE_TAG.match(ln.strip())
        if m:
            marks.append((i, m))

    lesson, step, step_min, step_idx = 0, None, None, 0
    steps_seen = []
    # 逐行扫一遍，记录每行所处的课时/环节，供页标归属
    ctx = []
    for i, ln in enumerate(lines):
        lm = LESSON_H.match(ln)
        if lm:
            lesson += 1
            step, step_min, step_idx = None, None, 0
        elif APPENDIX_H.match(ln):
            lesson = 99
            step, step_min, step_idx = None, None, 0
        else:
            sm = STEP_H.match(ln)
            if sm:
                raw = sm.group(1)
                mm = MINUTES.search(raw)
                step = MINUTES.sub("", raw).strip()
                step_min = int(mm.group(1)) if mm else None
                step_idx += 1
                if lesson != 99:
                    steps_seen.append((lesson, step_idx, step, step_min))
        ctx.append((lesson, step, step_min, step_idx))

    shots = []
    for k, (i, m) in enumerate(marks):
        p1 = int(m.group(1))
        p2 = int(m.group(2)) if m.group(2) else p1
        pages = list(range(p1, p2 + 1))
        stop = marks[k + 1][0] if k + 1 < len(marks) else len(lines)
        # 口径 2：正文归集不得越过附录标题
        if appendix_at is not None and i < appendix_at < stop:
            stop = appendix_at
        body = lines[i + 1:stop]

        buckets = {"teacher_lines": [], "stage": [], "materials": [],
                   "refs": [], "interaction": [], "quote": []}
        for ln in body:
            s = ln.strip()
            if not s or s.startswith("#") or s.startswith("【图位:"):
                continue
            if s.startswith("师："):
                buckets["teacher_lines"].append(s)
            elif s.startswith("["):
                buckets["stage"].append(s)
            elif s.startswith("〔"):
                buckets["materials"].append(s)
            elif s.startswith("参考："):
                buckets["refs"].append(s)
            elif s in ("学生互动分享", "学生自由分享",
                       "学生动笔写作"):
                buckets["interaction"].append(s)
            elif s.startswith(">"):
                buckets["quote"].append(s.lstrip("> ").strip())

        lesson_i, step_i, min_i, idx_i = ctx[i]
        # 页标常写在 ### 环节标题的前一两行（先出这一页、再进那个环节），这时该归到后面那个环节。
        # 但**只认紧贴的**：页标与 ### 之间除空行外还有正文，说明这一页属于上一个环节、
        # 下一个 ### 只是碰巧落在本页区间内。不设这道限制的话，第10页（学写法）会被归到
        # 13 行之后的「四、人物特点卡」——旁白就会拿错环节的设计意图去讲。
        # 探测范围要比正文范围多一行：末页的 stop 已被附录标题截断（口径 2），
        # 用 stop 当上界就永远看不到 `## 附：` 那一行，讲评页会挂回第 2 课时的末环节。
        for look in range(i + 1, min(stop + 1, len(lines))):
            raw = lines[look].strip()
            if not raw:
                continue
            # 认 ### 环节，也认 ## 第N课时 / ## 附：——课时分隔页与讲评页正是紧贴在
            # 这两种标题之前的，不认就会挂到上一课时的最后一个环节上。
            if (STEP_H.match(lines[look]) or LESSON_H.match(lines[look])
                    or APPENDIX_H.match(lines[look])):
                lesson_i, step_i, min_i, idx_i = ctx[look]
            break

        kicker, title = m.group(3).strip(), m.group(4).strip()
        shots.append({
            "id": "s%02d" % p1,
            "pages": pages,
            "line": i + 1,
            "kind": kind_of(kicker),
            "lesson": lesson_i,
            "kicker": kicker,
            "title": title,
            "step": {"name": step_i, "minutes": min_i, "index": idx_i},
            "source": buckets,
            "chars_shihua": sum(len(x) for x in buckets["teacher_lines"]),
        })

    if shots:
        shots[0]["kind"] = shots[0]["kind"]
    return shots, {"steps": steps_seen, "appendix_line": appendix_at}


def attach_envs(shots, teacher, report):
    """把 delivery.envs 的 keep/free/why 按环节挂到 shot 上。对不上标 match=null，不猜。"""
    envs = (teacher.get("delivery") or {}).get("envs") or []
    parsed = []
    for i, e in enumerate(envs):
        h = e.get("heading", "")
        m = ENV_NUM.match(h)
        nums, name = [], h
        if m:
            nums = [CIRCLED.index(m.group(1)) + 1]
            if m.group(2):
                nums = list(range(nums[0], CIRCLED.index(m.group(2)) + 2))
            name = m.group(3).strip()
        rows = {"keep": [], "free": [], "why": []}
        for r in e.get("rows", []):
            t = r.get("type")
            if t in rows:
                rows[t].append(strip_bold(r.get("text")))
        parsed.append({"i": i, "heading": h, "nums": nums, "name": name, "rows": rows})

    # 详案环节的全局序号（跨课时连续），与 envs 的 ①②③ 对应
    seq = {}
    n = 0
    for sh in shots:
        key = (sh["lesson"], sh["step"]["index"], sh["step"]["name"])
        if sh["lesson"] == 99 or not sh["step"]["name"]:
            continue
        if key not in seq:
            n += 1
            seq[key] = n

    for sh in shots:
        key = (sh["lesson"], sh["step"]["index"], sh["step"]["name"])
        gi = seq.get(key)
        hit, how = None, None
        if gi:
            for e in parsed:
                if gi in e["nums"]:
                    hit, how = e, ("exact" if len(e["nums"]) == 1 else "merged")
                    break
            if hit is None:
                sname = (sh["step"]["name"] or "").replace("、", "")
                sname = re.sub(r"^[一二三四五六七八九十]+", "", sname)
                for e in parsed:
                    a = set(sname) & set(e["name"])
                    if len(a) >= 3:
                        hit, how = e, "fuzzy"
                        break
            if hit is None and parsed and len(parsed[-1]["nums"]) > 1 and gi > max(parsed[-1]["nums"]):
                # 末条 envs 是合并条（如「⑥–⑦ 当堂写作 + 评改」），而详案第 2 课时实际有 3 个环节
                # ——配套 json 这一条本就是「整个第 2 课时」的旁注，写少了一个序号而已。
                # 延伸到末环节、标 merged-tail 并记进报告，比整页空着强；留痕让人工能看见。
                hit, how = parsed[-1], "merged-tail"
                report.setdefault("env_merged_tail", []).append(
                    {"shot": sh["id"], "step": sh["step"]["name"],
                     "env": parsed[-1]["heading"]})
        sh["env"] = ({"index": hit["i"], "heading": hit["heading"],
                      "minutes": sh["step"]["minutes"], "match": how}
                     if hit else
                     {"index": None, "heading": None,
                      "minutes": sh["step"]["minutes"], "match": None})
        sh["delivery"] = hit["rows"] if hit else {"keep": [], "free": [], "why": []}
        # 课时分隔/片头片尾本就不属于任何环节，不算失配
        if hit is None and sh["lesson"] != 99 and sh["kind"] not in ("lesson_split", "cover", "outro"):
            report.setdefault("env_unmatched", []).append(
                {"shot": sh["id"], "step": sh["step"]["name"], "title": sh["title"][:20]})


def attach_warns(shots, teacher, report):
    """warns/nono 是整课级的，不是页级。关键词命中挂页，挂不上的归到该课时末页。
    这步允许粗糙——生成旁白时模型还会拿到整课全表。"""
    warns = [strip_bold(w) for w in (teacher.get("overview") or {}).get("warns", [])]
    nono = [strip_bold(w) for w in (teacher.get("delivery") or {}).get("nono", [])]
    for sh in shots:
        sh["warns"], sh["nono"] = [], []
    teach_shots = [s for s in shots if s["lesson"] != 99]
    unplaced = {}

    def place(items, field):
        for w in items:
            best, score = None, 0
            # ① 文本里常直接点名「环节⑤」——这是最硬的指向，优先认
            em = re.search(r"环节\s*([①-⑳])", w)
            if em:
                want = CIRCLED.index(em.group(1))
                for sh in teach_shots:
                    if (sh.get("env") or {}).get("index") == want:
                        best, score = sh, 99
                        break
            # ② 引号里的短语（含直引号——教师用 json 里两种引号都出现过）
            if best is None:
                keys = re.findall(r"[“「\"]([^”」\"]{2,12})[”」\"]", w)
                for sh in teach_shots:
                    blob = sh["title"] + "".join(sh["source"]["teacher_lines"])
                    hit = sum(1 for k in keys if k in blob)
                    if hit > score:
                        best, score = sh, hit
            # 没有高置信度线索就**不挂**。曾按「冒号前那截与眉标的字面重合」退化匹配，
            # 结果把「开场三轮猜人」挂到了学写法页（共用了「三、猜、人」三个字）——
            # 错挂会让旁白在错的页讲错的提醒，比不挂更糟。挂不上的进整课全表，
            # 生成旁白时模型拿得到全表，自己判断哪条与本页相关。
            if best is not None and score > 0:
                best[field].append(w)
            else:
                unplaced.setdefault(field, []).append(w)

    place(warns, "warns")
    place(nono, "nono")
    # 整课全表随分镜单一起交给模型：页级归属只是提示，全表才是事实源
    report["course_warns"] = warns
    report["course_nono"] = nono
    if unplaced:
        report["unplaced"] = {k: len(v) for k, v in unplaced.items()}
    return {"warns": warns, "nono": nono}


def drop_appendix(shots, report, keep=False):
    """把「附：习作讲评」那几页从分镜里拿掉。

    **讲评环节只存在于详案，实际 PPT 里是删掉的**（0921 用户告知）。
    详案给它留了页标，于是分镜会把它当成一页真实的 PPT——
    三上一实测过一次：详案标 P25 是讲评页，而 PPT 的 P25 其实是 THE END，
    成片里就出现了「THE END 配着讲评环节旁白」。

    拿掉它还顺手修好了尾页对齐：腾出来的那一页会回到 tail_unused，
    由 outro 认领，正好对上 THE END。

    `--keep-appendix` 给“确实做了讲评页”的课次留一个口子
    （仓内版有些课次是有那一页的，比如五上一 P27）。"""
    if keep:
        return shots
    dropped = [s for s in shots if s["kind"] == "appendix"]
    if dropped:
        report["appendix_dropped"] = [
            {"shot": s["id"], "page": s["pages"], "title": s["title"][:24]} for s in dropped]
        report["appendix_note"] = (
            "讲评环节只在详案里、实际 PPT 已删，"
            "已从分镜移除（要保留用 --keep-appendix）")
    return [s for s in shots if s["kind"] != "appendix"]


def add_bookend_shots(shots, ac, sl_meta, teacher, anim=None):
    """给片头封面页与片尾页补上 cover / outro 两个 shot。

    详案只从 P2 起有页标，封面与 THE END 天然没有 shot。不补的话这两页只能放静场——
    片子一开场是六秒无声的封面，老师不知道自己打开的是什么、要看多久。
    cover 的旁白正是交代「这是哪一课的备课视频、多长、怎么用」的唯一位置。"""
    goal = re.sub(r"\{/?b\}", "", ((teacher.get("overview") or {}).get("goal") or ""))
    head = ac.get("head_unused") or []
    tail = ac.get("tail_unused") or []
    slides = ((anim or {}).get("slides") or {})

    def face(pages, fallback):
        """取那一页 PPT 上真正印着的字。不取课次标识段——那是文件名，
        模型会照着念出「三上-第一单元-XXX」，而封面上根本没这行。"""
        info = slides.get(str(pages[0])) if pages else None
        if info:
            bits = [x for x in (info.get("kicker"), info.get("title")) if x]
            t = " · ".join(b.replace(chr(10), " ").strip() for b in bits)
            if t:
                return t[:60]
        return fallback

    if head:
        shots.insert(0, {
            "id": "s00", "pages": head, "line": 0, "kind": "cover", "lesson": 0,
            "kicker": "片头", "title": face(head, sl_meta),
            "step": {"name": None, "minutes": None, "index": 0},
            "source": {"teacher_lines": [], "stage": [], "materials": [],
                       "refs": [], "interaction": [],
                       "quote": [goal] if goal else []},
            "chars_shihua": 0,
            "visual": {"slides": head, "frames": ["frames/p%02d.png" % x for x in head],
                       "clicks": 0, "ppt_title": "", "copyright": "unknown"},
            "env": {"index": None, "heading": None, "minutes": None, "match": None},
            "delivery": {"keep": [], "free": [], "why": []}, "warns": [], "nono": [],
        })
    if tail:
        shots.append({
            "id": "s99", "pages": tail, "line": 0, "kind": "outro", "lesson": 0,
            "kicker": "片尾", "title": face(tail, sl_meta),
            "step": {"name": None, "minutes": None, "index": 0},
            "source": {"teacher_lines": [], "stage": [], "materials": [],
                       "refs": [], "interaction": [],
                       "quote": [goal] if goal else []},
            "chars_shihua": 0,
            "visual": {"slides": tail, "frames": ["frames/p%02d.png" % x for x in tail],
                       "clicks": 0, "ppt_title": "", "copyright": "unknown"},
            "env": {"index": None, "heading": None, "minutes": None, "match": None},
            "delivery": {"keep": [], "free": [], "why": []}, "warns": [], "nono": [],
        })


def tune_budget(shots, total_target=None):
    """按页型定基准，信息量多的页加码，只在超出片长区间时才归一化。

    ⚠ **总量不再以师话字数为锚（0921 改）**。旧版把全片归一化到
    「师话字数 × 0.93」，而张祖庆视角报出的病灯正好是
    **「旁白总字数是详案师话的 94%」**——拿师话字数当锚，
    等于从预算层就把「陪老师读一遍详案」写死了。
    片长该由**增量素材的多少**决定：卡点多的课就该长，素材薄的课就该短。
    所以现在只在算出来的总量溢出 4300–6800 区间时才按比例拉回。
    `--target-chars` 仍可手动指定（要压片长时用）。
    """
    FILM_LO, FILM_HI = 4300, 6800
    for sh in shots:
        base = BUDGET.get(sh["kind"], 190)
        extra = 0
        d = sh.get("delivery") or {}
        # 信息量要把**卡点素材**算进来。0921 六上四实测：
        # 判据改成「risk 有素材就必须写」之后，`参考：` 密集的页必然写得长，
        # 而旧公式只数 keep/why/warns，结果 29 页里 10 页报「偏离预算」——
        # **不是旁白写长了，是预算没数它们。**
        n_info = (len(d.get("keep", [])) + len(d.get("why", []))
                  + len(sh.get("warns", [])) + len(sh.get("nono", []))
                  + len(sh["source"].get("refs") or [])
                  + len([m for m in (sh["source"].get("materials") or [])
                         if m.startswith("〔应答")]))
        # ⚠ 加码得轻——素材条数多不等于这一页能讲那么多。
        # 0921 第一版给 20×/上限 90，s08–s13 六页预算被推到 320–345，
        # 而实写只有 137–204——**不是写短了，是预算虚高**。
        extra += min(60, 12 * max(0, n_info - 1))
        if (sh.get("visual") or {}).get("clicks", 0) >= 5:
            extra += 20
        sh["budget_chars"] = base + extra
    tot = sum(s["budget_chars"] for s in shots) or 1
    if total_target:
        k = total_target / float(tot)
    elif tot > FILM_HI:
        k = FILM_HI / float(tot)
    elif tot < FILM_LO:
        k = FILM_LO / float(tot)
    else:
        k = 1.0
    for sh in shots:
        sh["budget_chars"] = int(round(sh["budget_chars"] * k / 5.0) * 5)
    return sum(s["budget_chars"] for s in shots)


def slides_from_pptx(path):
    """直接从画面源 pptx 抽每页的眉标与标题，形状同 anim.json 的 slides。

    ⚠ **有外部终稿时必须用这个，不能用 anim.json**。
    anim.json 是本仓那份 pptx 的快照，而终稿是外部平台后改的——两边页数就不一样。
    0921 五上四实测：终稿 27 页、anim 26 页，从 P15 起整体错一页，
    拿 anim 比对就报出 9 条「重合 0%」——**那不是详案错了，是拿旧快照量新片子**。

    版面规律（两条线都实测过）：按 top 升序，第一个非空文本框是眉标（top≈610000），
    第二个是标题；课时分隔页的第一个是大号数字（「02」），第二个仍是标题。
    """
    from pptx import Presentation
    out = {}
    prs = Presentation(path)
    for i, sl in enumerate(prs.slides, 1):
        txt = []
        for sh in sl.shapes:
            if not sh.has_text_frame:
                continue
            t = (sh.text_frame.text or "").strip()
            if t:
                txt.append((sh.top if sh.top is not None else 0, t))
        txt.sort(key=lambda x: x[0])
        head = [t for _, t in txt[:2]]
        out[str(i)] = {
            "kicker": head[0] if head else "",
            "title": head[1] if len(head) > 1 else "",
            "groups": [], "skip": False, "existing_clicks": 0,
        }
    return {"slides": out}


def attach_visual(shots, anim, slide_count, report):
    """anim.json 的 slides 是「页号字符串 → {groups, kicker, title, skip, existing_clicks}」，
    页号是 PPT 真实页号，拿它补点击数与页面标题，用来判断这一页信息量。"""
    slides = (anim or {}).get("slides") or {}
    for sh in shots:
        p = sh["pages"][0]
        info = slides.get(str(p)) or {}
        clicks = info.get("existing_clicks") or len(info.get("groups") or [])
        sh["visual"] = {
            "slides": sh["pages"],
            "frames": ["frames/p%02d.png" % x for x in sh["pages"]],
            "clicks": clicks,
            "ppt_title": (info.get("title") or "").replace("\n", " ")[:60],
            "copyright": "unknown",
        }
        for x in sh["pages"]:
            if x > slide_count:
                report.setdefault("page_out_of_range", []).append(
                    {"shot": sh["id"], "page": x, "slide_count": slide_count})


def absorb_unlabeled(shots, anim, slide_count, report):
    """详案没给页标的页，若与前一个 shot 眉标相同，就并进去。

    写作课件常有「同一环节跨两页」的情形（比如自由写作的第二页只剩一个眉标、
    学生在写），详案不会为它单独标一个页标。不并的话那一页就被整页跳过，
    成片里直接不出现——而老师手里的课件是有那一页的。
    眉标不同的不并，那是真漏标，继续报。
    """
    slides = (anim or {}).get("slides") or {}
    taken = set()
    for sh in shots:
        taken.update(sh["pages"])
    absorbed = []
    for pg in range(1, slide_count + 1):
        if pg in taken:
            continue
        prev = [sh for sh in shots if sh["pages"] and max(sh["pages"]) == pg - 1]
        if not prev:
            continue
        k_here = (slides.get(str(pg)) or {}).get("kicker") or ""
        k_prev = (slides.get(str(pg - 1)) or {}).get("kicker") or ""
        if k_here and k_here == k_prev:
            prev[0]["pages"].append(pg)
            taken.add(pg)
            absorbed.append({"page": pg, "into": prev[0]["id"], "kicker": k_here})
    if absorbed:
        report["absorbed_pages"] = absorbed
    return absorbed


def check_title_match(shots, anim, report, source_kind="final"):
    r"""详案页标的眉标 vs 那一页 PPT 上实际印的字，对不上就报。

    **这道比 align_check 更要紧**：align_check 只数页数，数得对不代表对得上。
    三上一实测——详案说 P25 是「讲评环节（可选）」，而 PPT 第 25 页其实是 THE END
    （讲评是可选环节，做课件时压根没做那一页）。页数正好对得上，于是一路静默，
    直到成片里 THE END 那页配着讲评环节的旁白才被看出来。

    anim.json 的 kicker/title 抽自 PPT 本身，是这件事唯一的机读凭据。
    只报「完全不沾边」的，不报措辞差异——anim 的 title 有时是截断的正文，不能苛求。

    ⚠ **失配怎么读，完全取决于画面源是哪一份**（0921 用户当场纠正过一次误判）：
    详案页标是对着**外部终稿**回注的，而 `写作课件PPT输出\` 里那份不一定是终稿。
    用 --fallback-repo-pptx 时大片失配是**预期**（版本不同），不是详案错了，
    **更不要据此去改详案页标**；用终稿时才是真信号。
    """
    slides = (anim or {}).get("slides") or {}
    bad = []
    for sh in shots:
        if sh["kind"] in ("cover", "outro"):      # 这两个本就是本脚本给结构页补的
            continue
        p = sh["pages"][0]
        info = slides.get(str(p))
        if not info:
            continue
        ppt_txt = (info.get("kicker") or "") + (info.get("title") or "")
        plan_txt = sh["kicker"] + sh["title"]
        # ① note 是最硬的凭据：标了封面/末页的，是没有教学内容的结构页，
        #    详案页标认领到它身上，基本就是详案标错了页。
        note = info.get("note") or ""
        if ("封面" in note or "末页" in note or "END" in note):
            bad.append({"shot": sh["id"], "page": p,
                        "详案页标": plan_txt[:26],
                        "PPT 实际": (note + "：" + ppt_txt.replace("\n", " "))[:30],
                        "重合": "结构页"})
            continue
        # ② 标题比对兜底。⚠ anim 的 kicker/title 抽自版面，常抓到编号之类的噪声
        #    （实测 P24 抓成了「3」「1」），太短或纯数字的一律不比，免得误报。
        clean = re.sub(r"[\s\d]", "", ppt_txt)
        if len(clean) < 4:
            continue
        a, b = set(clean), set(re.sub(r"[\s\d]", "", plan_txt))
        if not a or not b:
            continue
        overlap = len(a & b) / float(min(len(a), len(b)))
        if overlap < 0.25:
            bad.append({"shot": sh["id"], "page": p,
                        "详案页标": plan_txt[:26],
                        "PPT 实际": ppt_txt.replace("\n", " ")[:26],
                        "重合": "%.0f%%" % (overlap * 100)})
    if bad:
        report["title_mismatch"] = bad
        if source_kind == "repo":
            report["title_mismatch_note"] = (
                "这 %d 页与仓内 pptx 对不上——"
                "用 --fallback-repo-pptx 时这是**预期**："
                "详案页标对的是外部终稿，仓内这份可能是更早版本。"
                "**别据此改详案页标**；要真对齐就把终稿拷进工作区" % len(bad))
        else:
            report["title_mismatch_note"] = (
                "这 %d 页的详案页标与**终稿** PPT 对不上——"
                "**旁白会讲错页**，这是真信号，逐页看一遍再往下走" % len(bad))


def check_pagemap(shots, pagemap, report):
    """页标映射 json 只作交叉校验：逐条比对它的 page 与详案页标，差异全部记进报告。"""
    if not pagemap:
        report["pagemap"] = "缺页标映射 json，跳过交叉校验"
        return
    by_title = {}
    for e in pagemap:
        by_title.setdefault((e.get("kicker", ""), e.get("title", "")), e)
    off = []
    for sh in shots:
        e = by_title.get((sh["kicker"], sh["title"]))
        if e and e.get("page") != sh["pages"][0]:
            off.append({"shot": sh["id"], "plan_page": sh["pages"][0],
                        "pagemap_page": e.get("page"), "title": sh["title"][:24]})
    if off:
        report["pagemap_offset"] = off
        report["pagemap_note"] = (
            "映射 json 与详案页标不一致 %d 处——**以详案为准**（映射不展开跨页标记时会整体偏移）"
            % len(off))


def align_check(shots, slide_count, source_kind):
    """详案覆盖页 + 封面 + 片尾 == pptx 实际页数？对不上只报差异请人确认，不硬失败。
    这道校验因选用外部终稿而必设：外部人工嵌图时加页/拆页是常态，而失配的唯一表现是
    「旁白讲的和画面不是同一页」——不校验就只能靠看完 20 分钟片子才发现。"""
    covered = sorted({p for s in shots for p in s["pages"]})
    lo, hi = (covered[0], covered[-1]) if covered else (0, 0)
    missing = [p for p in range(lo, hi + 1) if p not in covered]
    head = list(range(1, lo))
    tail = list(range(hi + 1, slide_count + 1))
    over = [p for p in covered if p > slide_count]
    ok = not missing and not over and len(head) <= 1 and len(tail) <= 1
    why = []
    if missing:
        why.append("详案漏标 P%s —— 这几页 PPT 没有对应页标，视频里会被跳过"
                   % ",".join(str(p) for p in missing))
    if over:
        why.append("详案标到 P%d 但画面源只有 %d 页 —— 页标与这份 pptx 不是同一版"
                   % (max(over), slide_count))
    if len(head) > 1:
        why.append("片头 %d 页（P%s）无旁白" % (len(head), ",".join(str(p) for p in head)))
    if len(tail) > 1:
        why.append("片尾 %d 页（P%s）无旁白" % (len(tail), ",".join(str(p) for p in tail)))
    return {"result": "pass" if ok else "review",
            "source": source_kind,
            "plan_pages": len(covered), "covered": [lo, hi],
            "pptx_slides": slide_count,
            "head_unused": head, "tail_unused": tail, "gaps": missing, "over": over,
            "confirmed_by": None,
            "note": ("详案覆盖 P%d–P%d，片头 %d 页、片尾 %d 页未被旁白覆盖（正常：封面 + THE END）"
                     % (lo, hi, len(head), len(tail))) if ok else
                    "；".join(why) + "。确认无误后把确认人写进 confirmed_by 再往下走"}


def main():
    ap = argparse.ArgumentParser(description="备课视频分镜单（零模型）")
    ap.add_argument("unit", help="课次标识，如 三上-第一单元-猜猜他是谁")
    ap.add_argument("--dry-run", action="store_true", help="只打印归属表，不落盘")
    ap.add_argument("--fallback-repo-pptx", action="store_true",
                    help="外部终稿缺失时改用仓内 anim 版出画面")
    ap.add_argument("--keep-appendix", action="store_true",
                    help="保留「附：习作讲评」页（缺省移除：实际 PPT 里没有这一页）")
    # ⚠ 默认必须是 None（同 CLAUDE.md §3 那条红线）——给硬默认会永久盖住
    # tune_budget 里「按增量素材定片长」的新逻辑，而且完全静默。
    ap.add_argument("--target-chars", type=int, default=None)
    a = ap.parse_args()

    root = project_root()
    src = Sources.of(a.unit, root)
    report = {}

    shots, pmeta = parse_plan(src.plan_md)
    if not shots:
        raise SystemExit("详案里没有 〖PPT 第N页 …〗 页标——该课次未做页标回注，"
                         "先跑 laojohn-ppt 后处理链第 4 步 pageback_annotate.py")

    visual_path, source_kind = src.visual_pptx(fallback=a.fallback_repo_pptx)
    from pptx import Presentation
    slide_count = len(Presentation(visual_path).slides)

    anim = src.load_json("anim_json")
    # 画面源是外部终稿时，眉标／标题改从终稿本人抽（见 slides_from_pptx 头注释）。
    # anim.json 仍留给 add_bookend_shots 取封面文字以外的场合，但页级比对一律以终稿为准。
    if source_kind == "final":
        try:
            anim = slides_from_pptx(visual_path)
        except Exception as e:                      # noqa: BLE001
            print("⚠ 从终稿抽标题失败，回落 anim.json：%s" % e)
    teacher = src.load_json("teacher_json")
    attach_visual(shots, anim, slide_count, report)
    attach_envs(shots, teacher, report)
    course_level = attach_warns(shots, teacher, report)
    # 先丢讲评页再做两项比对——它注定不进片，留着只会在报告里刷一条注定的失配
    shots = drop_appendix(shots, report, keep=a.keep_appendix)
    absorb_unlabeled(shots, anim, slide_count, report)
    attach_visual(shots, anim, slide_count, report)   # 并页后重算 frames
    check_title_match(shots, anim, report, source_kind)
    check_pagemap(shots, src.load_json("pagemap"), report)
    ac = align_check(shots, slide_count, source_kind)
    add_bookend_shots(shots, ac, sl_meta=a.unit, teacher=teacher, anim=anim)
    total = tune_budget(shots, a.target_chars)

    doc = {
        "meta": {
            "course": a.unit,
            "visual_source": {"kind": source_kind, "path": os.path.relpath(visual_path, root)},
            "slide_count": slide_count,
            "align_check": ac,
            "rate_cps": 5.29,      # 火山渊博小叔 speed=1.0 实测；换通道后拿配音清单的实测值回写
            "budget_chars_total": total,
            "est_minutes": round(total / 4.7 / 60.0 + 1.5, 1),
            "shihua_chars_total": sum(s["chars_shihua"] for s in shots),
            "course_level": course_level,
            "report": report,
            **src.digest(),
        },
        "shots": shots,
    }

    print("课次 %s ｜ 画面源 %s ｜ pptx %d 页 ｜ 分镜 %d 个"
          % (a.unit, source_kind, slide_count, len(shots)))
    print("%-5s %-6s %-8s %-28s %-5s %-6s %s"
          % ("id", "页", "页型", "眉标", "预算", "env", "环节"))
    for sh in shots:
        pg = "-".join(str(x) for x in sh["pages"])
        env = sh["env"]["match"] or "✗"
        print("%-5s %-6s %-8s %-28s %-5d %-6s %s"
              % (sh["id"], pg, sh["kind"], sh["title"][:26], sh["budget_chars"],
                 env, sh["step"]["name"] or "-"))
    print("\n师话原文 %d 字（照念约 %.1f 分钟）→ 旁白预算 %d 字（约 %.1f 分钟）"
          % (doc["meta"]["shihua_chars_total"], doc["meta"]["shihua_chars_total"] / 4.7 / 60,
             total, total / 4.7 / 60))
    print("align_check: %s — %s" % (ac["result"], ac["note"]))
    for k, v in report.items():
        if k.startswith("course_"):          # 整课全表，不是问题，不刷屏
            continue
        print("  [%s] %s" % (k, v if not isinstance(v, list) else "%d 条" % len(v)))
        if isinstance(v, list):
            for item in v[:6]:
                print("      %s" % (item,))

    if a.dry_run:
        print("\n--dry-run，未落盘")
        return
    os.makedirs(src.out_dir, exist_ok=True)
    out = os.path.join(src.out_dir, a.unit + "-分镜单.json")
    with open(out, "w", encoding="utf-8", newline="\n") as f:
        json.dump(doc, f, ensure_ascii=False, indent=2)
    print("\n✓ %s" % os.path.relpath(out, root))


if __name__ == "__main__":
    main()
